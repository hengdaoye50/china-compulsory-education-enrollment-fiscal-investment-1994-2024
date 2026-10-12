"""Audit free denominator evidence without reconstructing published fiscal ratios."""
import json, hashlib, re
from decimal import Decimal as D
from pathlib import Path
from docx import Document
from pypdf import PdfReader
from openpyxl import load_workbook
from parse import ROOT, save_csv

names=['finance_denominator_sources','finance_denominator_followup_sources','finance_denominator_attachment_sources','finance_denominator_alternate_sources','finance_denominator_form_sources']
sources=[r for name in names for r in json.loads((ROOT/(name+'.json')).read_text(encoding='utf-8'))]
for r in sources:
    if r['status']=='downloaded':
        p=ROOT/r['path']; assert hashlib.sha256(p.read_bytes()).hexdigest()==r['sha256']
        if p.suffix=='.pdf':
            pdf=PdfReader(p)
            (ROOT/'data'/(r['key']+'_text.txt')).write_text('\n\n'.join(page.extract_text() for page in pdf.pages),encoding='utf-8')
save_csv(ROOT/'data/finance_denominator_source_receipt.csv',[{k:r.get(k) for k in ['key','url','status','path','bytes','sha256','error']} for r in sources])
doc=Document(ROOT/'raw/finance_denominator_moe_2020_docx.docx')
t='\n'.join(p.text for p in doc.paragraphs)
assert '农村指镇区' in t and '与乡村' in t and '初中包括普通初中和职业初中' in t
fiscal_definitions=t[t.index('107.'):t.index('110.')]
assert '年初' not in fiscal_definitions and '年末' not in fiscal_definitions
save_csv(ROOT/'data/finance_denominator_definition_checks.csv',[
    {'source':'https://hudong.moe.gov.cn/srcsite/A03/s182/202101/W020210113546624017867.docx','locator':'说明1','finding':'2020 indicator system defines rural as town plus village, city as urban district','status':'verified for this reference indicator system; not all annual fiscal forms'},
    {'source':'https://hudong.moe.gov.cn/srcsite/A03/s182/202101/W020210113546624017867.docx','locator':'说明2','finding':'default junior includes ordinary and vocational junior; individual formula must also be checked','status':'verified; do not substitute default for ordinary junior fiscal series'},
    {'source':'https://hudong.moe.gov.cn/srcsite/A03/s182/202101/W020210113546624017867.docx','locator':'107-109','finding':'in-school student denominator is named without explicit year-start/year-end weights','status':'weights unresolved; no fiscal ratio reconstruction'},
])
wb=load_workbook(ROOT/'raw/finance_denominator_zhejiang_2023_xlsx.xlsx',data_only=True)
inventory=[]
for s in wb:
    labels=[str(v) for row in s.iter_rows(values_only=True) for v in row if isinstance(v,str)]
    inventory.append({'sheet':s.title,'rows':s.max_row,'columns':s.max_column,'header':' / '.join(labels[:17]),'has_student_count_label':any('学生数' in v or '在校生' in v for v in labels),'has_per_student_label':any('生均' in v for v in labels)})
assert len(inventory)==4 and not any(r['has_student_count_label'] for r in inventory)
save_csv(ROOT/'data/finance_denominator_zhejiang_field_inventory.csv',inventory)

# Public web extraction retained separately: raw HTML download returned 403.
evidence=ROOT/'data/finance_denominator_shucheng_web_extract.txt'
e=evidence.read_text(encoding='utf-8')
url='https://www.shucheng.gov.cn/public/6626413/37309387.html'
blocks=[('combined','414',430,454,'14913623.00','15413623.00','15124019.00','15114019.00','10000.00'),
        ('primary','61',294,313,'10582262.66','10582262.66','10389734.66','10383086.69','6647.97'),
        ('junior','413',136,141,'4331360.34','4831360.34','4734284.34','4730932.31','3352.03')]
rows=[]
for stage,category,begin,end,income,spend,subsidy,budget,fund in blocks:
    for value in [begin,end,income,spend,subsidy,budget,fund]:
        assert format(D(str(value)),',') in e or str(value) in e,str(value)
    rows.append({'year':2023,'unit_code':'3134003910S','stage':stage,'school_category_code':category,'students_begin':begin,'students_end':end,'total_income_yuan':income,'total_expenditure_yuan':spend,'fiscal_subsidy_expenditure_yuan':subsidy,'general_budget_arranged_education_expenditure_yuan':budget,'government_fund_expenditure_yuan':fund,'source':url,'evidence_type':'web extracted public table; original HTML retrieval 403','evidence_path':str(evidence.relative_to(ROOT)),'evidence_sha256':hashlib.sha256(evidence.read_bytes()).hexdigest(),'eligibility':'single-year source bridge; weights and stage-allocation instructions unresolved; not analytical panel'})
save_csv(ROOT/'data/finance_denominator_shucheng_candidates.csv',rows)
checks=[]
for field in ['students_begin','students_end','total_income_yuan','total_expenditure_yuan','fiscal_subsidy_expenditure_yuan','general_budget_arranged_education_expenditure_yuan','government_fund_expenditure_yuan']:
    combined=D(str(rows[0][field])); calculated=D(str(rows[1][field]))+D(str(rows[2][field]))
    assert combined==calculated
    checks.append({'check':'primary_plus_junior_vs_combined','field':field,'computed':str(calculated),'reported':str(combined),'difference':str(calculated-combined),'status':'pass'})
for row in rows:
    calculated=D(row['general_budget_arranged_education_expenditure_yuan'])+D(row['government_fund_expenditure_yuan'])
    assert calculated==D(row['fiscal_subsidy_expenditure_yuan'])
    checks.append({'check':'budget_plus_fund_vs_fiscal_subsidy','stage':row['stage'],'computed':str(calculated),'reported':row['fiscal_subsidy_expenditure_yuan'],'difference':'0','status':'pass'})
save_csv(ROOT/'data/finance_denominator_shucheng_checks.csv',checks)
summary={'download_records':len(sources),'saved_originals':sum(r['status']=='downloaded' for r in sources),'web_extracted_candidate_rows':len(rows),'source_sum_checks':len(checks),'matched_continuous_weighted_fiscal_pairs':0,'original_html_saved_for_shucheng':False,'year_start_end_weights_verified':False,'main_panel_modified':False,'visualization_modified':False,'yantai_pdf_years':'reporting dates extracted; data years not verified from publication pages; no numeric panel entry'}
(ROOT/'data/finance_denominator_audit.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(summary,ensure_ascii=False))
