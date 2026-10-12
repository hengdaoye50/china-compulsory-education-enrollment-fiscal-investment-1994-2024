"""Province finance source screening and narrowly scoped 2017 label repair.

Does not build school/province joins or reconstruct fiscal denominators.
Requires project-local xlrd 2.0.2 for published XLS attachment inspection.
"""
import csv, hashlib, io, json, re, sys
from decimal import Decimal as D
from lxml import html
from openpyxl import load_workbook
from parse import ROOT, save_csv
sys.path.insert(0,str(ROOT/'.tmp/finance_xls_reader'))
import xlrd

def read_csv(path):
    return list(csv.DictReader(path.open(encoding='utf-8-sig',newline='')))
def text(path):
    doc=html.fromstring(path.read_bytes().decode('utf-8'))
    return ' '.join(' '.join(doc.xpath('//text()[not(ancestor::script) and not(ancestor::style)]')).split())
def region(value):
    value=re.sub(r'\s','',value)
    if value=='全国':return value
    for p in provinces:
        if value.startswith(p):return p
    return None
finance_path=ROOT/'data/finance_long.csv'
finance=read_csv(finance_path)
provinces=sorted({r['province'] for r in finance if r['province']!='全国'})
assert len(provinces)==31
before_hash=hashlib.sha256(finance_path.read_bytes()).hexdigest()
doc=html.fromstring((ROOT/'converted/finance_2017_36.html').read_bytes().decode('utf-8'))
tables=doc.xpath('//table')
correct2017={}
for index, metric, expected_caption in [(1,'career_per_student','事业费'),(3,'public_per_student','公用经费')]:
    table=tables[index]; prev=table.getprevious(); caption=''
    while prev is not None:
        val=' '.join(prev.text_content().split())
        if '各级教育生均' in val:caption=val;break
        prev=prev.getprevious()
    assert expected_caption in caption,caption
    for row in table.xpath('./tr|./tbody/tr')[2:]:
        cells=[''.join(td.itertext()).strip() for td in row.xpath('./td|./th')]
        p=region(cells[0])
        if not p:continue
        for stage,col in [('primary',2),('junior',5)]:
            correct2017[(p,stage,metric)]=D(cells[col].replace(',','').replace(' ',''))
assert len(correct2017)==128

# 2018 reports carry 2017 prior-year values. Keep this publication vintage separate.
prior_doc=html.fromstring((ROOT/'converted/finance_2018_36.html').read_bytes().decode('utf-8'))
prior_tables=prior_doc.xpath('//table');backrefs=[];backchecks=[]
source_2018=next(r['source'] for r in finance if r['year']=='2018' and r['metric']=='budget_per_student')
for index,metric in [(1,'budget_per_student'),(3,'career_per_student'),(5,'public_per_student')]:
    for row in prior_tables[index].xpath('./tr|./tbody/tr')[2:]:
        cells=[''.join(c.itertext()).strip() for c in row.xpath('./td|./th')]
        p=region(cells[0])
        if not p:continue
        for stage,col in [('primary',4),('junior',7)]:
            value=D(cells[col].replace(',','').replace(' ',''))
            if metric=='budget_per_student':
                backrefs.append(dict(observation_year=2017,publication_data_year=2018,province=p,stage=stage,
                                     metric=metric,value=str(value),source=source_2018,table_index=index,column_index=col,
                                     vintage='2018 attachment prior-year column; separate from 2017 original release',
                                     eligibility='published per-student value only; not paired expenditure/denominator'))
            else:
                assert value==correct2017[(p,stage,metric)]
                backchecks.append(dict(province=p,stage=stage,metric=metric,value=str(value),status='pass',
                                       check='2017 original corrected labels vs 2018 prior-year column'))
assert len(backrefs)==64 and len(backchecks)==128
save_csv(ROOT/'data/province_finance_2017_budget_backward_reference.csv',backrefs)
save_csv(ROOT/'data/province_finance_2017_backward_checks.csv',backchecks)

# Repair only demonstrated table-label errors. Preserve historical affected rows once.
affected=[r for r in finance if r['year']=='2017' and r['stage']!='all']
needs_repair=any(r['metric']=='budget_per_student' for r in affected)
if needs_repair:
    backup=ROOT/'data/province_finance_2017_before_label_repair.csv'
    assert not backup.exists(), 'Do not overwrite an earlier repair record'
    save_csv(backup,affected)
    for r in affected:
        r['metric']={'budget_per_student':'career_per_student','career_per_student':'public_per_student'}[r['metric']]
    save_csv(finance_path,finance)
    panel_path=ROOT/'data/education_resource_panel.csv'
    panel=read_csv(panel_path)
    snapshot=[{k:v for k,v in r.items() if '_per_student' not in k} for r in panel]
    for r in panel:
        if r['year']!='2017':continue
        for metric in ['career_per_student','public_per_student']:
            r[metric]=str(correct2017[(r['province'],r['stage'],metric)])
            r[metric+'_source']=next(f['source'] for f in affected if f['province']==r['province'] and f['stage']==r['stage'] and f['metric']==metric)
        r['budget_per_student']='';r['budget_per_student_source']=''
    assert snapshot==[{k:v for k,v in r.items() if '_per_student' not in k} for r in panel]
    save_csv(panel_path,panel)
    jp=ROOT/'data/panel.json';records=json.loads(jp.read_text(encoding='utf-8'))
    for r in records:
        if r['year']!=2017:continue
        for metric in ['career_per_student','public_per_student']:
            r[metric]=float(correct2017[(r['province'],r['stage'],metric)])
            r[metric+'_source']=next(f['source'] for f in affected if f['province']==r['province'] and f['stage']==r['stage'] and f['metric']==metric)
        r.pop('budget_per_student',None);r.pop('budget_per_student_source',None)
    jp.write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
for r in affected:
    assert D(r['value'])==correct2017[(r['province'],r['stage'],r['metric'])]

sources=[]
for name in ['province_finance_sources','province_finance_attachment_sources','province_finance_followup_sources','province_finance_deep_sources']:
    sources+=json.loads((ROOT/(name+'.json')).read_text(encoding='utf-8'))
bykey={r['key']:r for r in sources}
for r in sources:
    if r['status']=='downloaded':
        p=ROOT/r['path'];assert len(p.read_bytes())==r['bytes']
        assert hashlib.sha256(p.read_bytes()).hexdigest()==r['sha256']
save_csv(ROOT/'data/province_finance_source_receipt.csv', [{k:r.get(k) for k in ['key','url','status','path','bytes','sha256','error']} for r in sources])
def source_text(key):return text(ROOT/bykey[key]['path'])

scope=[]
for y in [2021,2022,2023]:
    k=f'province_finance_liaoning_{y}';t=source_text(k)
    for phrase in ['地方属教育部门和其他部门办学校','不含企业办、民办、中外合作办以及中央属学校','且是各单位当年实际支出的数据','其中数']:
        assert phrase in ''.join(t.split()), (k,phrase)
    scope.append(dict(province='辽宁',year=y,source=bykey[k]['url'],locator='注2-5',
                      finding='生均限地方属教育及其他部门办；事业费为205剔除附加及同级发改委基建后的实际支出；公用为其中数',
                      boundary='只核准这三份公告；未取得普通小学/初中学生权重；不推广所有省年'))
save_csv(ROOT/'data/province_finance_scope_evidence.csv',scope)

# Table inventories inspect complete spreadsheet cells, retaining blanks as blanks.
inventory=[];comparisons=[]
national={(int(r['year']),r['province'],r['stage'],r['metric']):r for r in finance}
def compare(province,year,stage,metric,value,source,locator,tolerance='0.005'):
    reported=D(national[(year,province,stage,metric)]['value'])
    difference=D(str(value))-reported
    assert abs(difference)<=D(tolerance),(province,metric,difference)
    comparisons.append(dict(province=province,year=year,stage=stage,metric=metric,local_value=str(value),
                            moe_value=str(reported),difference=str(difference),tolerance=tolerance,source=source,locator=locator,status='pass'))
for key,province,fmt in [('province_finance_heilongjiang_2023_xlsx','黑龙江','xlsx'),
                         ('province_finance_jiangsu_2023_republication_attachment','江苏','xls'),
                         ('province_finance_zhejiang_2023_attachment','浙江','xlsx')]:
    raw=(ROOT/bykey[key]['path']).read_bytes()
    if fmt=='xlsx':
        book=load_workbook(io.BytesIO(raw),data_only=True)
        sheets=[(s.title,list(s.iter_rows(values_only=True))) for s in book]
    else:
        book=xlrd.open_workbook(file_contents=raw)
        sheets=[(s.name,[s.row_values(i) for i in range(s.nrows)]) for s in book.sheets()]
    for name,rows in sheets:
        labels=[str(v) for row in rows for v in row if isinstance(v,str)]
        inventory.append(dict(province=province,year=2023,source=bykey[key]['url'],sheet=name,rows=len(rows),
                              header=' / '.join(labels[:12]),student_count_field_found=any('学生数' in v or '在校生' in v for v in labels),
                              note='人均经费不是财政学生数字段；不据生均反算分母'))
        metric=next((m for token,m in [('公用经费','public_per_student'),('事业费','career_per_student'),('生均一般公共预算教育经费','budget_per_student')]
                     if any(token in v for v in labels[:5])),None)
        if province=='黑龙江' and metric:
            row=next(row for row in rows if row[0]=='黑龙江省')
            for stage,col in [('primary',9),('junior',7)]:compare(province,2023,stage,metric,row[col],bykey[key]['url'],name+' 省合计')
        if province=='江苏' and metric:
            row=next(row for row in rows if row[0]=='江苏省')
            for stage,col in [('primary',14),('junior',11)]:compare(province,2023,stage,metric,row[col],bykey[key]['url'],name+' 省合计')
assert not any(r['student_count_field_found'] for r in inventory)
save_csv(ROOT/'data/province_finance_spreadsheet_inventory.csv',inventory)

# Direct provincial HTML versus MOE attachment, matching named metrics without mixing units.
for key,province,year in [('province_finance_fujian_2023','福建',2023),('province_finance_fujian_2022','福建',2022),
                         ('province_finance_fujian_2021','福建',2021),('province_finance_fujian_2020','福建',2020),
                         ('province_finance_fujian_2019','福建',2019),('province_finance_liaoning_2023','辽宁',2023),
                         ('province_finance_liaoning_2022','辽宁',2022),('province_finance_liaoning_2021','辽宁',2021),
                         ('province_finance_zhejiang_2023','浙江',2023)]:
    t=source_text(key)
    for stage,label in [('primary','普通小学'),('junior','普通初中')]:
        matches=re.findall(label+r'(?:(?:生均)?一般公共预算(?:教育经费|教育事业费支出|公用经费支出))?为?([\d,]+\.\d+)元',''.join(t.split()))
        assert len(matches)==3,(key,label,matches)
        for metric,value in zip(['budget_per_student','career_per_student','public_per_student'],matches):
            compare(province,year,stage,metric,value.replace(',',''),bykey[key]['url'],'公告分学段段落')
save_csv(ROOT/'data/province_finance_cross_source_checks.csv',comparisons)

# A separate parser rechecks all selected provincial cells and their table captions.
cell_checks=[]
for year in range(2017,2024):
    path=ROOT/'converted'/f'finance_{year}_36.html'
    parsed=html.fromstring(path.read_bytes().decode('utf-8')).xpath('//table')
    for r in finance:
        if int(r['year'])!=year or r['province']=='全国' or r['stage']=='all':continue
        table=parsed[int(r['table_index'])];prev=table.getprevious();caption=''
        while prev is not None:
            candidate=''.join(prev.text_content().split())
            if '生均' in candidate:caption=candidate;break
            prev=prev.getprevious()
        expected={'budget_per_student':'一般公共预算教育经费','career_per_student':'事业费','public_per_student':'公用经费'}[r['metric']]
        assert expected in caption,(year,r['metric'],caption)
        if r['metric']=='budget_per_student':assert '事业费' not in caption and '公用经费' not in caption
        cells=next([''.join(c.itertext()).strip() for c in row.xpath('./td|./th')]
                   for row in table.xpath('./tr|./tbody/tr')[2:] if row.xpath('./td|./th') and
                   region(''.join(row.xpath('./td|./th')[0].itertext()))==r['province'])
        value=D(cells[int(r['column_index'])].replace(',','').replace(' ',''))
        assert value==D(r['value'])
        cell_checks.append(dict(year=year,province=r['province'],stage=r['stage'],metric=r['metric'],
                                caption=caption,value=str(value),table_index=r['table_index'],column_index=r['column_index'],status='pass'))
save_csv(ROOT/'data/province_finance_source_cell_checks.csv',cell_checks)

registry=[]
deep={(r['province'],r['year']) for r in comparisons}
for province in provinces:
    for year in range(2017,2024):
        selected=[r for r in finance if r['province']==province and int(r['year'])==year]
        registry.append(dict(province=province,year=year,published_stage_metric_cells=sum(r['stage']!='all' for r in selected),
                             all_stage_budget_total_found=any(r['metric']=='regional_budget_education_100m' for r in selected),
                             stage_actual_spending_verified=False,paired_fiscal_students_verified=False,student_weights_verified=False,
                             local_source_depth='交叉核验公告或电子表' if (province,year) in deep else '仅既有全国公告附件；省级深查待完成',
                             classification='只可描述已发布生均指标；不做支出/分母分解',
                             note='2017附件仅事业费与公用；预算指标未从该附件取得' if year==2017 else ('2021预算增幅含同口径调整；不直接用2020原值复算' if year==2021 else ''),
                             source=' | '.join(sorted({r['source'] for r in selected}))))
save_csv(ROOT/'data/province_finance_2017_2023_registry.csv',registry)
summary=dict(source_records=len(sources),saved_originals=sum(r['status']=='downloaded' for r in sources),
             failed_downloads=sum(r['status']!='downloaded' for r in sources),province_year_registry=len(registry),
             existing_stage_metric_cells=sum(r['published_stage_metric_cells'] for r in registry),
             spreadsheet_sheets_inspected=len(inventory),cross_source_checks=len(comparisons),
             source_cell_and_caption_checks=len(cell_checks),
             separate_2017_budget_backward_reference_cells=len(backrefs),backward_checks=len(backchecks),
             repaired_2017_metric_labels=128,matched_continuous_stage_spending_student_pairs=0,
             all_provincial_local_year_sources_fully_screened=False,resource_numeric_fields_unchanged=True,
             visualizations_unchanged=True,existing_finance_visuals_use_2022_2023_unaffected=True,
             repair_backup_path='data/province_finance_2017_before_label_repair.csv',
             finance_hash_after=hashlib.sha256(finance_path.read_bytes()).hexdigest())
(ROOT/'data/province_finance_audit.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(summary,ensure_ascii=False))
