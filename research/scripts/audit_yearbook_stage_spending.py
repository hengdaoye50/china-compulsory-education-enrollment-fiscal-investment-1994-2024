"""Audit user-provided yearbooks separately from the main panel.
No denominator reconstruction, no implicit blank-to-zero conversion.
"""
import csv, hashlib, json, re, sys
from decimal import Decimal as D, InvalidOperation
from pathlib import Path
import pdfplumber
from openpyxl import load_workbook
from parse import ROOT, save_csv
sys.stdout.reconfigure(encoding='utf-8')
DATA=ROOT/'data';PROJECT=ROOT.parent;BASE=ROOT/'用户提供资料/年鉴解压'
def read(p):return list(csv.DictReader(p.open(encoding='utf-8-sig',newline='')))
def norm(v):return re.sub(r'\s+','',str(v or ''))
def number(v):
 if v is None or str(v).strip()=='':return None
 try:return D(str(v).replace(',','').strip())
 except InvalidOperation:return None
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
protected=json.loads((DATA/'province_finance_next_protected_baseline.json').read_text(encoding='utf-8'))
assert all(sha(ROOT/p)==s for p,s in protected.items())
old_inventory={r['path']:r for r in read(DATA/'yearbook_extracted_file_inventory.csv')}
national={(r['province'],int(r['year']),r['stage'],r['metric']):r for r in read(DATA/'finance_long.csv')}
provinces=sorted({k[0] for k in national if k[1]==2023 and k[0]!='全国'})
assert len(provinces)==31
cities=['大连','宁波','厦门','青岛','深圳']
tables=[];rawcells=[];amounts=[];perstudent=[];checks=[];rollups=[];missing=[]
year_evidence=[]
for volume in range(2018,2025):
 year=volume-1
 notes=list((BASE/f'卷{volume}').rglob('*.pdf'))
 method={2018:'human visual: note physical3 / printed634, statistical time 2017',2019:'human visual: note physical4 / printed634, statistical time 2018',2020:'official NBS 2020 book description states data year 2019; package note absent',2021:'note text physical3 / printed639, statistical time 2020',2022:'note text physical3 / printed639, statistical time 2021',2023:'human visual: note physical3 / printed639, statistical time 2022',2024:'NII/CiNii 2024 volume subject year 2023; package note absent'}[volume]
 year_evidence.append(dict(volume_year=volume,data_year=year,basis=method,local_note=str(notes[0].relative_to(PROJECT)) if notes else '',local_note_sha256=sha(notes[0]) if notes else '',boundary='calendar data-year mapping; package authenticity not independently certified; cross-source cells checked separately'))
 selected=[]
 for p in (BASE/f'卷{volume}').rglob('*.xlsx'):
  if not re.match(r'[4567]-',p.name):continue
  part=int(p.name[0])
  if part!=7 and not any(s in p.name for s in ['普通小学','普通初中']):continue
  w=load_workbook(p,read_only=True,data_only=True)
  for ws in w:
   head=''.join(norm(v) for row in ws.iter_rows(max_row=3,values_only=True) for v in row if v is not None)
   stage='primary' if '普通小学' in head else 'junior' if '普通初中' in head else None
   if not stage or '农村' in head:continue
   if part==7 and not ('生均一般公共预算教育事业费和基本建设支出' in head):continue
   scope='local_affiliation' if '地方普通' in head else 'all_affiliations'
   if part==7:assert scope=='local_affiliation'
   rel=str(p.relative_to(PROJECT));assert sha(p)==old_inventory[rel]['sha256']
   selected.append((part,stage,scope))
   rows=[];occupied=[]
   for row in ws:
    for c in row:
     if c.value is not None:
      occupied.append((c.coordinate,str(c.value)))
      rawcells.append(dict(volume_year=volume,data_year=year,file=rel,sheet=ws.title,cell=c.coordinate,value=str(c.value)))
    region=norm(row[0].value)
    if region=='合计':region='全国'
    if region not in provinces+cities+['全国']:continue
    rows.append((region,row))
   observed=[a for a,_ in rows]
   assert len(observed)==len(set(observed)),(p,'duplicate regions')
   assert set(observed)==set(provinces+cities+['全国']),(p,set(observed))
   header=' | '.join(f'{a}={v}' for a,v in occupied if int(re.sub(r'\D','',a))<min(row[0].row for _,row in rows))
   assert ('千元' in header) if part!=7 else ('元' in header)
   hits=';'.join(sorted(set(re.findall(r'.{0,12}(?:学生数|学生人数|年初学生|年末学生|加权|权重).{0,20}',' '.join(v for _,v in occupied)))))
   tables.append(dict(volume_year=volume,data_year=year,part=part,stage=stage,affiliation_scope=scope,file=rel,sheet=ws.title,sha256=sha(p),province_rows=31,city_rows=5,national_rows=1,headers=header,student_field_hits=hits,source_type='user supplied Excel compilation',vintage='annual yearbook volume; original publisher cell version not separately verified'))
   components={}
   for region,row in rows:
    common=dict(volume_year=volume,year=year,province=region,stage=stage,affiliation_scope=scope,part=part,file=rel,sheet=ws.title,row=row[0].row,region_cell=row[0].coordinate,source_type='user supplied Excel compilation',eligibility='separate candidate; fiscal student denominator unavailable')
    if part==7:
     # Header may occupy six rows in 2019, but data column meanings stay explicit.
     assert '事业费支出' in header and '公用部分' in header
     for metric,col in [('career_per_student',3),('public_per_student',5)]:
      c=row[col-1];v=number(c.value);assert v is not None,(p,c.coordinate)
      perstudent.append(dict(**common,metric=metric,value=str(v),unit='元',cell=c.coordinate))
      if region not in cities:
       n=national[(region,year,stage,metric)];diff=v-D(n['value'])
       status=('different_national_affiliation_scope_not_comparable' if region=='全国' else 'match' if abs(diff)<=D('.005') else 'difference_to_investigate')
       checks.append(dict(volume_year=volume,year=year,province=region,stage=stage,metric=metric,yearbook_value=str(v),announcement_value=n['value'],difference=str(diff),status=status,comparison_eligible=region!='全国',file=rel,cell=c.coordinate,announcement_source=n['source']))
    else:
     assert all(t in header for t in ['个人部分','公用部分','基本建设'])
     cv={name:number(row[col-1].value) for name,col in [('total',2),('personal',3),('wage_benefit',4),('personal_family_subsidy',5),('public',7),('goods_services',8),('capital_other',9),('capital_equipment',10),('capital_project',11),('construction',12)]}
     assert all(cv[x] is not None for x in ['total','personal','public']),p
     components[region]=cv
     for metric,col in [('total',2),('personal',3),('public',7),('construction',12)]:
      c=row[col-1];v=cv[metric]
      amounts.append(dict(**common,metric=metric,value='' if v is None else str(v),unit='千元',cell=c.coordinate,value_status='blank as original; not assigned zero' if v is None else 'numeric original cell',amount_definition={4:'all-source education spending',5:'fiscal subsidy spending, including general public budget career/construction and government fund budget etc',6:'general public budget career and construction spending'}[part]))
     # Arithmetic diagnostics only, never manufacture a missing original cell.
     for kind,left,right in [('total_parts','total',['personal','public','construction']),('personal_parts','personal',['wage_benefit','personal_family_subsidy']),('public_parts','public',['goods_services','capital_other']),('capital_parts','capital_other',['capital_equipment','capital_project'])]:
      vals=[cv[left]]+[cv[x] for x in right]
      if any(v is None for v in vals):
       missing.append(dict(volume_year=volume,year=year,province=region,stage=stage,part=part,scope=scope,kind=kind,file=rel,status='arithmetic skipped for blank source component'))
       continue
      diff=vals[0]-sum(vals[1:],D(0))
      rollups.append(dict(volume_year=volume,year=year,province=region,stage=stage,part=part,scope=scope,kind=kind,difference=str(diff),tolerance='2千元 per-row component rounding',status='pass' if abs(diff)<=D(2) else 'unresolved',file=rel))
   if part!=7:
    for metric in ['total','personal','public']:
     diff=sum(components[x][metric] for x in provinces)-components['全国'][metric]
     rollups.append(dict(volume_year=volume,year=year,province='全国',stage=stage,part=part,scope=scope,kind='31_province_sum_'+metric,difference=str(diff),tolerance='16千元 sum of 32 rounded cells',status='pass' if abs(diff)<=D(16) else 'unresolved',file=rel))
  w.close()
 expected={(part,stage,scope) for part in [4,5,6] for stage in ['primary','junior'] for scope in ['all_affiliations','local_affiliation']}|{(7,s,'local_affiliation') for s in ['primary','junior']}
 assert set(selected)==expected and len(selected)==14,(volume,selected)
 print(f'volume {volume}: 14 stage tables checked',flush=True)
for name,rows in [('yearbook_stage_table_inventory.csv',tables),('yearbook_stage_source_cells.csv',rawcells),('yearbook_stage_spending_candidates.csv',amounts),('yearbook_stage_per_student_candidates.csv',perstudent),('yearbook_announcement_cell_checks.csv',checks),('yearbook_stage_arithmetic_checks.csv',rollups),('yearbook_stage_blank_components.csv',missing),('yearbook_data_year_evidence.csv',year_evidence)]:save_csv(DATA/name,rows)
assert all(sha(ROOT/p)==s for p,s in protected.items())
audit=dict(volumes=7,data_years=list(range(2017,2024)),selected_tables=len(tables),amount_fields_including_blanks=len(amounts),numeric_original_amount_cells=sum(r['value_status']=='numeric original cell' for r in amounts),blank_original_amount_fields=sum(r['value']=='' for r in amounts),province_amount_fields_including_blanks=sum(r['province'] in provinces for r in amounts),per_student_cells=len(perstudent),announcement_comparison_records=len(checks),province_announcement_checks=sum(c['province']!='全国' for c in checks),province_announcement_matches=sum(c['status']=='match' for c in checks),province_announcement_unresolved=sum(c['status']=='difference_to_investigate' for c in checks),national_scope_noncomparisons=sum(c['province']=='全国' for c in checks),arithmetic_checks=len(rollups),arithmetic_pass=sum(c['status']=='pass' for c in rollups),arithmetic_unresolved=sum(c['status']!='pass' for c in rollups),national_province_sum_checks=sum(c['kind'].startswith('31_province_sum') for c in rollups),skipped_blank_component_checks=len(missing),matched_spending_student_pairs=0,denominators_reconstructed=False,protected_files_unchanged=True,publisher_original_cells_independently_verified=False,student_weights_found=False)
(DATA/'yearbook_stage_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(audit,ensure_ascii=False,indent=2))
