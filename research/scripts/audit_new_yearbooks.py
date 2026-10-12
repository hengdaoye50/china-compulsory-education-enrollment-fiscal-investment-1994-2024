from pathlib import Path
import zipfile, io, json, re, csv, hashlib, sys
from collections import Counter
import openpyxl
ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'research' / 'data' / '年鉴资料核查_20261010'
OUT.mkdir(exist_ok=True)
sys.path.insert(0,str(ROOT/'research'/'.tmp'/'yearbook_read_deps'))
sys.stdout.reconfigure(encoding='utf-8')
families = ['中国统计年鉴','中国财政年鉴','中国教育统计年鉴']
patterns = {
 '学生':r'(各地区|分地区).*(小学|初中).*(基本情况|学生)|(?:小学|初中).*学生数',
 '分省小学':r'(各地区|分地区).*(小学).*|(?:小学).*(各地区|分地区)',
 '分省初中':r'(各地区|分地区).*(初中).*|(?:初中).*(各地区|分地区)',
 '财政本级收支':r'(各地区|分地区).*(财政收入|财政支出|预算收入|预算支出)|(?:省|市|区).*收支.*总表',
 '财政补助平衡':r'税收返还|转移支付|(?:中央|财政).*补助|财政收支平衡|收支决算总表',
 '物价':r'(各地区|分地区).*(居民消费价格|消费价格指数)',
 '人口经济':r'(各地区|分地区).*(人口|生产总值)',
 '口径说明':r'编者|指标解释|统计说明|统计范围|关于年份',
}
def decode_name(info):
 s=info.filename
 if not (info.flag_bits & 0x800):
  try:s=s.encode('cp437').decode('gbk')
  except (UnicodeError,LookupError):pass
 return s
def inspect(b,name):
 suffix=Path(name).suffix.lower()
 if suffix not in ['.xls','.xlsx']:return {'status':'not_excel'}
 try:
  if b[:2]==b'PK':
   w=openpyxl.load_workbook(io.BytesIO(b),read_only=True,data_only=True)
   sheets=[]
   for s in w.worksheets[:2]:
    rows=[]; region=[]; terms=[]
    for i,row in enumerate(s.iter_rows(values_only=True),1):
     vals=[str(v) if v is not None else '' for v in row]
     if i<=15 and any(vals):rows.append({'row':i,'cells':vals[:22]})
     line=' '.join(vals)
     if re.search('北京|天津|河北|重庆|四川|新疆|全省|本级|全区|全市|总计|合计',line):region.append({'row':i,'cells':vals[:22]})
     if re.search('补助|返还|转移|上解|结转|调入|债务|收入总计|支出总计|决算|单位',line):terms.append({'row':i,'cells':vals[:22]})
    sheets.append({'sheet':s.title,'rows':s.max_row,'cols':s.max_column,'headers':rows,'regions':region[:12],'fiscal_terms':terms[:28]})
   w.close();return {'status':'read','format':'xlsx','sheets':sheets}
  import xlrd
  w=xlrd.open_workbook(file_contents=b)
  sheets=[]
  for s in w.sheets()[:2]:
   rows=[];region=[];terms=[]
   for i in range(s.nrows):
    vals=[str(v) for v in s.row_values(i)]; line=' '.join(vals)
    if i<15 and any(vals):rows.append({'row':i+1,'cells':vals[:22]})
    if re.search('北京|天津|河北|重庆|四川|新疆|全省|本级|全区|全市|总计|合计',line):region.append({'row':i+1,'cells':vals[:22]})
    if re.search('补助|返还|转移|上解|结转|调入|债务|收入总计|支出总计|决算|单位',line):terms.append({'row':i+1,'cells':vals[:22]})
   sheets.append({'sheet':s.name,'rows':s.nrows,'cols':s.ncols,'headers':rows,'regions':region[:12],'fiscal_terms':terms[:28]})
  return {'status':'read','format':'xls','sheets':sheets}
 except Exception as e:return {'status':'error','error':str(e)[:300],'magic':b[:12].hex()}
inventory=[]; candidates=[]; samples=[]
for family in families:
 for p in sorted((ROOT/family).glob('*.zip')):
  year=int(re.search(r'(19|20)\d{2}',p.name)[0]);rec={'family':family,'volume_year':year,'archive':str(p.relative_to(ROOT)),'bytes':p.stat().st_size}
  try:
   with zipfile.ZipFile(p) as z:
    infos=[i for i in z.infolist() if not i.is_dir()];rec['entries']=len(infos);rec['extensions']=dict(Counter(Path(decode_name(i)).suffix.lower() for i in infos));count=Counter(); chosen={}
    for info in infos:
     name=decode_name(info)
     cats=[k for k,pat in patterns.items() if re.search(pat,name)]
     for k in cats:
      count[k]+=1;candidates.append({'family':family,'volume_year':year,'archive':str(p.relative_to(ROOT)),'category':k,'table':name})
      if k not in chosen and Path(name).suffix.lower() in ['.xls','.xlsx']:chosen[k]=info
     # Same province across all fiscal volumes for scope checks.
     if family=='中国财政年鉴' and re.search('河北.*收支.*总表',name) and not re.search('续表|基金|资本|保险',name):chosen['河北收支总表']=info
    rec['candidate_counts']=dict(count)
    for cat,info in chosen.items():
     if cat not in ['学生','分省小学','分省初中','财政本级收支','财政补助平衡','物价','河北收支总表']:continue
     sample={'family':family,'volume_year':year,'archive':str(p.relative_to(ROOT)),'category':cat,'table':decode_name(info)}
     sample.update(inspect(z.read(info),decode_name(info)));samples.append(sample)
  except Exception as e:rec['error']=str(e)
  inventory.append(rec)
  print(f'{family} {year}: {rec.get("entries",0)} files {rec.get("candidate_counts",{})}',flush=True)
(OUT/'archive_inventory.json').write_text(json.dumps(inventory,ensure_ascii=False,indent=2),encoding='utf8')
(OUT/'table_samples.json').write_text(json.dumps(samples,ensure_ascii=False,indent=2),encoding='utf8')
with (OUT/'key_table_catalog.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.DictWriter(f,fieldnames=['family','volume_year','archive','category','table']);w.writeheader();w.writerows(candidates)
extra={}
for family in families:
 extra[family]=[str(p.relative_to(ROOT)) for p in (ROOT/family).glob('*') if p.is_file() and p.suffix.lower()!='.zip']
(OUT/'extra_files.json').write_text(json.dumps(extra,ensure_ascii=False,indent=2),encoding='utf8')
print('SAMPLE_STATUS',dict(Counter(s['status'] for s in samples)),flush=True)
