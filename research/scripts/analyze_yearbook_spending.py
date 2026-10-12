"""Conditional fiscal-only description. No student denominators or main-panel writes."""
import csv,json,hashlib,re,sys
from pathlib import Path
from decimal import Decimal as D
from collections import defaultdict
from fractions import Fraction
from openpyxl import load_workbook
from parse import ROOT,save_csv
sys.stdout.reconfigure(encoding='utf-8')
DATA=ROOT/'data';PROJECT=ROOT.parent;BASE=ROOT/'用户提供资料/年鉴解压'
def read(n):return list(csv.DictReader((DATA/n).open(encoding='utf-8-sig',newline='')))
def num(v):return None if v is None or str(v).strip()=='' else D(str(v).replace(',',''))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def fmt(v):return '' if v is None else str(v)
def pct(a,b):return (a/b-1)*100
def clean(v):return re.sub(r'\s+','',str(v or ''))
protected=json.loads((DATA/'province_finance_next_protected_baseline.json').read_text(encoding='utf-8'))
def protect():assert all(sha(ROOT/p)==s for p,s in protected.items())
protect()
inventory={r['path']:r for r in read('yearbook_extracted_file_inventory.csv')}
for r in read('yearbook_stage_table_inventory.csv'):assert sha(PROJECT/r['file'])==r['sha256']
panel=read('yearbook_province_stage_spending_candidate_panel.csv')
keys=[(r['province'],r['year'],r['stage'],r['affiliation_scope']) for r in panel]
assert len(keys)==len(set(keys))==868
assert {r['unit'] for r in panel}=={'千元'}
bykey=dict(zip(keys,panel));funds=['all_source','fiscal_subsidy','general_budget'];metrics=['total','personal','public','construction']
qc=[]
def check(kind,context,a,b):
 if a is None or b is None:status='not_checked_original_blank';difference=''
 else:difference=str(a-b);status='pass' if a-b>=-D(2) else 'unresolved'
 qc.append(dict(kind=kind,**context,difference=difference,status=status,tolerance='2千元 for separate rounded cells'))
for r in panel:
 ctx={k:r[k] for k in ['province','year','stage','affiliation_scope']}
 for upper,lower in zip(funds,funds[1:]):
  for m in metrics:check(upper+'_contains_'+lower+'_'+m,ctx,num(r[upper+'_'+m]),num(r[lower+'_'+m]))
 if r['affiliation_scope']=='all_affiliations':
  local=bykey[r['province'],r['year'],r['stage'],'local_affiliation']
  for f in funds:
   for m in metrics:check('all_contains_local_'+f+'_'+m,ctx,num(r[f+'_'+m]),num(local[f+'_'+m]))

# Annual national organizer tables distinguish affiliation from organizer explicitly.
org=[];orgraw=[];ambiguous=[]
for volume in range(2018,2025):
 found=[]
 for p in (BASE/f'卷{volume}').rglob('*.xlsx'):
  if not p.name.startswith('1-') or not '支出明细' in p.name:continue
  w=load_workbook(p,read_only=True,data_only=True);ws=w.active
  original_title=''.join(clean(v) for row in ws.iter_rows(max_row=2,values_only=True) for v in row)
  title=original_title.replace('其它','其他')
  category=next((c for c in ['地方教育和其他部门','中央教育和其他部门','教育和其他部门','全国','地方'] if '('+c+')' in title or '（'+c+'）' in title),None)
  if category=='中央教育和其他部门':w.close();continue
  fund='general_budget' if '一般公共预算教育事业费和基本建设' in title else 'fiscal_subsidy' if '财政补助支出明细' in title else 'all_source' if '教育经费支出明细' in title else None
  if not category or not fund:w.close();continue
  ambiguous_block=volume==2024 and 23<=int(re.match(r'1-(\d+)',p.name)[1])<=32
  if ambiguous_block:
   assert sha(p)==inventory[str(p.relative_to(PROJECT))]['sha256']
   ambiguous.append(dict(volume_year=volume,file=str(p.relative_to(PROJECT)),original_title=original_title,issue='second block repeats education-spending title; fiscal-subsidy label absent',treatment='not assigned a funding category; excluded from organizer coverage inference'))
   w.close();continue
  rel=str(p.relative_to(PROJECT));assert sha(p)==inventory[rel]['sha256'];found.append((fund,category))
  stagehits=[]
  for row in ws:
   label=clean(row[0].value)
   stage='primary' if re.fullmatch(r'(?:1[.．、])?普通小学',label) else 'junior' if label=='普通初中' else None
   if not stage:continue
   stagehits.append(stage)
   for metric,col in [('total',2),('personal',3),('public',7),('construction',12)]:
    c=row[col-1];v=num(c.value)
    org.append(dict(year=volume-1,volume_year=volume,stage=stage,fund=fund,organizer_affiliation_category=category,metric=metric,value=fmt(v),unit='千元',file=rel,sheet=ws.title,cell=c.coordinate,sha256=sha(p),original_title=original_title))
   for c in row:
    if c.value is not None:orgraw.append(dict(file=rel,sheet=ws.title,cell=c.coordinate,value=str(c.value)))
  assert set(stagehits)=={'primary','junior'} and len(stagehits)==2,(p,stagehits)
  w.close()
 expected=8 if volume==2024 else 12
 assert len(found)==expected and len(set(found))==expected,(volume,found)
orgmap={(str(r['year']),r['stage'],r['fund'],r['organizer_affiliation_category'],r['metric']):r for r in org}
amounts=read('yearbook_stage_spending_candidates.csv')
reconcile=[]
for r in amounts:
 if r['province']!='全国':continue
 f=funds[int(r['part'])-4];cat='全国' if r['affiliation_scope']=='all_affiliations' else '地方'
 o=orgmap.get((r['year'],r['stage'],f,cat,r['metric']))
 if o is None:
  reconcile.append(dict(year=r['year'],stage=r['stage'],fund=f,scope=r['affiliation_scope'],metric=r['metric'],status='not_checked_part1_2024_title_conflict',regional_file=r['file'],regional_cell=r['cell']))
  continue
 a=num(r['value']);b=num(o['value'])
 reconcile.append(dict(year=r['year'],stage=r['stage'],fund=f,scope=r['affiliation_scope'],metric=r['metric'],part1_value=fmt(b),regional_national_value=fmt(a),difference=fmt(None if a is None or b is None else a-b),status='not_checked_original_blank' if a is None or b is None else 'pass' if abs(a-b)<=1 else 'unresolved',part1_file=o['file'],part1_cell=o['cell'],regional_file=r['file'],regional_cell=r['cell']))
coverage=[]
for year in range(2017,2024):
 for stage in ['primary','junior']:
  for fund in funds:
   for allcat,govcat in [('全国','教育和其他部门'),('地方','地方教育和其他部门')]:
    if (str(year),stage,fund,allcat,'total') not in orgmap:continue
    a=num(orgmap[str(year),stage,fund,allcat,'total']['value']);b=num(orgmap[str(year),stage,fund,govcat,'total']['value'])
    assert b<=a+2
    coverage.append(dict(year=year,stage=stage,fund=fund,affiliation='all_affiliations' if allcat=='全国' else 'local_affiliation',all_organizer_total=str(a),government_organizer_total=str(b),outside_government_organizer_amount=str(a-b),outside_government_organizer_pct=str((a-b)/a*100),unit='千元',limit='national coverage diagnostic only; never allocate this share to provinces'))

# National annual CPI, visually transcribed from official 2024 yearbook table 5-1.
cpi={2017:D('101.6'),2018:D('102.1'),2019:D('102.9'),2020:D('102.5'),2021:D('100.9'),2022:D('102.0'),2023:D('100.2')}
factor={2017:D(1)}
for year in range(2018,2024):factor[year]=factor[year-1]*cpi[year]/100
price=ROOT/'raw/fiscal_analysis_cpi_history_2024.jpg'
save_csv(DATA/'fiscal_analysis_cpi.csv',[dict(year=y,cpi_previous_year_100=str(cpi[y]),factor_2017_1=str(factor[y]),source_url='https://www.stats.gov.cn/sj/ndsj/2024/html/C05-01.jpg',source_file=str(price.relative_to(PROJECT)),sha256=sha(price),method='same-analyst visual transcription; national annual CPI, not province or education-input price index') for y in cpi])

series=[]
for r in panel:
 for f in funds:
  v={m:num(r[f+'_'+m]) for m in metrics};career=v['personal']+v['public'];prev=bykey.get((r['province'],str(int(r['year'])-1),r['stage'],r['affiliation_scope']))
  b=bykey[r['province'],'2017',r['stage'],r['affiliation_scope']];y=int(r['year'])
  s=dict(province=r['province'],year=y,stage=r['stage'],affiliation_scope=r['affiliation_scope'],fund=f,total=str(v['total']),personal=str(v['personal']),public=str(v['public']),construction=fmt(v['construction']),career=str(career),unit='千元',cpi_adjusted_total_2017_prices=str(v['total']/factor[y]),cpi_adjusted_career_2017_prices=str(career/factor[y]),nominal_total_index_2017_100=str(v['total']/D(b[f+'_total'])*100),cpi_adjusted_total_index_2017_100=str(v['total']/D(b[f+'_total'])/factor[y]*100),personal_share_total_pct=str(v['personal']/v['total']*100),public_share_total_pct=str(v['public']/v['total']*100),construction_share_total_pct=fmt(None if v['construction'] is None else v['construction']/v['total']*100),personal_share_career_pct=str(v['personal']/career*100),source_file=r[f+'_total_source'],status='conditional descriptive series; user Excel compilation; organizer/year comparability pending',denominator_used=False)
  for m in ['total','career']:
   previous=None if prev is None else D(prev[f+'_total']) if m=='total' else D(prev[f+'_personal'])+D(prev[f+'_public'])
   cur=v['total'] if m=='total' else career
   s['nominal_'+m+'_yoy_pct']=fmt(None if previous is None else pct(cur,previous))
   s['cpi_adjusted_'+m+'_yoy_pct']=fmt(None if previous is None else pct(cur/previous,cpi[y]/100))
  series.append(s)
assert len(series)==2604
national=[]
for y in range(2017,2024):
 for stage in ['primary','junior']:
  for scope,cat in [('all_affiliations','全国'),('local_affiliation','地方')]:
   for f in funds:
    if (str(y),stage,f,cat,'total') in orgmap:
     v={m:num(orgmap[str(y),stage,f,cat,m]['value']) for m in metrics};basis='Part1 national source cells'
    else:
     v={m:num(next(r for r in amounts if r['province']=='全国' and r['year']==str(y) and r['stage']==stage and r['affiliation_scope']==scope and int(r['part'])==funds.index(f)+4 and r['metric']==m)['value']) for m in metrics};basis='regional Part5 national row; Part1 title conflict excluded'
    career=v['personal']+v['public']
    national.append(dict(year=y,stage=stage,affiliation_scope=scope,fund=f,total=str(v['total']),personal=str(v['personal']),public=str(v['public']),construction=fmt(v['construction']),career=str(career),total_100m_yuan=str(v['total']/100000),cpi_adjusted_total_2017_prices_100m_yuan=str(v['total']/factor[y]/100000),personal_share_career_pct=str(v['personal']/career*100),unit_original='千元',source=basis+'; never adds city rows'))
endpoints=[]
for province in sorted({r['province'] for r in panel}|{'全国'}):
 pool=national if province=='全国' else [r for r in series if r['province']==province]
 for stage in ['primary','junior']:
  for scope in ['all_affiliations','local_affiliation']:
   for f in funds:
    rows={int(r['year']):r for r in pool if r['stage']==stage and r['affiliation_scope']==scope and r['fund']==f};a=rows[2017];b=rows[2023]
    endpoints.append(dict(province=province,stage=stage,scope=scope,fund=f,total_2017=a['total'],total_2023=b['total'],nominal_total_change_pct=str(pct(D(b['total']),D(a['total']))),cpi_adjusted_total_change_pct=str(pct(D(b['total'])/factor[2023],D(a['total']))),nominal_career_change_pct=str(pct(D(b['career']),D(a['career']))),cpi_adjusted_career_change_pct=str(pct(D(b['career'])/factor[2023],D(a['career']))),personal_share_career_2017=a['personal_share_career_pct'],personal_share_career_2023=b['personal_share_career_pct'],personal_share_career_change_pp=str(D(b['personal_share_career_pct'])-D(a['personal_share_career_pct'])),status='conditional source series description; not causal'))
outliers=[r for r in series if r['nominal_total_yoy_pct'] and abs(D(r['nominal_total_yoy_pct']))>30]
verification=[]
factor_fraction=Fraction(1)
for y in range(2018,2024):factor_fraction*=Fraction(str(cpi[y]))/100
for r in endpoints:
 exact=(Fraction(r['total_2023'])/Fraction(r['total_2017'])/factor_fraction-1)*100
 observed=D(r['cpi_adjusted_total_change_pct'])
 assert abs(observed-D(exact.numerator)/D(exact.denominator))<D('1e-20')
 verification.append(dict(province=r['province'],stage=r['stage'],scope=r['scope'],fund=r['fund'],kind='Decimal endpoint price adjustment vs exact rational arithmetic',status='pass'))
summary=[]
for stage in ['primary','junior']:
 for scope in ['all_affiliations','local_affiliation']:
  for fund in funds:
   rows=[r for r in endpoints if r['province']!='全国' and r['stage']==stage and r['scope']==scope and r['fund']==fund]
   values=sorted(D(r['cpi_adjusted_total_change_pct']) for r in rows)
   summary.append(dict(stage=stage,scope=scope,fund=fund,provinces=31,cpi_adjusted_total_increase_provinces=sum(D(r['cpi_adjusted_total_change_pct'])>0 for r in rows),cpi_adjusted_total_decrease_provinces=sum(D(r['cpi_adjusted_total_change_pct'])<0 for r in rows),median_cpi_adjusted_total_change_pct=str(values[15]),personal_share_career_increase_provinces=sum(D(r['personal_share_career_change_pp'])>0 for r in rows),limit='equal-weight count/median, not national expenditure growth; annual coverage pending'))
save_csv(DATA/'fiscal_analysis_province_summary.csv',summary)
save_csv(DATA/'fiscal_analysis_formula_checks.csv',verification)
save_csv(DATA/'fiscal_analysis_title_conflicts.csv',ambiguous)
for name,rows in [('fiscal_analysis_quality_checks.csv',qc),('fiscal_analysis_organizer_cells.csv',org),('fiscal_analysis_organizer_source_rows.csv',orgraw),('fiscal_analysis_national_reconciliation.csv',reconcile),('fiscal_analysis_organizer_coverage.csv',coverage),('fiscal_analysis_province_series.csv',series),('fiscal_analysis_national_series.csv',national),('fiscal_analysis_endpoint_changes.csv',endpoints),('fiscal_analysis_large_changes.csv',outliers)]:save_csv(DATA/name,rows)
if not outliers:
 with (DATA/'fiscal_analysis_large_changes.csv').open('w',encoding='utf-8-sig',newline='') as output:
  csv.DictWriter(output,fieldnames=list(series[0])).writeheader()
protect()
audit=dict(candidate_rows=868,province_year_stage=434,series_rows=2604,national_series_rows=len(national),scope_tables_usable=80,scope_tables_title_conflict=len(ambiguous),organizer_cells=len(org),hierarchy_checks=len(qc),hierarchy_pass=sum(r['status']=='pass' for r in qc),hierarchy_unresolved=sum(r['status']=='unresolved' for r in qc),hierarchy_blank_skips=sum(r['status'].startswith('not_checked') for r in qc),national_reconciliation_records=len(reconcile),national_reconciliation_pass=sum(r['status']=='pass' for r in reconcile),national_reconciliation_title_skips=sum(r['status']=='not_checked_part1_2024_title_conflict' for r in reconcile),national_reconciliation_unresolved=sum(r['status']=='unresolved' for r in reconcile),exact_rational_formula_checks=len(verification),large_nominal_changes_over_30_pct=len(outliers),cpi_factor_2023_2017=str(factor[2023]),protected_files_unchanged=True,source_workbooks_unchanged=True,student_denominator_used=False,resource_finance_merge_performed=False,charts_updated=False)
(DATA/'fiscal_analysis_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(audit,ensure_ascii=False,indent=2))
print('National endpoint descriptions:')
for r in endpoints:
 if r['province']=='全国' and r['scope']=='all_affiliations':print(json.dumps(r,ensure_ascii=False))
# A companion notebook exposes the same calculations without creating another path.
nb={'nbformat':4,'nbformat_minor':5,'metadata':{'kernelspec':{'display_name':'Python 3','language':'python','name':'python3'}},'cells':[{'cell_type':'markdown','metadata':{},'source':['# 年鉴实际支出描述核查\n','仅描述整理版财政序列；不反推学生、不合并资源主表、不作因果解释。\n','national CPI仅为通用价格调整敏感性，非教育投入成本指数。\n']},{'cell_type':'code','metadata':{},'execution_count':None,'outputs':[],'source':['from pathlib import Path\n','import runpy\n','script = Path.cwd() / "scripts/analyze_yearbook_spending.py"\n','if not script.exists(): script = Path.cwd() / "research/scripts/analyze_yearbook_spending.py"\n','runpy.run_path(str(script.resolve()))\n']}]}
nb['cells'][1]['source']=['from pathlib import Path\n','import sys\n','script = Path.cwd() / "scripts/analyze_yearbook_spending.py"\n','if not script.exists(): script = Path.cwd() / "research/scripts/analyze_yearbook_spending.py"\n','sys.path.insert(0, str(script.resolve().parent))\n']
nb['cells'].append({'cell_type':'code','metadata':{},'execution_count':None,'outputs':[],'source':(ROOT/'scripts/analyze_yearbook_spending.py').read_text(encoding='utf-8').split('# A companion notebook')[0].splitlines(keepends=True)})
(ROOT/'notebooks/财政支出描述核查.ipynb').write_text(json.dumps(nb,ensure_ascii=False,indent=2),encoding='utf-8')
