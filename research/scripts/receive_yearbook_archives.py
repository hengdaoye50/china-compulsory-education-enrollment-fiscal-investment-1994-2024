"""Receive user-provided annual Excel compilations, preserving every original.
This inventory does not certify publisher provenance or map student coverage.
"""
import csv, hashlib, json, re, stat, subprocess, sys, zipfile
from pathlib import Path, PureWindowsPath
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
PROJECT=ROOT.parent
SOURCE=PROJECT/'中国教育经费统计年鉴'
DEST=ROOT/'用户提供资料'/'年鉴解压'
sys.stdout.reconfigure(encoding='utf-8')
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def target(base,name):
 name=name.replace('\\','/')
 wp=PureWindowsPath(name)
 assert not wp.drive and not name.startswith('/') and '..' not in wp.parts, name
 p=(base/name).resolve();assert p.is_relative_to(base.resolve()), name
 return p
def csvsave(name,rows):
 with (ROOT/'data'/name).open('w',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
archives=[];files=[];headers=[]
for year in range(2018,2025):
 a=next(SOURCE.glob(f'*{year}*'));original_sha=sha(a)
 base=DEST/f'卷{year}';base.mkdir(parents=True,exist_ok=True)
 if a.suffix.lower()=='.zip':
  with zipfile.ZipFile(a) as z:
   members=z.infolist()
   for i in members:
    p=target(base,i.filename)
    assert not stat.S_ISLNK(i.external_attr>>16),i.filename
    assert not i.flag_bits&1,'encrypted archive needs password'
    if i.is_dir():p.mkdir(parents=True,exist_ok=True);continue
    assert p.suffix.lower() in ['.xlsx','.xls','.csv','.pdf','.txt','.doc','.docx','.html','.htm','.jpg','.png'],p
    b=z.read(i) # checks CRC
    p.parent.mkdir(parents=True,exist_ok=True)
    if p.exists():assert p.read_bytes()==b, 'existing extraction differs'
    else:p.write_bytes(b)
 elif a.suffix.lower()=='.rar':
  command=['C:/Program Files/WinRAR/UnRAR.exe']
  listing=subprocess.run(command+['lb','-p-',str(a)],capture_output=True,check=True).stdout.decode('gbk')
  for n in listing.splitlines():
   p=target(base,n)
   assert p.suffix.lower() in ['','.xlsx','.xls','.csv','.pdf','.txt','.doc','.docx','.html','.htm','.jpg','.png'],n
  result=subprocess.run(command+['x','-p-','-o-','-idq',str(a),str(base)+'\\'],capture_output=True)
  assert result.returncode==0,(result.returncode,result.stderr.decode('gbk',errors='replace'))
 else:raise ValueError(a)
 extracted=sorted(p for p in base.rglob('*') if p.is_file())
 assert sha(a)==original_sha
 archives.append(dict(volume_year=year,expected_data_year_pending_check=year-1,archive=str(a.relative_to(PROJECT)),archive_bytes=a.stat().st_size,archive_sha256=original_sha,extracted_files=len(extracted),extracted_bytes=sum(p.stat().st_size for p in extracted),destination=str(base.relative_to(PROJECT)),evidence_type='user-provided Excel compilation; publisher original provenance not yet verified'))
 for p in extracted:
  files.append(dict(volume_year=year,path=str(p.relative_to(PROJECT)),bytes=p.stat().st_size,sha256=sha(p),extension=p.suffix.lower()))
  if p.suffix.lower()=='.xlsx' and ('普通小学' in p.name or '普通初中' in p.name):
   w=load_workbook(p,data_only=True,read_only=True)
   for ws in w:
    top=[f'{c.coordinate}={c.value}' for row in ws.iter_rows(max_row=12) for c in row if c.value is not None]
    headers.append(dict(volume_year=year,file=str(p.relative_to(PROJECT)),sheet=ws.title,rows=ws.max_row,columns=ws.max_column,first_12_rows=' | '.join(top),eligibility='inventory only; actual statistical year and coverage pending source checks'))
   w.close()
 print(f'卷{year}: {len(extracted)} files received',flush=True)
csvsave('yearbook_archive_receipt.csv',archives)
csvsave('yearbook_extracted_file_inventory.csv',files)
csvsave('yearbook_stage_table_headers.csv',headers)
(ROOT/'data/yearbook_archive_receipt.json').write_text(json.dumps(archives,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(dict(volumes=len(archives),files=len(files),stage_table_sheets=len(headers),all_original_archives_preserved=True),ensure_ascii=False))
