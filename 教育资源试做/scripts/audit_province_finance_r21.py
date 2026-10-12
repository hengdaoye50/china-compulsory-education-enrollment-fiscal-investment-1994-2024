"""Round 21 source audit. Writes separate ledgers; never reconstructs denominators.

Guizhou four amounts and four ratios are human visual transcriptions from
physical PDF pages 4/5. Re-running arithmetic is not independent visual review.
"""
import csv, hashlib, json, re, sys
from decimal import Decimal as D
import pdfplumber
from openpyxl import load_workbook
from parse import ROOT, save_csv
sys.stdout.reconfigure(encoding='utf-8')
DATA=ROOT/'data'
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def read(p): return list(csv.DictReader(p.open(encoding='utf-8-sig',newline='')))
baseline=json.loads((DATA/'province_finance_next_protected_baseline.json').read_text(encoding='utf-8'))
assert all(sha(ROOT/p)==v for p,v in baseline.items())
sources=[]
for n in ['province_finance_r21_sources','province_finance_r21_attachment_sources','province_finance_r21_access_sources']:
 sources+=json.loads((ROOT/(n+'.json')).read_text(encoding='utf-8'))
src={s['key']:s for s in sources}
receipt=[]
for s in sources:
 row={k:s.get(k,'') for k in ['key','url','status','bytes','sha256','path','error']}
 row['evidence_type']='official institutional republication' if '_cee' in s['key'] else 'publisher bibliographic metadata' if 'publisher' in s['key'] else 'library catalog metadata' if 'library_catalog' in s['key'] else 'library access page attempt' if 'xmu_cnki' in s['key'] else 'government original'
 if s['status']=='downloaded':
  p=ROOT/s['path'];assert sha(p)==s['sha256'] and p.stat().st_size==s['bytes']
  if p.suffix=='.pdf': assert p.read_bytes().startswith(b'%PDF')
  if p.suffix=='.xlsx': assert p.read_bytes().startswith(b'PK')
 receipt.append(row)
save_csv(DATA/'province_finance_r21_source_receipt.csv',receipt)
for province in ['anhui_2023','guangdong_2022']:
 assert src[f'province_finance_r21_{province}_xlsx']['sha256']==src[f'province_finance_r21_{province}_cee_xlsx']['sha256']

inventory=[];export=[];books={}
for key in ['province_finance_r21_anhui_2023_xlsx','province_finance_r21_guangdong_2022_xlsx']:
 w=load_workbook(ROOT/src[key]['path'],data_only=True);books[key]=w
 for ws in w:
  occupied=[(c.coordinate,str(c.value)) for row in ws for c in row if c.value is not None]
  joined=' '.join(v for _,v in occupied)
  inventory.append(dict(source_key=key,sheet=ws.title,rows=ws.max_row,columns=ws.max_column,nonempty_cells=len(occupied),headers=' | '.join(f'{c.coordinate}={c.value}' for row in list(ws)[:7] for c in row if c.value is not None),student_field_hits=';'.join(sorted(set(re.findall(r'.{0,15}(?:学生数|学生人数|年初学生|年末学生|加权|权重).{0,25}',joined)))),finding='all-stage amount table and stage per-student tables; no stage expenditure total/student denominator/weight column found in inspected workbook'))
  export+= [dict(source_key=key,sheet=ws.title,cell=a,value=v) for a,v in occupied]
save_csv(DATA/'province_finance_r21_sheet_inventory.csv',inventory)
save_csv(DATA/'province_finance_r21_xlsx_cells.csv',export)
pdf_inventory=[];nx={}
for i in range(1,5):
 key=f'province_finance_r21_ningxia_2023_table{i}_pdf'
 with pdfplumber.open(ROOT/src[key]['path']) as p:
  assert len(p.pages)==1;t=p.pages[0].extract_text();nx[i]=t
  pdf_inventory.append(dict(source_key=key,physical_page=1,text=t,finding='all-stage total' if i==1 else 'per-student stage table; no total/student/weight column in reviewed page'))
save_csv(DATA/'province_finance_r21_pdf_text_inventory.csv',pdf_inventory)
cells=[]
def add(province,year,stage,metric,value,growth,key,locator,method):
 cells.append(dict(province=province,year=year,stage=stage,metric=metric,value=str(value),published_growth_pct=str(growth),unit='元',source=src[key]['url'],source_key=key,locator=locator,method=method,vintage='current-year column in presently available release; first-release identity not separately verified',eligibility='published per-student only; fiscal denominator not verified'))
metrics=['budget_per_student','career_per_student','public_per_student']
key='province_finance_r21_anhui_2023_xlsx'
for i,metric in enumerate(metrics,1):
 ws=books[key].worksheets[i];assert '生均' in ws['A1'].value and ws['D3'].value=='普通小学' and ws['F3'].value=='普通初中'
 for stage,col in [('primary',4),('junior',6)]:add('安徽',2023,stage,metric,ws.cell(5,col).value,ws.cell(5,col+1).value,key,f'{ws.title}!{ws.cell(5,col).coordinate}; A1/D3/F3/D4/F4','original XLSX cell')
key='province_finance_r21_guangdong_2022_xlsx';ws=books[key].worksheets[1]
assert ws['F3'].value=='普通初中' and ws['H3'].value=='普通小学' and '2022年' in ws['A1'].value
for stage,col in [('primary',8),('junior',6)]:add('广东',2022,stage,metrics[0],ws.cell(5,col).value,ws.cell(5,col+1).value,key,f'{ws.title}!{ws.cell(5,col).coordinate}; F3/H3; reverse stage order','original XLSX cell')
for i,metric in enumerate(metrics,2):
 t=nx[i];assert '普通初中' in t and '普通小学' in t and '单位：元' in t
 values=next(l for l in t.splitlines() if l.startswith('宁夏回族自治区 ')).split()[1:];assert len(values)==12
 for stage,col in [('primary',8),('junior',6)]:add('宁夏',2023,stage,metric,values[col].replace(',',''),values[col+1],f'province_finance_r21_ningxia_2023_table{i}_pdf',f'physical page 1; 宁夏回族自治区 row; 普通初中 precedes 普通小学','PDF text with original page visual header/row verification')
save_csv(DATA/'province_finance_r21_candidate_cells.csv',cells)
national={(r['province'],r['year'],r['stage'],r['metric']):r for r in read(DATA/'finance_long.csv')}
checks=[]
for c in cells:
 n=national[(c['province'],str(c['year']),c['stage'],c['metric'])]
 for f in ['value','published_growth_pct']:
  diff=D(c[f])-D(n[f]);checks.append(dict(province=c['province'],year=c['year'],stage=c['stage'],metric=c['metric'],field=f,local_value=c[f],national_value=n[f],difference=str(diff),status='pass' if abs(diff)<=D('0.005') else 'unresolved',source=c['source'],locator=c['locator']))
save_csv(DATA/'province_finance_r21_cross_source_checks.csv',checks)

# Full-province functional-classification finals, separate from education-finance survey.
functional=[]
for year,amounts,ratios in [(2022,[3750238,2136362],['99.7','100.0']),(2023,[3919844,2361334],['104.5','110.5'])]:
 key=f'province_finance_r21_guizhou_functional_{year}_pdf'
 with pdfplumber.open(ROOT/src[key]['path']) as p:
  pages=len(p.pages);textchars=sum(len(x.extract_text() or '') for x in p.pages)
 for stage,name,code,page,value,ratio in zip(['primary','junior'],['小学教育','初中教育'],['2050202','2050203'],[4,5],amounts,ratios):
  functional.append(dict(province='贵州',year=year,stage=stage,metric='general_public_budget_functional_final_expenditure',code=code,code_basis='page 4: 205教育支出/02普通教育/02小学教育; page 5 continuation 03初中教育',row_name=name,value=value,unit='万元',published_current_over_prior_final_pct=ratio,source=src[key]['url'],source_sha256=src[key]['sha256'],physical_page=page,table_caption=f'表2 贵州省{year}年全省一般公共预算支出决算表',amount_column='决算数',coverage='full province as original table caption; not provincial government only',vintage='current final column of currently available original government PDF',pdf_pages=pages,pdf_extracted_text_chars=textchars,method='human visual transcription at 2x; no OCR and no independent second reviewer',eligibility='separate functional expenditure candidate; not education-finance survey total; no matched student denominator',definition_status='annual functional-code definition and organizer/school-stage coverage bridge not yet verified',publication_bridge_status='government PDF directly obtained; parent publication page not yet verified'))
save_csv(DATA/'province_finance_r21_functional_candidates.csv',functional)
continuity=[]
for stage in ['primary','junior']:
 a=next(x for x in functional if x['stage']==stage and x['year']==2022);b=next(x for x in functional if x['stage']==stage and x['year']==2023)
 ratio=D(b['value'])/D(a['value'])*100;diff=ratio-D(b['published_current_over_prior_final_pct'])
 continuity.append(dict(province='贵州',stage=stage,prior_year=2022,current_year=2023,prior_original_final=a['value'],current_original_final=b['value'],computed_current_over_prior_final_pct=str(ratio),published_ratio_pct=b['published_current_over_prior_final_pct'],difference_percentage_points=str(diff),tolerance_percentage_points='0.05',status='pass' if abs(diff)<=D('.05') else 'unresolved',boundary='two original files match one-decimal published ratio; does not verify fiscal students or classify all education expenditure'))
save_csv(DATA/'province_finance_r21_functional_continuity_checks.csv',continuity)
bridges=[]
for pk,ak,token in [
 ('province_finance_r21_anhui_2023_original','province_finance_r21_anhui_2023_xlsx','wKg86menG1qAZrN7AADme8oj7LQ95.xlsx'),
 ('province_finance_r21_guangdong_2022','province_finance_r21_guangdong_2022_xlsx','539530/4324695.xlsx'),
 *[('province_finance_r21_ningxia_2023',f'province_finance_r21_ningxia_2023_table{i}_pdf',src[f'province_finance_r21_ningxia_2023_table{i}_pdf']['url'].split('/')[-1]) for i in range(1,5)]]:
 assert token in (ROOT/src[pk]['path']).read_text(encoding='utf-8')
 bridges.append(dict(page_source=src[pk]['url'],attachment_source=src[ak]['url'],token=token,status='pass'))
save_csv(DATA/'province_finance_r21_publication_bridge.csv',bridges)
coverage=[]
for province,year,depth in [('安徽',2023,'government original notice + four-sheet original; mirrored attachment identical'),('广东',2022,'government notice + two-sheet original; mirrored attachment identical'),('宁夏',2023,'government notice + four one-page PDF tables')]:
 coverage.append(dict(province=province,year=year,local_source_depth=depth,stage_actual_spending_verified=False,paired_fiscal_students_verified=False,student_weights_verified=False,classification='per-student only',boundary='only inspected listed attachments; no exhaustive negative claim'))
for year in [2022,2023]:coverage.append(dict(province='贵州',year=year,local_source_depth='full-province functional final PDF; physical pages 4/5 visually verified',stage_actual_spending_verified='functional code only; survey total not verified',paired_fiscal_students_verified=False,student_weights_verified=False,classification='separate functional expenditure candidate',boundary='do not merge functional candidates into education-finance main series'))
save_csv(DATA/'province_finance_r21_coverage.csv',coverage)
registry=read(DATA/'province_finance_next_2017_2023_registry.csv')
for r in registry:
 match=next((c for c in coverage if c['province']==r['province'] and str(c['year'])==r['year']),None)
 if match:
  r['local_source_depth']=match['local_source_depth'];r['note']=(r['note']+'; round21 audit: '+match['classification']+'; see separate r21 ledgers').strip('; ')
save_csv(DATA/'province_finance_r21_2017_2023_registry.csv',registry)
scope=[dict(province='广东',year=2022,finding='notice footnote: education-finance statistic 3863.13亿元 vs fiscal-final education expenditure 3871.14亿元',boundary='different statistical coverage; neither is stage expenditure; preserve prior same-scope adjustment footnote',source=src['province_finance_r21_guangdong_2022']['url']),dict(province='贵州',year='2022/2023',finding='full-province finals directly report functional codes 2050202/2050203; 20509 education surcharge and 20599 other education shown separately',boundary='stage function row not exhaustive education-finance stage spending; annual code definition/coverage bridge/student denominator pending',source=src['province_finance_r21_guizhou_functional_2023_pdf']['url'])]
save_csv(DATA/'province_finance_r21_scope_evidence.csv',scope)
assert all(sha(ROOT/p)==v for p,v in baseline.items())
fiscal=[s for s in sources if not any(t in s['key'] for t in ['yearbook','xmu_cnki'])]
audit=dict(download_records=len(sources),saved_download_records=sum(s['status']=='downloaded' for s in sources),unique_saved_contents=len({s['sha256'] for s in sources if s['status']=='downloaded'}),failed_download_records=sum(s['status']!='downloaded' for s in sources),fiscal_source_download_records=len(fiscal),fiscal_source_unique_contents=len({s['sha256'] for s in fiscal if s['status']=='downloaded'}),fiscal_source_all_downloaded=all(s['status']=='downloaded' for s in fiscal),workbook_sheets=len(inventory),ningxia_original_pdf_pages=4,local_per_student_province_years=3,candidate_cells=len(cells),cross_source_checks=len(checks),cross_source_pass=sum(c['status']=='pass' for c in checks),cross_source_unresolved=[c for c in checks if c['status']!='pass'],verified_publication_bridges=len(bridges),functional_direct_amounts=len(functional),functional_continuous_stage_candidates=2,functional_ratio_checks=len(continuity),functional_ratio_all_pass=all(c['status']=='pass' for c in continuity),matched_continuous_survey_spending_student_pairs=0,complete_student_weight_rule_found=False,registry_rows=len(registry),protected_files_unchanged=True,protected_hashes=baseline,review_type='same analyst source/cell/arithmetic and visual checks; no independent review')
(DATA/'province_finance_r21_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(audit,ensure_ascii=False,indent=2))
