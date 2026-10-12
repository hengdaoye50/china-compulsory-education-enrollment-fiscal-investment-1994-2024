"""Audit integrated deliverables and make page contact sheets for human review."""
import hashlib,json,re,shutil
from pathlib import Path
import pandas as pd
from PIL import Image,ImageDraw
from pypdf import PdfReader
from docx import Document
ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'最终综合研究报告_20261011'
QA=OUT/'最终版式审校'
name='生源变动与义务教育经费长期演进最终报告'
doc=Document(OUT/(name+'.docx'))
pdf=QA/(name+'.pdf')
reader=PdfReader(pdf)
text='\n'.join(p.extract_text() for p in reader.pages)
compact=re.sub(r'\s+','',text)
checks=[]
def check(label,ok):
    checks.append({'check':label,'passed':bool(ok)})
    assert ok,label
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
receipt=json.loads((OUT/'报告合稿记录.json').read_text(encoding='utf8'))
check('23_pages',len(reader.pages)==23)
check('no_blank_or_orphan_pages',all(len(p.extract_text())>300 for p in reader.pages))
check('15_figures_15_tables',len(doc.inline_shapes)==15 and len(doc.tables)==15)
for kind,count in [('图',15),('表',15)]:
    labels=[int(re.match(kind+r'(\d+)\s',p.text).group(1)) for p in doc.paragraphs if re.match(kind+r'(\d+)\s',p.text)]
    check(kind+'_numbering',labels==list(range(1,count+1)))
for p,h in receipt['input_sha256'].items():check('unchanged:'+p,sha(ROOT/p)==h)
for f in receipt['figure_map']:check('figure_hash:'+str(f['number']),sha(ROOT/f['path'])==f['sha256'])
checks_from_build=pd.read_csv(OUT/'合稿数值检查.csv')
check('all_28_build_checks',len(checks_from_build)==28 and checks_from_build.passed.all())
check('old_awkward_region_label_absent','同一28省' not in compact)
check('regression_equation_rendered','βS' in compact and 'βR' in compact and 'ε' in compact)
for value in ['0.601','0.0144','0.766','0.506','0.0008','0.0033','399.27','257.35','594','1,140,343','472','300','436','71.4%']:
    check('key_number_rendered:'+value,value in compact)
for n in range(1,9):
    old=pd.read_csv(ROOT/f'教育资源试做/正式学术图组_20261011/figure_data/fig{n:02d}.csv')
    new=pd.read_csv(OUT/f'图表_合稿版/figure_data/fig{n:02d}.csv')
    pd.testing.assert_frame_equal(old,new)
    check('reexport_data_identity:'+str(n),True)
manifest=json.loads((OUT/'图表_合稿版/figure_manifest.json').read_text(encoding='utf8'))
check('plot_export_checks',all(manifest['checks'].values()))
for f in receipt['figure_map']:
    with Image.open(ROOT/f['path']) as im:
        dpi=im.info.get('dpi',(0,0))
        check('figure_600dpi:'+str(f['number']),all(abs(x-600)<1 for x in dpi))
        check('figure_white_corner:'+str(f['number']),im.convert('RGB').getpixel((0,0))==(255,255,255))
pages=sorted(QA.glob('page-*.png'),key=lambda p:int(p.stem.split('-')[1]))
check('all_page_images',len(pages)==23 and all(p.stat().st_size>10000 for p in pages))
overview=QA/'逐页总览';overview.mkdir(exist_ok=True)
for first in range(0,len(pages),4):
    canvas=Image.new('RGB',(1200,1610),'white')
    draw=ImageDraw.Draw(canvas)
    for k,p in enumerate(pages[first:first+4]):
        with Image.open(p) as im:canvas.paste(im.resize((600,776)),((k%2)*600,(k//2)*805+25))
        draw.text(((k%2)*600+8,(k//2)*805+5),str(first+k+1),fill='black')
    canvas.save(overview/f'总览{first//4+1}.png')
pd.DataFrame(checks).to_csv(OUT/'最终交付检查.csv',index=False,encoding='utf-8-sig')
shutil.copy2(pdf,OUT/(name+'.pdf'))
result=dict(pages=23,figures=15,tables=15,checks=len(checks),all_passed=True,visual_review='pending',docx_sha256=sha(OUT/(name+'.docx')),pdf_sha256=sha(OUT/(name+'.pdf')),no_new_estimation=True)
(OUT/'最终交付记录.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(result,ensure_ascii=False))
