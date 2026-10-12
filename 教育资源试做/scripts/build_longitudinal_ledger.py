"""Build separate, traceable candidate ledgers; never interpolate or run models.
Run with the bundled Python. Source values and accepted values are separate.
"""
from pathlib import Path
from decimal import Decimal, InvalidOperation
from collections import Counter, defaultdict
import csv, hashlib, io, json, re, sys, zipfile
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT / '教育资源试做'
AUDIT = BASE / 'data/年鉴资料核查_20261010'
OUT = BASE / 'data/长期财政数据建设_20261010'
OUT.mkdir(exist_ok=True)
sys.path.insert(0, str(BASE / '.tmp/yearbook_read_deps'))
sys.stdout.reconfigure(encoding='utf-8')
PROVS = '北京 天津 河北 山西 内蒙古 辽宁 吉林 黑龙江 上海 江苏 浙江 安徽 福建 江西 山东 河南 湖北 湖南 广东 广西 海南 重庆 四川 贵州 云南 西藏 陕西 甘肃 青海 宁夏 新疆'.split()
CITIES = '大连 宁波 厦门 青岛 深圳'.split()
STAGES = {'小学': 'primary', '初中': 'junior'}
PROTECTED = [BASE/'data/finance_long.csv', BASE/'data/education_resource_panel.csv', BASE/'data/panel.json', BASE/'教育资源试做.html']
def sha(b): return hashlib.sha256(b).hexdigest()
protected_before = {str(p.relative_to(ROOT)): sha(p.read_bytes()) for p in PROTECTED}
def norm(v): return re.sub(r'\s+', '', str(v or ''))
def raw(v): return '' if v is None else str(v)
def parse(v):
    s = raw(v).strip()
    if not s: return None, 'missing_blank'
    if s in ['-', '--', '—', '…', '...', 'nan']: return None, 'missing_symbol_unresolved'
    # Only explicit numeric notation; do not replace letters, punctuation, or internal spaces.
    if not re.fullmatch(r'[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?', s): return None, 'invalid_numeric_text'
    try: n = Decimal(s.replace(',', ''))
    except InvalidOperation: return None, 'invalid_numeric_text'
    return n, 'zero_observed' if n == 0 else 'numeric_observed'
def write(name, rows, fields=None):
    if fields is None: fields = list(dict.fromkeys(k for r in rows for k in r))
    with (OUT/name).open('w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=fields); w.writeheader(); w.writerows(rows)
def readcsv(p):
    with p.open(encoding='utf-8-sig', newline='') as f: return list(csv.DictReader(f))
def decoded(i):
    s = i.filename
    if not i.flag_bits & 0x800:
        try: s = s.encode('cp437').decode('gbk')
        except UnicodeError: pass
    return s
archives = {}
def getbytes(archive, table):
    if archive not in archives:
        z = zipfile.ZipFile(ROOT/archive)
        archives[archive] = (z, {decoded(i): i for i in z.infolist() if not i.is_dir()})
    z, idx = archives[archive]
    return z.read(idx[table])
def workbook(b):
    if b[:2] == b'PK':
        w = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
        result = [(s.title, list(s.iter_rows(values_only=True))) for s in w.worksheets]
        w.close(); return result
    import xlrd
    w = xlrd.open_workbook(file_contents=b)
    return [(s.name, [s.row_values(i) for i in range(s.nrows)]) for s in w.sheets()]
def region(r):
    s = norm(' '.join(raw(v) for v in r[:2]))
    if re.match(r'^(?:总计|合计|全国|总数|Total)', norm(r[0]), re.I): return '全国'
    # Full-name starts prevent a province name appearing in an unrelated label.
    return next((p for p in PROVS+CITIES if s.startswith(p)), None)
ledger, sources, checks, mappings, issues = [], [], [], [], []
source_seen = set()
def source(family, volume, year, archive, table, b, sheet, rows, year_status, scope, block=0):
    sid = sha((archive+'|'+table+'|'+sheet+'|'+str(block)).encode())[:20]
    if sid not in source_seen:
        sources.append(dict(source_id=sid, family=family, volume_year=volume, data_year=year,
                            data_year_status=year_status, archive=archive, table=table, sheet=sheet,
                            table_sha256=sha(b), row_count=len(rows), table_block=block, scope=scope,
                            source_quality='user_supplied_excel_conversion; publisher_cells_not_independently_verified'))
        source_seen.add(sid)
    return sid
def cell(sid, year, stage, province, metric, row, col, v, unit, school, funding, label='', extra=''):
    n, status = parse(v)
    rid = sha(f'{sid}|{row}|{col}|{metric}'.encode()).__str__()[:24]
    rec = dict(record_id=rid, source_id=sid, data_year=year, province=province, stage=stage,
               metric=metric, source_row=row, source_column=col, cell=f'{openpyxl.utils.get_column_letter(col)}{row}',
               raw_value=raw(v), numeric_candidate='' if n is None else str(n), accepted_value='',
               value_status=status, unit=unit, school_scope=school, funding_scope=funding,
               original_label=label, arithmetic_status='not_checked', bridge_status='unverified',
               model_eligible=False, verification_basis='', notes=extra)
    ledger.append(rec); return rec
def issue(kind, sid, year, stage, province, detail, severity='high'):
    issues.append(dict(issue_id=f'I{len(issues)+1:06d}', kind=kind, source_id=sid, data_year=year,
                       stage=stage, province=province, severity=severity, status='open', detail=detail,
                       correction='', evidence_required='original image/independent source/scope note as applicable'))

# A. Students: re-read source workbooks, not truncated sample rows.
edu = json.loads((AUDIT/'education_province_samples.json').read_text(encoding='utf8'))
student_rows = []
for spec in sorted(edu, key=lambda a:(a['data_year_label'],a['stage'])):
    year = spec['data_year_label']; stage = STAGES[spec['stage']]
    if year < 1995: continue
    b = getbytes(spec['archive'], spec['table']); sheet, rows = workbook(b)[0]
    rr = [(i+1,r,region(r)) for i,r in enumerate(rows) if region(r)]
    first = min(i for i,r,p in rr)
    headers = rows[:first-1]
    # Early junior files combine junior and senior. Verified grade-column layout overrides merged-title offsets.
    if stage == 'junior' and year <= 2002:
        total_col = {1995:4,1996:4,1997:3,1998:4,1999:4,2000:3,2001:3,2002:4}[year]
        grades = list(range(total_col+1,total_col+5))
        basis = 'explicit junior subtotal and four grades; inspected header rows; not secondary total'
    else:
        hits = [(i+1,j+1,raw(v)) for i,r in enumerate(headers) for j,v in enumerate(r)
                if re.search(r'在校(?:学)?生|在校学', norm(v)) and j>0]
        hits = [h for h in hits if not re.search('预计|毕业班', norm(h[2]))]
        assert len(hits)==1, (year,stage,hits)
        total_col = hits[0][1]
        grades = sorted(set(j+1 for r in headers for j,v in enumerate(r)
                            if re.search(r'[一二三四五六]年级', norm(v)) and j+1>total_col))
        basis = 'unique enrolment header with separately labelled grade columns'
    assert len(grades) == (6 if stage=='primary' else 4), (year,stage,grades)
    sid = source('education_students',year,year,spec['archive'],spec['table'],b,sheet,rows,
                 'archive_label_current_year; table_time_definition_pending','all_organizers_candidate')
    mappings.append(dict(source_id=sid, metric='S', total_column=total_col, component_columns=json.dumps(grades),
                         basis=basis, year=year, stage=stage, header_evidence=json.dumps(headers,ensure_ascii=False)))
    vals = {}; records = {}
    for row,r,p in rr:
        if p in CITIES: continue
        s = cell(sid,year,stage,p,'S',row,total_col,r[total_col-1],'人','all_organizers_candidate','not_applicable','enrolment',
                 'ordinary/other school types, attached classes and timing pending scope notes')
        gs = [cell(sid,year,stage,p,f'S_grade{k}',row,col,r[col-1],'人',s['school_scope'],'not_applicable',f'grade{k}')
              for k,col in enumerate(grades,1)]
        nums = [parse(r[col-1])[0] for col in [total_col]+grades]
        diff = nums[0]-sum(nums[1:],Decimal(0)) if all(n is not None for n in nums) else None
        state = 'pass_exact' if diff==0 else 'unresolved_difference' if diff is not None else 'uncheckable_invalid_or_missing'
        s['arithmetic_status'] = state
        checks.append(dict(check='student_grade_sum',source_id=sid,year=year,stage=stage,province=p,
                           difference='' if diff is None else str(diff),tolerance=0,status=state))
        if state != 'pass_exact': issue('student_grade_sum',sid,year,stage,p,f'{state}; difference={diff}; raw total={s["raw_value"]}')
        if nums[0] is not None and (nums[0]<=0 or nums[0] != nums[0].to_integral_value()):
            issue('invalid_student_domain',sid,year,stage,p,'student total must be positive integer')
        vals[p] = nums[0]; records[p] = s
        student_rows.append(s)
    expected = [p for p in PROVS if not(year<1997 and p=='重庆')]
    assert sorted(p for p in vals if p!='全国') == sorted(expected), (year,stage,'region coverage')
    diff = vals.get('全国')-sum((vals[p] for p in expected),Decimal(0)) if vals.get('全国') is not None and all(vals[p] is not None for p in expected) else None
    state = 'pass_exact' if diff==0 else 'unresolved_difference' if diff is not None else 'uncheckable_invalid_or_missing'
    checks.append(dict(check='student_province_sum',source_id=sid,year=year,stage=stage,province='全国',difference='' if diff is None else str(diff),tolerance=0,status=state))
    if state!='pass_exact': issue('student_province_sum',sid,year,stage,'全国',f'{state}; difference={diff}')
    for p,s in records.items():
        if s['arithmetic_status']=='pass_exact' and state=='pass_exact':
            s['accepted_value']=s['numeric_candidate'];s['verification_basis']='grade sum and province sum exact; source-internal numeric acceptance only'
            s['bridge_status']='numeric_internal_checks_pass; school_scope_pending'
    print('students',year,stage,len(expected),state,flush=True)

# B. Supplement original sheet: record every source row and actual year, including blanks.
panel_spec=json.loads((AUDIT/'supplement_panel_coverage.json').read_text(encoding='utf8'))
panel_path=ROOT/panel_spec['source']
w=openpyxl.load_workbook(panel_path,read_only=True,data_only=True)
ws=w['原始数据-长面板']; iterator=ws.iter_rows(values_only=True); head=next(iterator)
panel_hash=sha(panel_path.read_bytes())
sid='supplement_original_'+panel_hash[:16]
sources.append(dict(source_id=sid,family='supplement_original',volume_year='',data_year='1995-2024',data_year_status='column_header',archive=panel_spec['source'],table=panel_spec['source'],sheet=ws.title,table_sha256=panel_hash,row_count=ws.max_row,scope='province_annual_original_sheet',source_quality='user_compiled_NBS_panel; original_sheet_only; revisions_pending'))
metrics={'地方财政一般预算收入':'R','地方财政一般预算支出':'fiscal_expenditure','地区生产总值':'GDP','居民消费价格指数(上年=100)':'CPI_yoy'}
panel_rows=[]
for ri,r in enumerate(iterator,2):
    if r[4] not in metrics: continue
    # Multiple CPI categories share a short name: select only the overall series.
    if r[4]=='居民消费价格指数(上年=100)' and str(r[5]) != '价格指数-居民消费价格指数和商品零售价格指数(上年=100)-居民消费价格指数(上年=100)':
        continue
    p=next((p for p in PROVS if norm(r[1]).startswith(p)),None)
    if not p:continue
    for ci,h in enumerate(head,1):
        if not str(h).isdigit() or not 1995<=int(h)<=2024: continue
        s=cell(sid,int(h),'province_only',p,metrics[r[4]],ri,ci,r[ci-1],raw(r[6]) if r[4]!='居民消费价格指数(上年=100)' else '上年=100','not_applicable','general_budget_label_pending' if metrics[r[4]] in ['R','fiscal_expenditure'] else 'not_applicable',raw(r[5]),'compiled original layer; cross-source verification pending')
        if p=='重庆' and int(h)<1997:
            s['value_status']='geography_not_applicable_observation_present' if s['numeric_candidate'] else 'not_applicable_geography'
            issue('pre_1997_chongqing',sid,int(h),'province_only',p,'original compiled row cannot represent independent Chongqing before establishment')
        elif s['numeric_candidate'] and Decimal(s['numeric_candidate'])<=0:
            issue('nonpositive_panel_value',sid,int(h),'province_only',p,f'{s["metric"]}={s["raw_value"]}')
        panel_rows.append(s)
w.close()
print('supplement records',len(panel_rows), 'metrics',dict(Counter(s['metric'] for s in panel_rows)),flush=True)

# C. Fiscal final-account tables: source labels and final-account cells, with budget columns excluded.
fiscal_specs=json.loads((AUDIT/'fiscal_province_scope.json').read_text(encoding='utf8'))
fiscal_records=[]
for spec in fiscal_specs:
    if not 1996<=spec['volume_year']<=2025:continue
    b=getbytes(spec['archive'],spec['table']); sheet,rows=workbook(b)[0]
    p=spec['province']; title=' '.join(raw(v) for r in rows[:4] for v in r)
    years=set(int(v) for v in re.findall(r'(?:19|20)\d{2}',title))
    filename_years=set(int(v) for v in re.findall(r'(?:19|20)\d{2}',Path(spec['table']).name))
    yset=years or filename_years
    year=next(iter(yset)) if len(yset)==1 else spec['volume_year']-1
    if not 1995<=year<=2024:continue
    year_state='table_or_filename_explicit' if len(yset)==1 else 'volume_minus_one_pending'
    sid=source('fiscal_final_account',spec['volume_year'],year,spec['archive'],spec['table'],b,sheet,rows,year_state,'province_named; full_province_vs_level_and_cities_pending')
    fulltext=' '.join(raw(v) for r in rows for v in r)
    unit_hits=re.findall(r'单位\s*[:：]?\s*(亿\s*元|万\s*元|千\s*元|元)',fulltext)
    unit=norm(unit_hits[0]) if len(set(norm(u) for u in unit_hits))==1 else 'unit_pending'
    active_cols=[]; block=1
    for ri,r in enumerate(rows,1):
        if any('决算数' in norm(v) for v in r): active_cols=[ci+1 for ci,v in enumerate(r) if '决算数' in norm(v)]
        if ri>4 and any(re.search('收支决算总表|FINALACCOUNT',norm(v),re.I) for v in r): block+=1
        for ci,v in enumerate(r,1):
            label=norm(v)
            if not re.search('中央补助|税收返还|一般性转移|专项转移|上解中央|地方财政收入|一般公共预算收入|一般预算收入|本年收入合计|收入总计',label):continue
            if re.search('决算总表|统计表|单位|说明',label):continue
            next_label=next((j for j in range(ci+1,len(r)+1) if re.search(r'[\u4e00-\u9fff]',raw(r[j-1]))),len(r)+1)
            candidates=[col for col in active_cols if ci<col<next_label]
            if len(candidates)!=1:
                issue('fiscal_final_column_unresolved',sid,year,'province_only',p,f'row={ri}; label={raw(v)}; final columns={candidates}')
                continue
            col=candidates[0]
            metric=('central_subsidy' if '中央补助' in label else 'tax_return' if '税收返还' in label else 'general_transfer' if '一般性转移' in label else 'special_transfer' if '专项转移' in label else 'remittance_central' if '上解中央' in label else 'revenue_total_not_available_resources' if '收入总计' in label else 'R_fiscal_candidate')
            s=cell(sid,year,'province_only',p,metric,ri,col,r[col-1],unit,'not_applicable','fiscal_table_scope_pending',raw(v),f'block={block}; containing items must not be added; geography, level and city coverage pending')
            if spec['volume_year']==2008 and p=='河北':
                s['value_status']='conflicting_table_blocks_pending';issue('hebei_tianjin_conflict',sid,year,'province_only',p,f'row={ri}; preserve both blocks, select neither pending independent check')
            fiscal_records.append(s)
    if unit=='unit_pending':issue('fiscal_unit_unresolved',sid,year,'province_only',p,'unit not uniquely present in converted table')
print('fiscal source tables',sum(s['family']=='fiscal_final_account' for s in sources),'cells',len(fiscal_records),flush=True)

# D. Long-run education spending and income source cells from archive tables.
catalog=readcsv(AUDIT/'all_archive_tables.csv')
expense_records=[]; expense_tables=[]
def finance_blocks(b):
    for sheet,fullrows in workbook(b):
        starts=[0];seen_regions=False
        for i,r in enumerate(fullrows):
            if region(r):seen_regions=True
            if seen_regions and norm(r[0]) in ['地区','地区Region','Region']:
                starts.append(i);seen_regions=False
        for block,start in enumerate(starts):
            end=starts[block+1] if block+1<len(starts) else len(fullrows)
            yield sheet,fullrows[start:end],start,block,fullrows
for spec in catalog:
    if spec['family']!='中国教育经费统计年鉴':continue
    name=Path(spec['table']).name
    if not re.search(r'\.(xls|xlsx)$',name,re.I) or not re.search('普通小学|普通初中|初级中学|普通中学',name):continue
    if re.search('中央|农村|生均',name) or not re.search('支出|收入',name):continue
    volume=int(spec['volume_year']); year=volume-1
    if not 1995<=year<=2024:continue
    b=getbytes(spec['archive'],spec['table'])
    for sheet,rows,offset,block,fullrows in finance_blocks(b):
        rr=[(i+1,r,region(r)) for i,r in enumerate(rows) if region(r)]
        if not rr:continue
        first=min(i for i,r,p in rr); headers=rows[:first-1]
        ht=' '.join(raw(v) for r in headers for v in r)
        stage='primary' if '普通小学' in name else 'junior' if re.search('普通初中|初级中学',name) else 'secondary_combined_excluded'
        school='local_affiliation' if '地方' in name else 'all_affiliations_candidate'
        funding='all_source' if (re.search('教育经费.*支出',name) and not re.search('预算|补助',name)) else 'education_income' if '收入' in name else 'budget_or_subsidy_original_classification'
        sid=source('education_finance',volume,year,spec['archive'],spec['table'],b,sheet,rows,
                   'prior_2017_2023_evidence' if 2017<=year<=2023 else 'volume_minus_one_pending_table_notes',school,block)
        fulltext=norm(' '.join(raw(v) for r in fullrows for v in r))
        unit='千元' if '千元' in fulltext else '万元' if '万元' in fulltext else 'unit_pending'
        column_labels={j+1:' | '.join(dict.fromkeys(raw(r[j]) for r in headers if j<len(r) and raw(r[j]))) for j in range(max(len(r) for r in rows))}
        mapping={}
        nested_budget_start=min([j+1 for r in headers for j,v in enumerate(r) if j>0 and '财政预算内教育经费支出' in norm(v)]+[999]) if funding=='all_source' else 999
        for col,label in column_labels.items():
            if col>=nested_budget_start:continue
            n=norm(label).replace('|','')
            if '个人部分' in n:mapping.setdefault('P',[]).append(col)
            if '公用部分' in n:mapping.setdefault('C',[]).append(col)
            if col>1 and not re.search('自筹|其中|#',n) and re.search('基建与专项支出|基本建设支出|基建支出',n):mapping.setdefault('K',[]).append(col)
        # Totals only when uniquely headed in an ordinary all-source expenditure table.
        total=[col for col,label in column_labels.items() if col>1 and norm(label) in ['总计','合计','教育经费支出|合计']]
        if funding=='all_source' and len(total)==1:mapping['E']=[total[0]]
        expense_tables.append(dict(source_id=sid,year=year,stage=stage,school_scope=school,funding_scope=funding,unit=unit,table_block=block,source_row_offset=offset,
                                   mappings=json.dumps(mapping),headers=json.dumps(headers,ensure_ascii=False),table=spec['table']))
        byprov={}
        for ri,r,p in rr:
            if p in CITIES:continue
            records={}
            for col in range(2,len(r)+1):
                if not raw(r[col-1]) and not column_labels.get(col):continue
                metric=next((m for m,cols in mapping.items() if len(cols)==1 and cols[0]==col),'finance_original_field')
                s=cell(sid,year,stage,p,metric,ri+offset,col,r[col-1],unit,school,funding,column_labels.get(col,''),
                       'raw column only; early/late classifications and continuation tables not joined')
                expense_records.append(s);records[col]=s
            byprov[p]=records
        # Source-local arithmetic only: E=P+C+remaining original construction field where explicitly mapped.
        for m,cols in mapping.items():
            if len(cols)>1:issue('finance_multiple_columns',sid,year,stage,'all',f'{m} maps to {cols}; no automatic choice')
        if stage=='secondary_combined_excluded':issue('secondary_scope_excluded',sid,year,stage,'all','ordinary secondary cannot serve as junior series')
print('finance tables',len(expense_tables),'cells',len(expense_records),flush=True)

# Reconcile only exact original components in the same workbook/sheet and province.
# This also verifies stacked continuation blocks without applying the first block's headers to the second.
srcmap={s['source_id']:s for s in sources}
finance_groups=defaultdict(list)
for s in expense_records:
    if s['metric'] in ['E','P','C','K']:
        src=srcmap[s['source_id']]
        finance_groups[(src['archive'],src['table'],src['sheet'],s['province'])].append(s)
for key,group in finance_groups.items():
    components={m:[s for s in group if s['metric']==m] for m in ['E','P','C','K']}
    if not all(len(v)==1 for v in components.values()):continue
    nums={m:parse(v[0]['raw_value'])[0] for m,v in components.items()}
    diff=nums['E']-nums['P']-nums['C']-nums['K'] if all(v is not None for v in nums.values()) else None
    state='pass_rounding' if diff is not None and abs(diff)<=Decimal(2) else 'unresolved_difference' if diff is not None else 'uncheckable_invalid_or_missing'
    ref=group[0]
    checks.append(dict(check='finance_E_P_C_K_original_sum',source_id=ref['source_id'],year=ref['data_year'],stage=ref['stage'],province=ref['province'],difference='' if diff is None else str(diff),tolerance=2,status=state))
    for s in group:
        s['arithmetic_status']=state
        if state=='pass_rounding':s['accepted_value']=s['numeric_candidate'];s['verification_basis']='same workbook original components reconcile within 2 original monetary units; cross-year scope still pending'
    if state!='pass_rounding':issue('finance_original_component_sum',ref['source_id'],ref['data_year'],ref['stage'],ref['province'],f'{state}; E-P-C-K={diff}; source-local only')

# RAR 2024 already has previously extracted, traced source cells; preserve their source inventory.
old_table_inv={r['file']+'|'+r['sheet']:r for r in readcsv(BASE/'data/yearbook_stage_table_inventory.csv')}
oldcells=readcsv(BASE/'data/yearbook_stage_spending_candidates.csv')
for r in oldcells:
    if r['volume_year']!='2024' or r['province'] not in PROVS+['全国']:continue
    inv=old_table_inv[r['file']+'|'+r['sheet']]; file=ROOT/r['file']; b=file.read_bytes()
    assert sha(b)==inv['sha256'],file
    sid=sha((r['file']+'|'+r['sheet']).encode())[:20]
    if sid not in source_seen:
        sources.append(dict(source_id=sid,family='education_finance',volume_year=2024,data_year=2023,data_year_status='prior_year_evidence',archive='2024 RAR; extracted file traced in prior receipt',table=r['file'],sheet=r['sheet'],table_sha256=sha(b),row_count='',scope=r['affiliation_scope'],source_quality='previously_extracted_user_excel'))
        source_seen.add(sid)
    col=openpyxl.utils.cell.coordinate_from_string(r['cell'])[0]
    cell(sid,2023,r['stage'],r['province'],{'total':'E','personal':'P','public':'C','construction':'K'}[r['metric']],int(r['row']),openpyxl.utils.column_index_from_string(col),r['value'],r['unit'],r['affiliation_scope'],r['amount_definition'],r['metric'],'2024 RAR values imported only after source hash recheck')

# E. Availability match is not a model-ready common sample.
index=defaultdict(list)
for s in ledger:
    if s['province'] in PROVS:index[(s['data_year'],s['province'],s['stage'],s['metric'])].append(s)
availability=[]
for y in range(1995,2025):
    for p in PROVS:
        for stage in ['primary','junior']:
            sr=index.get((y,p,stage,'S'),[])
            fr={m:[s for s in index.get((y,p,stage,m),[]) if s['funding_scope']=='all_source' and s['school_scope']=='all_affiliations_candidate'] for m in ['E','P','C']}
            # Include 2023 extracted all-source/all-affiliation records with the same semantic scope.
            for m in fr:fr[m]+=[s for s in index.get((y,p,stage,m),[]) if s['school_scope']=='all_affiliations' and s['funding_scope']=='all-source education spending']
            pr={m:index.get((y,p,'province_only',m),[]) for m in ['R','CPI_yoy']}
            local={m:[s for s in index.get((y,p,stage,m),[]) if s['funding_scope']=='all_source' and s['school_scope']=='local_affiliation'] for m in ['E','P','C']}
            availability.append(dict(year=y,province=p,stage=stage,geography_status='not_applicable' if p=='重庆' and y<1997 else 'boundary_pending' if p=='四川' and y<=1997 else 'current_province',student_records=len(sr),student_numeric_accepted=sum(bool(s['accepted_value']) for s in sr),E_candidates=len(fr['E']),P_candidates=len(fr['P']),C_candidates=len(fr['C']),local_E_candidates=len(local['E']),local_P_candidates=len(local['P']),local_C_candidates=len(local['C']),R_candidates=len(pr['R']),CPI_candidates=len(pr['CPI_yoy']),candidate_fields_present=all(sr and fr[m] for m in fr) and all(pr[m] for m in pr),model_eligible=False,reason='school/time/funding/price/geographic bridges not yet jointly approved'))
coverage=[]
for y in range(1995,2025):
    for stage in ['primary','junior']:
        group=[a for a in availability if a['year']==y and a['stage']==stage and a['geography_status']!='not_applicable']
        coverage.append(dict(year=y,stage=stage,expected_provinces=len(group),student_rows=sum(a['student_records']>0 for a in group),student_internal_numeric_accepted=sum(a['student_numeric_accepted']>0 for a in group),E_provinces=sum(a['E_candidates']>0 for a in group),P_provinces=sum(a['P_candidates']>0 for a in group),C_provinces=sum(a['C_candidates']>0 for a in group),local_E_provinces=sum(a['local_E_candidates']>0 for a in group),local_P_provinces=sum(a['local_P_candidates']>0 for a in group),local_C_provinces=sum(a['local_C_candidates']>0 for a in group),R_provinces=sum(a['R_candidates']>0 for a in group),CPI_provinces=sum(a['CPI_candidates']>0 for a in group),candidate_fields_present=sum(a['candidate_fields_present'] for a in group),model_ready_provinces=0))
write('source_registry.csv',sources)
write('raw_value_ledger.csv',ledger)
write('student_column_dictionary.csv',mappings)
write('student_candidate_series.csv',student_rows)
write('fiscal_original_candidates.csv',fiscal_records)
write('supplement_original_candidates.csv',panel_rows)
write('education_finance_table_dictionary.csv',expense_tables)
write('education_finance_original_candidates.csv',expense_records)
write('arithmetic_checks.csv',checks)
write('exception_ledger.csv',issues)
write('matching_availability.csv',availability)
write('coverage_by_year_stage.csv',coverage)
definitions = [
    ('S','普通小学/普通初中在校生','人','province × stage × actual_year','same school type, organizer, attached classes and census time','source-internal numeric acceptance is separate from cross-source comparability'),
    ('P','原分类个人部分支出','original monetary unit','province × stage × year × funding × affiliation','original component definition and classification-era bridge','not pure teacher wages; positive values only for logs'),
    ('C','原分类公用部分支出','original monetary unit','province × stage × year × funding × affiliation','capital and service items may be nested; inspect original hierarchy','not pure routine teaching cost; cannot infer quality'),
    ('E','公布支出总额','original monetary unit','province × stage × year × funding × affiliation','same population as S; preserve source total and component discrepancy','do not reconstruct total from incompatible components'),
    ('K','建设及历史相关分类','original monetary unit','province × stage × year × original definition','early construction and special spending differs from later construction','no forced long-run bridge; no log of zero'),
    ('R','全省地方自身财政收入','亿元 or verified original unit','province × actual_year','all subordinate cities/counties; historical budget classification; geographic scope','not province-government level alone; not resources after transfers'),
    ('central_subsidy','中央补助原项目','verified source unit','province × year × table block','tax return/general/special transfers containment and city exclusions','do not add containing and contained items'),
    ('remittance_central','上解中央支出原项目','verified source unit','province × year × table block','full-province vs government level; unit and final column','cannot subtract from unverified resource total'),
    ('CPI_yoy','综合居民消费价格指数','上年=100','province × actual_year','aggregate index only; original layer; valid positive observations','purchasing-power approximation; chain not built until gaps resolved'),
    ('GDP','地区生产总值','亿元','province × actual_year','version/revisions, geography, lagged control choice','not automatically a real GDP growth series'),
    ('D','Δln(C)−Δln(P)','log points','province × stage × actual_year','same positive C/P, school/funding scope and adjacent years','not yet constructed; not indicator of proven crowding-out'),
    ('Z','max(0,−Δln(real R))','log points','province × actual_year','verified R and continuous CPI; comparable adjacent years','not yet constructed; own-revenue contraction only'),
    ('E_per_S','同范围实际生均支出','verified real unit/person','province × stage × year','approved E/S school scope and statistical timing','not yet constructed; never recover S by dividing E by published per-student figure'),
]
write('variable_dictionary.csv',[dict(variable=v,definition=d,unit=u,grain=g,admission_rule=a,interpretation_boundary=n) for v,d,u,g,a,n in definitions])
write('status_dictionary.csv',[
    dict(field='value_status',code=c,meaning=m) for c,m in [
        ('numeric_observed','explicit source numeric notation; not itself accepted'),('zero_observed','source explicitly zero; never treat as missing or add constant'),
        ('missing_blank','source blank; do not replace by zero'),('missing_symbol_unresolved','source symbol; missing/not-applicable meaning pending'),
        ('invalid_numeric_text','letters/internal spaces/other invalid notation retained; no automatic repair'),
        ('not_applicable_geography','independent administrative unit did not exist'),
        ('geography_not_applicable_observation_present','compiled numeric observation present before unit existed; exclude'),
        ('conflicting_table_blocks_pending','inconsistent region titles; no block selected')]] + [
    dict(field='bridge_status',code='unverified',meaning='cross-year and cross-source scope not approved; model_eligible=false'),
    dict(field='bridge_status',code='numeric_internal_checks_pass; school_scope_pending',meaning='total and grades reconcile internally, but student/finance scope and timing remain pending'),
    dict(field='accepted_value',code='nonempty',meaning='only documented numeric acceptance; does not imply model admission'),
    dict(field='model_eligible',code='False',meaning='this milestone does not certify any jointly comparable model sample')])
for z,idx in archives.values():z.close()
protected_after={str(p.relative_to(ROOT)):sha(p.read_bytes()) for p in PROTECTED}
assert protected_after==protected_before,'protected source outputs changed'
assert len({s['record_id'] for s in ledger})==len(ledger),'duplicate source cell metric key'
assert len(student_rows)==1916,'60 student tables: 1856 province and 60 national rows'
summary=dict(ledger_rows=len(ledger),source_tables=len(sources),student_rows=len(student_rows),student_province_rows=sum(s['province']!='全国' for s in student_rows),student_numeric_internal_accepted=sum(bool(s['accepted_value']) and s['province']!='全国' for s in student_rows),student_check_status=dict(Counter(s['status'] for s in checks if s['check'].startswith('student'))),finance_check_status=dict(Counter(s['status'] for s in checks if s['check'].startswith('finance'))),panel_metric_counts=dict(Counter(s['metric'] for s in panel_rows)),fiscal_cells=len(fiscal_records),education_finance_cells=len(expense_records),issue_counts=dict(Counter(i['kind'] for i in issues)),matching_rows=len(availability),model_ready_rows=0,protected_files_unchanged=True,protected_hashes=protected_after,interpolation_used=False,regression_run=False,scope='first construction milestone; candidate values are not final comparable panel')
(OUT/'build_receipt.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(summary,ensure_ascii=False,indent=2))
