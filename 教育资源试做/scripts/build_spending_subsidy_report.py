"""Illustrated research supplement, retaining exploratory qualifications."""
from pathlib import Path
import json
import pandas as pd
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[2];DATA=ROOT/'教育资源试做/data/支出细项与补助拓展_20261011';FIG=ROOT/'教育资源试做/支出细项与补助图组_20261011';OUT=ROOT/'支出细项与补助拓展报告_20261011';OUT.mkdir(exist_ok=True)
t=pd.read_csv(DATA/'component_primary_tests.csv');c=pd.read_csv(DATA/'all_coefficients.csv');sub=pd.read_csv(DATA/'subsidy_exploratory_tests.csv');checks=json.loads((DATA/'validation_receipt.json').read_text(encoding='utf8'))
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.8);sec.header_distance=sec.footer_distance=Inches(.3)
for name,size in [('Normal',11),('Title',18),('Subtitle',10),('Heading 1',13),('Caption',9)]:
    st=doc.styles[name];st.font.name='黑体' if name in ['Title','Heading 1'] else '宋体';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.font.italic=False;st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),st.font.name);st.paragraph_format.line_spacing=1.18;st.paragraph_format.space_after=Pt(7)
for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
hp=sec.header.paragraphs[0];hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT;hp.add_run('义务教育经费长期研究  支出细项与中央补助拓展').font.size=Pt(8)
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER;field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');fp._p.append(field)
textlog=[]
def para(text,small=False):
    textlog.append(text);p=doc.add_paragraph(text);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if small:
        for r in p.runs:r.font.size=Pt(9)
    return p
def heading(text):textlog.append('\n## '+text);doc.add_heading(text,1)
def figure(name,caption,width=6.6):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(FIG/f'{name}.png'),width=Inches(width));doc.add_paragraph(caption,'Caption');textlog.append(caption)
def page(title):doc.add_page_break();heading(title)
def row(stage,reg):return t[(t.stage==stage)&(t.regressor==reg)].iloc[0]
def table(headers,rows):
    tb=doc.add_table(rows=1,cols=len(headers));tb.autofit=False
    for cell,label in zip(tb.rows[0].cells,headers):
        cell.text=label;sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EAEAEA');cell._tc.get_or_add_tcPr().append(sh)
    for values in rows:
        for cell,value in zip(tb.add_row().cells,values):cell.text=str(value)
    for rr in tb.rows:
        for cell in rr.cells:
            cell.width=Inches(6.6/len(headers))
            for p in cell.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
                for r in p.runs:r.font.size=Pt(9)
    textlog.append('| '+' | '.join(headers)+' |\n'+'\n'.join('| '+' | '.join(map(str,r))+' |' for r in rows))
doc.add_paragraph('义务教育支出细项与中央补助拓展报告','Title');doc.add_paragraph('三十年研究的解释性补充  2008—2023年省级变化  2026年10月11日','Subtitle')
para('本轮进一步区分日常运行、人员相关项目与资本性项目，并检验加入中央补助后，生源与支出调整的关系是否改变。主要发现是：小学与初中的商品服务支出对学生规模变化的响应均高于工资福利支出。这一关联并非只来自公用支出中的资本性项目；中央补助的条件关联则未获得明确统计支持。')
heading('样本与方法')
para('整体研究仍以1995—2024年为历史主线，1994年另列条件估算。本轮依据2007—2023年31个省级地区细项金额计算2008—2023年变化；不跨2006—2007年分类边界，不接入2024年待核的资本性分类。两学段各472条完整省年，不插值、不补零、不缩尾。金额为全隶属全来源支出，不等同财政性支出。')
para('对数变化以省级同比CPI折算实际值，控制学生变化、实际自身收入及省份、年份效应，各省年等权。直接估计“商品服务对数变化减工资福利对数变化”。两学段的学生与收入差系数共4项，采用9999次约束野生省级聚类自助、检验反演单项95%区间及Holm校正。方案先于本轮拟合记录，但研究为自适应探索，不称为预注册确认性检验。')
figure('图F3_商品服务与工资响应差','图1  两学段商品服务与工资福利的直接响应差。线为单项检验反演95%区间，非四项同时置信区间；图中p值为本轮四项Holm校正。每学段472条省年、31个聚类。',6.5)
for s,label in [('primary','小学'),('junior','初中')]:
    r=row(s,'dS');para(f'{label}学生响应差为{r.coefficient:.3f}，95%区间为{r.WCR_95_CI_low:.3f}至{r.WCR_95_CI_high:.3f}，校正p＝{r.Holm_family_p:.4f}。学生变化提高1个对数百分点时，商品服务相对工资福利的变化约高出{r.coefficient:.3f}个对数百分点。这是条件关联，不能解释为学生数量的因果效应。',True)
page('区分人员相关项目与公用内部项目')
figure('图F1_支出内部构成','图2  2007—2023年31个省级地区名义金额合计的内部构成。左列分母为个人支出，右列为公用支出；该汇总不冒充官方全国序列。其他资本性属于公用支出，与另列基本建设支出区分。')
para('个人支出分为工资福利及对个人家庭补助；公用支出分为商品服务及其他资本性项目。2023年工资福利占个人支出约89.7%（小学）、90.3%（初中），商品服务占公用支出约58.5%、53.4%。因此，用“个人＝教师工资”“公用＝日常运行”替代原统计分类，会损失重要结构信息。')
para('图中2017—2018年前后工资福利份额明显变化，应与养老保险、人员支出列支方式及年鉴说明一并核准。账面份额不能直接等同教师人均工资或待遇变化。看图后追加剔除2017—2018年的检查，学生响应差仍为0.816（小学）、0.450（初中），但这并不能证明所有跨年分类问题已经消除。')
para('回读原始工作簿的4216个细项单元格，保留坐标、项目标签、原表哈希及千元单位；1054条省年均满足个人和公用的细项加总，在原表显示舍入容差内成立。本轮没有把范围待核记录升级为完整口径认证。',True)
page('把公用增长拆成两项金额贡献')
figure('图F2_公用增长贡献','图3  商品服务与其他资本性对公用实际增长的贡献。按同一省份前期公用金额作分母，计算两项实际金额变化贡献，再取当年模型样本省份的算术均值。两柱之和与黑线在原表舍入容差内一致；不是全国合计增长率。')
para('年度增长速度的变化，不只涉及日常运行项目。图示中2009—2012年两学段的其他资本性平均贡献普遍高于商品服务；部分后期年份出现商品服务为正而资本性为负的组合。只看公用总额会把这类内部抵消隐藏起来。')
para('这里的贡献是“实际分项金额变化／前期公用金额”，单位为百分点。它与回归中的对数变化不同；不以总增量作分母计算所谓贡献率，避免总增量接近零时的比例失真。年度均值使用各学段472条模型样本中属于当年的对应记录，部分年度省份数量因原财政缺口不同，应理解为样本描述。')
para('资本性项目可能涉及设备、信息化和其他资产购置，但该统计项目不足以区分每一用途。本轮只描述账面贡献；不将资本性增长解释为投资效率提升，也不将负贡献等同公共服务退步。',True)
page('支出细项响应与稳健性')
figure('图F4_支出细项响应','图4  四个细项对学生变化的分模型系数。控制实际自身收入和省份、年份效应；各学段均472条省年。线为省级CR1及t(30)近似95%区间，供辅助描述；直接差异以图1的检验为准。')
para('商品服务的学生响应系数为0.589（小学）、0.444（初中），工资福利为−0.177、−0.062。商品服务与工资福利之间的差异得到直接检验支持；工资福利系数本身的区间包含零，不能据此宣布生源减少会推高工资支出。商品服务也不是纯教学运行成本，工资福利不是纯专任教师工资。')
para('其他资本性的学生响应为0.446、0.901，个人家庭补助为0.061、0.339。这些属于辅助细项结果，未再设额外主假设；不能根据单个区间或两学段各自的显著性，宣称学段之间存在显著差异。')
sens=pd.read_csv(DATA/'sensitivity.csv');rows=[]
for label,spec in [('剔除2013—2014年','without_2013_14'),('剔除2015—2016年','without_2015_16'),('剔除2020—2022年','without_2020_22')]:rows.append([label,*[f'{sens[(sens.stage==s)&(sens.specification==spec)&(sens.outcome=="DGW")&(sens.regressor=="dS")].coefficient.iloc[0]:.3f}' for s in ['primary','junior']]])
table(['敏感性处理','小学学生响应差','初中学生响应差'],rows)
para('表1  预先列明的年份剔除敏感性。数值为商品服务减工资福利的学生响应差，控制条件与全窗口主模型相同；各行使用其自身的完整样本。',True)
para('逐省删除的学生响应差范围为0.707—0.821（小学）、0.457—0.562（初中），未因单一省份改变方向。敏感性表为点估计辅助检查，不替代主自助推断。结构上支持“项目调整响应不同”的解释，但尚不足以识别工资制度、编制或组织刚性中的具体因果路径。',True)
page('全窗口关系不等于各时期稳定规律')
figure('图F7_时期分段响应','图5  看图后追加的时期检查。全窗口、2008—2016年前期、2019—2023年后期分别估计，控制学生、自身实际收入、省份和年份效应。线统一为CR1及t(30)近似95%区间，不与图1的自助区间混称；分段样本不同，系数变化不是正式结构突变检验。')
para('前期学生响应差为1.011（小学）、0.502（初中）；后期为1.128、−0.207。初中后期近似95%区间为−1.339至0.925，精度明显不足。全窗口关联获得支持，并不等于后期也已证实相同方向；同样不能根据后期负向点估计宣布关系逆转。')
para('这项检查是在构成图揭示2017—2018年前后变化后追加的，并非拟合前方案中的确认性假设。它揭示应进一步核查的时期问题，没有替换原来的主检验。剔除过渡年份后的全窗口正向关系，亦不足以排除不同历史阶段的制度和项目差异。')
para('因此，报告中的结构性结论应限定为：在2008—2023年可连接样本中，商品服务与工资福利存在不同的生源响应；初中后期是否延续这一关系仍不确定。对三十年研究的贡献在于区分支出项目、财政来源和历史阶段，而非提出无时段边界的一般规律。')
page('自身收入下降与补助合计变化')
figure('图F5_自身收入与补助合计变化','图6  两学段共同的436个省年仅计一次，覆盖2008—2023年31个省级地区。金色为自身收入实际下降而自身收入与中央补助合计未下降的记录。坐标为实际对数变化乘100，非普通百分比增长。',6.5)
para('财政年鉴决算表提供495条年度中央补助记录，均按表内亿元或万元单位换算，且自身收入与既有核准值相符。与相邻年及原模型连接后，每学段436条省年，较细项模型少36条；全部排除另存。不把补助空白填为零，不将返还、补助分项与总项重复相加。')
para('共同样本中，自身收入实际下降的84条省年有60条（71.4%）合计指标未下降。这说明仅以自身收入下降代表财政资源同步下降，可能误判部分省年的预算背景。不过，中央补助含返还及专门用途，补助也会对地方困难和需求作出响应；上述并存状态不能解释为外生转移支付的因果缓冲效果。')
para('“自身收入与中央补助合计”不等于完整可用财力。它没有完整扣除上解等净结算，也没有纳入全部其他合法财源；债务、结转、调入及收入总计不混入中央补助。计划单列市和省市边界仍待充分桥接，相关省份剔除敏感性另存。',True)
page('加入补助后的模型比较与证据边界')
figure('图F6_中央补助关联','图7  同时控制学生、自身实际收入、省份及年份效应的中央补助系数。每学段436条省年，9999次约束野生聚类自助；线为单项检验反演95%区间。4项补助探索检验构成单独Holm家族，不是全项目确认性校正。',6.5)
para('中央补助对商品服务的系数为0.010（小学）、0.038（初中），对商品服务减工资福利的差系数为−0.072、−0.053；4项校正p均为1.000，区间均跨零。当前数据不能确认补助增长与这两类支出调整存在独立关联，也不能将证据不足写成补助没有作用。')
rows=[]
for s,label in [('primary','小学'),('junior','初中')]:
    vals=[c[(c.stage==s)&(c.specification==spec)&(c.outcome=='DGW')&(c.regressor=='dS')].iloc[0] for spec in ['same_sample_own','own_and_subsidy','own_plus_subsidy']];rows.append([label,*[f'{r.coefficient:.3f}\n({r.cluster_CR1_se:.3f})' for r in vals]])
table(['学生响应差','控制自身收入','另加中央补助','改用两项合计'],rows)
para('表2  三列均使用各学段同一436条省年，括号为省级CR1标准误。以同样本比较避免把样本变化归因于财力变量。学生响应差在补助进入后仍保持正向，但表中变化不构成中介解释比例或因果机制检验。',True)
para('本轮收获是把宽泛的个人—公用关系收窄为人员相关与商品服务项目的调整差异，同时验证自身收入并不足以概括补助后的财政背景。下一步应围绕工资列支、人员制度及教育专项补助寻找独立证据，而非继续增加回归直到出现显著结果。三十年描述与这一较短时期的条件关联应共同构成报告，不能把本轮结果外推到全部三十年。')
para(f'来源为用户提供的中国教育经费统计年鉴和中国财政年鉴原表。拟合前方案、原表回读、模型输入、所有系数、排除、敏感性及图表数据另存；{checks["checks"]}项独立检查通过。图片为白底600 dpi PNG，并附SVG及PDF。既有报告和保护数据未覆盖。',True)
path=OUT/'义务教育支出细项与中央补助拓展报告.docx';doc.save(path);(OUT/'报告正文.md').write_text('# 义务教育支出细项与中央补助拓展报告\n\n'+'\n\n'.join(textlog),encoding='utf8')
assert len(Document(path).inline_shapes)==7
(OUT/'报告制作记录.json').write_text(json.dumps(dict(figures=7,tables=2,validation_checks=checks['checks'],old_reports_preserved=True),ensure_ascii=False,indent=2),encoding='utf8');print(path)
