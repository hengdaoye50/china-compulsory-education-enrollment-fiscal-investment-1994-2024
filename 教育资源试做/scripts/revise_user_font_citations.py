"""Edit the user's font-adjusted document in place structurally, saving a new copy."""
from pathlib import Path
from copy import deepcopy
import json,re,hashlib,zipfile
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
ROOT=Path(__file__).resolve().parents[2]
SRC=Path('<LOCAL_USER>/Desktop/中国义务教育生源规模与财政投入趋势分析报告（1994—2024）.docx')
OUT=ROOT/'义务教育趋势分析报告_引文规范修订版_20261012';OUT.mkdir(exist_ok=True)
NAME=SRC.stem;doc=Document(SRC)
original_styles=doc.styles.element.xml
refs={
'5':'梁文艳, 孙雨婷, 杜育红. 教育强国背景下适应人口变动的义务教育师资配置——基于人口增减分化典型省域的研究[J]. 教育研究, 2026(2): 40-53.',
'6':'蔺海沣, 马玲. 人口结构变动对公共教育支出的影响效应——基于全球70项定量研究的三水平元分析[J]. 教育研究, 2026(2): 54-68.',
'7':'李涛, 黄嘉欣. 学龄人口减少对乡村小规模学校发展的影响及其疏解——基于东北地区人口负增长省份的调查[J]. 教育研究, 2025(5): 107-121.',
'8':'陈坤, 秦玉友. 人口变动趋势下基础教育师资配置机制研究——基于8省16县(区)的调查[J]. 西南民族大学学报(人文社科版), 2026, 47(5): 229-240.',
'1a':'国家教育委员会计划建设司. 中国教育事业统计年鉴: 1995-1997[M]. 北京: 人民教育出版社, 1996-1998.',
'1b':'教育部发展规划司. 中国教育统计年鉴: 1998-2014[M]. 北京: 人民教育出版社, 1999-2015.',
'1c':'教育部发展规划司. 中国教育统计年鉴: 2015-2024[M]. 北京: 中国统计出版社, 2016-2025.',
'2':'国家教育委员会财务司, 教育部财务司, 国家统计局. 中国教育经费统计年鉴: 1996-2025[M]. 北京: 中国统计出版社, 1996-2025.',
'3':'中国财政年鉴编辑委员会. 中国财政年鉴: 1996-2024[M]. 北京: 中国财政杂志社, 1996-2024.',
'4':'国家统计局. 中国统计年鉴: 1995-2025[M]. 北京: 中国统计出版社, 1995-2025.',
'16':'国务院. 国务院关于进一步完善城乡义务教育经费保障机制的通知: 国发〔2015〕67号[EB/OL]. (2015-11-25)[2026-10-12]. https://www.moe.gov.cn/jyb_xxgk/moe_1777/moe_1778/201511/t20151130_221655.html.',
'17':'国务院办公厅. 国务院办公厅关于进一步调整优化结构提高教育经费使用效益的意见: 国办发〔2018〕82号[EB/OL]. (2018-08-17)[2026-10-12]. https://www.moe.gov.cn/jyb_xxgk/moe_1777/moe_1778/201808/t20180827_346266.html.',
'18':'国务院办公厅. 国务院办公厅关于印发教育领域中央与地方财政事权和支出责任划分改革方案的通知: 国办发〔2019〕27号[EB/OL]. (2019-05-24)[2026-10-12]. https://www.moe.gov.cn/jyb_xxgk/moe_1777/moe_1778/201906/t20190604_384344.html.',
'15':'阿里云. DataV.GeoAtlas: 省级行政区边界数据[DS/OL]. [2026-10-12]. https://geo.datav.aliyun.com/areas_v3/bound/100000_full.json.',
'9':'CAMERON A C, GELBACH J B, MILLER D L. Bootstrap-based improvements for inference with clustered errors[J]. The Review of Economics and Statistics, 2008, 90(3): 414-427. DOI: 10.1162/rest.90.3.414.',
'12':'蒋鸣和. 中国教育经费投入的“八五”回顾与“九五”展望[J/OL]. 教育研究信息, 1996(9)[2026-10-12]. https://www.cnsaes.org.cn/homepage/Upfile/201029/2010020981343441.pdf.',
'13':'国务院新闻办公室. 中国的儿童状况[R/OL]. 北京: 国务院新闻办公室, 1996[2026-10-12]. https://www.fmprc.gov.cn/ziliao_674904/zt_674979/ywzt_675099/wzzt_675579/2296_675789/200207/t20020705_7961975.shtml.',
'14':'教育部, 国家统计局, 财政部. 关于1998年全国教育经费执行情况统计公告: 教财〔1999〕29号[EB/OL]. (1999-11-02)[2026-10-12]. https://hudong.moe.gov.cn/jyb_sjzl/sjzl_jfzxgg/201005/t20100527_88452.html.'
}
pat=re.compile(r'\[(\d+(?:\s*[-—–,，]\s*\d+)*)\]')
def ids(t):
    result=[]
    for q in re.split(r'[,，]',t):
        z=re.split(r'[-—–]',q)
        result.extend(range(int(z[0]),int(z[-1])+1))
    out=[]
    for i in result:
        keys=['1a','1b','1c'] if i==1 else ['4'] if i in [4,10,11] else [str(i)]
        for key in keys:
            if key not in out:out.append(key)
    return out
def fmt(nums):
    nums=sorted(set(nums));spans=[];i=0
    while i<len(nums):
        j=i
        while j+1<len(nums) and nums[j+1]==nums[j]+1:j+=1
        spans.append(str(nums[i])+'-'+str(nums[j]) if j-i>=2 else ','.join(str(x) for x in nums[i:j+1]));i=j+1
    return '['+','.join(spans)+']'
def put_text(p,text):
    properties=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    for r in list(p.runs):p._p.remove(r._r)
    r=p.add_run(text)
    if properties is not None:r._r.insert(0,properties)
def append_ref(p,text):
    properties=deepcopy(p.runs[-1]._r.rPr) if p.runs and p.runs[-1]._r.rPr is not None else None
    r=p.add_run(text)
    if properties is not None:r._r.insert(0,properties)
summary='全国小学在校生在1997年达到峰值，持续回落至2013年后回升，2024年再次下降；初中在2003年达到峰值、2015年降至低点后回升，两学段的需求变化呈现时序差异。2013—2024年，两学段财政性收入名义金额分别增长107.31%和111.27%，收入占比则回落，举办者收入权重变化参与了全国比例调整。地方支出实际增速在后期放缓，部分地区公用支出累计下降，公用内部商品服务与资本性项目的变化又可能相互抵消。省级模型进一步显示，商品服务增长与生源变化的关联强于工资福利，教师、班级和中央补助拓展尚未建立独立机制证据。长期投入扩张由此表现为学段时序、地区增速和支出项目的差异化变化。'
target=next(p for p in doc.paragraphs if p.text.startswith('三十年资料表明'))
put_text(target,summary)
heading_index=next(i for i,p in enumerate(doc.paragraphs) if p.text=='资料来源与参考文献')
body=doc.paragraphs[:heading_index]
# Explicitly cite the map source rather than listing an uncited database.
map_para=next(p for p in body if '边界来源：DataV.GeoAtlas' in p.text)
append_ref(map_para,'[15]')
order=[]
for p in body:
    for match in pat.finditer(p.text):
        for key in ids(match.group(1)):
            if key not in order:order.append(key)
assert set(order)==set(refs),(order,set(refs)-set(order))
number={key:i+1 for i,key in enumerate(order)}
changed=[];superscripts=0
for p in body:
    text=p.text;matches=list(pat.finditer(text))
    if not matches:continue
    runs=[];pos=0
    for r in p.runs:
        runs.append((pos,pos+len(r.text),r));pos+=len(r.text)
    assert pos==len(text),'Unexpected non-run text in citation paragraph'
    def property_at(at):
        return next((deepcopy(r._r.rPr) for a,b,r in runs if a<=at<b and r._r.rPr is not None),None)
    segments=[];at=0
    for match in matches:
        for a,b,r in runs:
            lo=max(at,a);hi=min(match.start(),b)
            if hi>lo:segments.append((text[lo:hi],deepcopy(r._r.rPr),False))
        segments.append((fmt(number[k] for k in ids(match.group(1))),property_at(match.start()),True));at=match.end()
    for a,b,r in runs:
        lo=max(at,a)
        if b>lo:segments.append((text[lo:b],deepcopy(r._r.rPr),False))
    for r in list(p.runs):p._p.remove(r._r)
    for t,pr,super_ in segments:
        r=p.add_run(t)
        if pr is not None:r._r.insert(0,pr)
        if super_:r.font.superscript=True;r.font.size=Pt(10.5);superscripts+=1
    changed.append(p.text)
# Replace the bibliography with unique entries ordered by first citation.
bibliography=doc.paragraphs[heading_index+1:]
assert len(bibliography)==18
for p,key in zip(bibliography,order):
    put_text(p,f'[{number[key]}] '+refs[key]);p.paragraph_format.left_indent=Pt(18);p.paragraph_format.first_line_indent=Pt(-18);p.paragraph_format.space_after=Pt(5)
    p.paragraph_format.keep_together=True
    for run in p.runs:run.font.superscript=False
assert doc.styles.element.xml==original_styles,'User styles changed'
dest=OUT/(NAME+'.docx');doc.save(dest)
# Preserve every original package part except the edited document body.
with zipfile.ZipFile(dest) as edited:updated_body=edited.read('word/document.xml')
temporary=OUT/'package_preserved.docx'
with zipfile.ZipFile(SRC) as original,zipfile.ZipFile(temporary,'w',compression=zipfile.ZIP_DEFLATED) as preserved:
    for item in original.infolist():preserved.writestr(item,updated_body if item.filename=='word/document.xml' else original.read(item.filename))
temporary.replace(dest)
with zipfile.ZipFile(SRC) as a,zipfile.ZipFile(dest) as b:
    names=[n for n in a.namelist() if n.startswith('word/media/') and not n.endswith('/')]
    assert all(a.read(n)==b.read(n) for n in names)
    assert a.read('word/styles.xml')==b.read('word/styles.xml')
receipt={'source':str(SRC),'source_sha256':hashlib.sha256(SRC.read_bytes()).hexdigest(),'output':str(dest),'styles_preserved':True,'images_preserved':True,'references':len(order),'reference_number_map':number,'superscript_markers':superscripts,'GB_standard':'GB/T 7714-2025','summary':summary,'merged_old_references':{'4':[4,10,11]},'education_yearbook_split':['1995-1997','1998-2014','2015-2024'],'article_1996_pagination':'Original issue pagination not obtained; no page range invented.'}
(OUT/'引文与字体修订记录.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf8')
(OUT/'参考文献核准与编号.md').write_text('\n\n'.join(f'[{number[k]}] '+refs[k] for k in order),encoding='utf8')
print(json.dumps(receipt,ensure_ascii=False))
