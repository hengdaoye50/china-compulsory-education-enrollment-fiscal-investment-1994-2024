from pathlib import Path
import hashlib, json, re, shutil, zipfile
from docx import Document
from pypdf import PdfReader

root = Path.cwd()
out = root / '义务教育趋势分析报告_引文规范修订版_20261012'
name = '中国义务教育生源规模与财政投入趋势分析报告（1994—2024）'
source = Path('<LOCAL_USER>/Desktop') / (name + '.docx')
docx = out / (name + '.docx')
pdf_source = out / '终稿核验' / (name + '.pdf')
a, b = Document(source), Document(docx)
with zipfile.ZipFile(source) as za, zipfile.ZipFile(docx) as zb:
    assert za.namelist() == zb.namelist()
    changed_parts = [n for n in za.namelist() if za.read(n) != zb.read(n)]
    assert changed_parts == ['word/document.xml'], changed_parts
assert len(a.tables) == len(b.tables) == 15
assert len(a.inline_shapes) == len(b.inline_shapes) == 18
assert len(a.paragraphs) == len(b.paragraphs) == 220
assert [[c.text for row in t.rows for c in row.cells] for t in a.tables] == [[c.text for row in t.rows for c in row.cells] for t in b.tables]
head = next(i for i,p in enumerate(b.paragraphs) if p.text == '资料来源与参考文献')
refs = b.paragraphs[head+1:]
assert [int(re.match(r'\[(\d+)\]', p.text)[1]) for p in refs] == list(range(1,19))
cited = set()
markers = 0
for p in b.paragraphs[:head]:
    for run in p.runs:
        matches = list(re.finditer(r'\[([\d,\-]+)\]',run.text))
        if matches:
            assert run.font.superscript is True, run.text
            markers += len(matches)
            for match in matches:
                for part in match[1].split(','):
                    ends = [int(x) for x in part.split('-')]
                    cited.update(range(ends[0], ends[-1]+1))
assert cited == set(range(1,19)), cited
pages = PdfReader(pdf_source).pages
assert len(pages) == 31
starts = {4:'分析框架与资料范围',5:'生源规模的长期变化',6:'财政性经费与收入构成',8:'支出扩张的区域差异',11:'支出项目的结构调整',14:'生源变化与支出响应',17:'办学资源与财政补助',18:'主要结论与政策启示',20:'指标定义与统计范围',23:'估计方法与稳健性分析',27:'早期经费与跨期衔接',29:'分省结果与资料管理',31:'资料来源与参考文献'}
for page,title in starts.items():
    assert title in re.sub(r'\s+', '', pages[page-1].extract_text()), (page,title)
pdf = out / (name + '.pdf')
shutil.copy2(pdf_source,pdf)
digest=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert digest(source)=='2e8d7ae743467f7d011a675d39d682e46120ab85222556e6769c4b9a02edeb34'
receipt={'source_sha256':digest(source),'docx_sha256':digest(docx),'pdf_sha256':digest(pdf),'changed_package_parts':changed_parts,'original_styles_fonts_headers_footers_images_preserved':True,'tables_unchanged':15,'figures_unchanged':18,'unique_references':18,'superscript_markers':markers,'all_references_cited':True,'pdf_pages':31,'toc_page_starts_verified':starts,'visual_review_pages':list(range(1,32)),'remaining_metadata_limitation':'1996年文献原刊页码未取得，未补造。'}
(out/'最终交付核验.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(receipt,ensure_ascii=False,indent=2))
