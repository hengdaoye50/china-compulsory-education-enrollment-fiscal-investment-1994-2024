"""Round 33: reviewed census snapshots; no projection, figures, or enrollment division."""
from pathlib import Path
import csv, hashlib, json, re
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
DATA=ROOT/'data'
def read(name):
 return list(csv.DictReader((DATA/name).open(encoding='utf-8-sig')))
def save(name,rows):
 with (DATA/name).open('w',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()

# Fujian data are retained verbatim in the combined evidence table.
rows=read('population_province2020_single_age.csv')
assert len(rows)==20
checks=[]
for r in rows:
 for k in ['year','age_completed','total','male','female']:r[k]=int(r[k])
 r['note_status']='yearbook editor note reviewed'
 r['military_exclusion_status']='explicit exclusion in editor note'
 r['source_format']='official image; direct reviewed transcription'
 r['source_locator']='3-1, original image, age row'

jl_url='https://tjj.jl.gov.cn/tjsj/qwfb/jlsdqcqgrkpcnj/zk/html/A0301.jpg'
jl=[(116439,60272,56167),(146716,75519,71197),(147720,76065,71655),
 (174101,89793,84308),(184163,94893,89270),(154854,79871,74983),
 (205365,106594,98771),(193471,100150,93321),(214322,111442,102880),
 (192275,100351,91924),(203903,105724,98179),(211155,110109,101046),
 (220704,114870,105834),(229218,119775,109443),(224317,116868,107449),
 (227487,118527,108960),(220445,115747,104698),(179181,92699,86482),
 (210661,108542,102119),(219137,110692,108445)]
def make(province,age,counts,url,path,format,locator,note,military,excludes):
 return dict(province=province,year=2020,reference_time='2020-11-01 00:00',age_completed=age,
 total=counts[0],male=counts[1],female=counts[2],unit='人',count_type='census resident population; full counts',
 table='3-1',url=url,source_path=path,excludes=excludes,annual_panel_eligible=False,
 school_enrollment_denominator_certified=False,note_status=note,military_exclusion_status=military,
 source_format=format,source_locator=locator)
for age,counts in enumerate(jl):
 rows.append(make('吉林',age,counts,jl_url,'raw/population_jilin2020_age.jpg',
 'official image; direct reviewed transcription','3-1, original image, age row',
 'yearbook editor note reviewed','explicit exclusion in editor note',
 'active servicemen; short-table households exclude HK/Macao/Taiwan residents and foreigners; no undercount correction'))

yn_url='https://stats.yn.gov.cn/zwgk/zfxxgk/fdzdgknr/tjsj/tjnj/202607/P020260716686337458968.rar'
yn_dir=ROOT/'raw/population_yunnan2020_excel/excel'
files=list(yn_dir.glob('a3-1*.xlsx'))
assert len(files)==4
def parse_age_table(p):
 w=openpyxl.load_workbook(p,read_only=True,data_only=True);s=w.active
 title=s.cell(1,1).value;assert '分年龄' in title and '性别' in title
 ages={};groups={};total=None
 for i,r in enumerate(s.iter_rows(values_only=True),1):
  label=str(r[0] or '').strip()
  if label=='合计':total=tuple(int(v) for v in r[1:4])
  if re.fullmatch(r'\d+',label) and int(label)<20:
   ages[int(label)]=(tuple(int(v) for v in r[1:4]),i)
  if re.fullmatch(r'\d+-\d+',label):groups[label]=tuple(int(v) for v in r[1:4])
 assert set(ages)==set(range(20))
 return ages,groups,total,s.title
yn_tables={}
for p in files:
 suffix=re.match(r'a3-1([abc]?)',p.name).group(1)
 yn_tables[suffix]=parse_age_table(p)+(p,)
assert set(yn_tables)=={'','a','b','c'}
yn_age,yn_groups,yn_total,yn_sheet,yn_file=yn_tables['']
for age,(counts,i) in yn_age.items():
 rows.append(make('云南',age,counts,yn_url,str(yn_file.relative_to(ROOT)).replace('\\','/'),
 'official Excel archive; original cell extraction',f'sheet {yn_sheet}; A{i}:D{i}',
 'official bulletin note reviewed; yearbook editor note pending','not certified from editor note',
 'official bulletin excludes HK/Macao/Taiwan residents and foreigners; other book exclusions pending'))

def check(province,label,actual,expected):
 gap=actual-expected;assert gap==0,(province,label,actual,expected)
 checks.append(dict(province=province,check=label,actual=actual,expected=expected,gap=gap))
group_expected={'福建':[(2549344,1380757,1168587),(2993870,1627920,1365950),(2482011,1342664,1139347),(2012096,1081131,930965)],
 '吉林':[(769139,396542,372597),(960287,498408,461879),(1089297,567346,521951),(1056911,546207,510704)],
 '云南':[yn_groups[f'{lo}-{lo+4}'] for lo in [0,5,10,15]]}
for province in ['福建','吉林','云南']:
 rr=sorted([r for r in rows if r['province']==province],key=lambda r:r['age_completed'])
 assert [r['age_completed'] for r in rr]==list(range(20))
 for r in rr:check(province,f"age {r['age_completed']}; male+female",r['male']+r['female'],r['total'])
 for gi,lo in enumerate([0,5,10,15]):
  for j,col in enumerate(['total','male','female']):
   check(province,f'age {lo}-{lo+4}; {col}',sum(r[col] for r in rr[lo:lo+5]),group_expected[province][gi][j])
 for hi,expected in ({'福建':[(14,8025225),(15,8453900)],'吉林':[(14,2818723),(15,3046210)],'云南':[(14,9237474),(15,9818670)]}[province]):
  check(province,f'age 0-{hi} versus official bulletin',sum(r['total'] for r in rr[:hi+1]),expected)
# Yunnan: all-area partition checked against city/town/village tables, not school locations.
for age in range(20):
 for j,col in enumerate(['total','male','female']):
  check('云南',f'age {age}; city+town+village {col}',sum(yn_tables[k][0][age][0][j] for k in ['a','b','c']),yn_age[age][0][j])
for j,col in enumerate(['total','male','female']):
 check('云南','all-age '+col+' versus official bulletin',yn_total[j],[47209277,24420924,22788353][j])

windows=[];composition=[]
for province in ['福建','吉林','云南']:
 rr=[r for r in rows if r['province']==province]
 for lo,hi in [(0,4),(5,9),(10,14),(15,19),(6,11),(12,14),(6,10),(11,14)]:
  selected=[r for r in rr if lo<=r['age_completed']<=hi]
  windows.append(dict(province=province,year=2020,age_window=f'{lo}-{hi}',age_count=hi-lo+1,
   total=sum(r['total'] for r in selected),male=sum(r['male'] for r in selected),female=sum(r['female'] for r in selected),
   interpretation='observed age window; not enrollment, certified education demand, or annual series',note_status=rr[0]['note_status']))
 g=[next(w['total'] for w in windows if w['province']==province and w['age_window']==k) for k in ['0-4','5-9','10-14']]
 composition.append(dict(province=province,year=2020,age0_4=g[0],age5_9=g[1],age10_14=g[2],
  age0_4_to_age5_9=g[0]/g[1],age5_9_to_age10_14=g[1]/g[2],
  interpretation='same-time equal-width age composition ratios; not growth rates or forecasts',comparability_status=rr[0]['note_status']))
save('population_census2020_expanded_single_age.csv',rows)
save('population_census2020_expanded_age_windows.csv',windows)
save('population_census2020_expanded_checks.csv',checks)
save('population_census2020_expanded_composition.csv',composition)

names=['北京','天津','河北','山西','内蒙古','辽宁','吉林','黑龙江','上海','江苏','浙江','安徽','福建','江西','山东','河南','湖北','湖南','广东','广西','海南','重庆','四川','贵州','云南','西藏','陕西','甘肃','青海','宁夏','新疆']
statuses={
 '福建':('extracted; editor note reviewed','20 ages; first pilot complete'),
 '吉林':('extracted; editor note reviewed','20 ages; province geography note retained; first timeout resolved'),
 '云南':('extracted; editor note pending','20 ages; Excel archive valid; bulletin matched; PDF archive truncated/unusable'),
 '黑龙江':('catalog located; original table inaccessible','prior direct download HTTP 403; web timeout'),
 '上海':('catalog located; original table not locally obtained','official table link confirmed; direct download HTTP 403'),
 '四川':('official source lead; catalog not obtained','catalog connection closed; no age table certified'),
 '湖南':('official yearbook entry located; table pending','official entry web timeout; no age table certified'),
 '陕西':('official census publication lead; age table pending','only unrelated/long-form volume located; not target table'),
 '广西':('search lead only; target table pending','no original age table certified'),
 '浙江':('search lead only; target table pending','no original province age table certified')}
coverage=[dict(province=p,year=2020,status=statuses.get(p,('not investigated in current bounded rounds','not evidence of unavailable data'))[0],
 reason=statuses.get(p,('','not evidence of unavailable data'))[1],extracted_ages='0-19' if p in ['福建','吉林','云南'] else '',
 annual_panel_eligible=False,school_enrollment_denominator_certified=False) for p in names]
save('population_census2020_province_coverage.csv',coverage)
# Hash the precise sources consumed, plus archive health separately.
sources=read('population_province2020_sources.csv')
for name,url,status in [
 ('population_jilin2020_catalog.html','https://tjj.jl.gov.cn/tjsj/qwfb/jlsdqcqgrkpcnj/left.htm','saved; catalog inspected'),
 ('population_jilin2020_age.jpg',jl_url,'saved; original reviewed'),
 ('population_jilin2020_note.jpg','https://tjj.jl.gov.cn/tjsj/qwfb/jlsdqcqgrkpcnj/zk/html/note.jpg','saved; original reviewed'),
 ('population_jilin2020_bulletin.pdf','https://tjj.jl.gov.cn/tjsj/qwfb/jlsdqcqgrkpcnj/zk/html/fu0103.pdf','saved; published age aggregate checked'),
 ('population_yunnan2020_entry.html','https://stats.yn.gov.cn/zwgk/zfxxgk/fdzdgknr/tjsj/tjnj/202607/t20260716_3205140.html','saved; archive link inspected'),
 ('population_yunnan2020_excel.rar',yn_url,'archive valid; 12 selected workbooks extracted'),
 ('population_yunnan2020_bulletin.html','https://stats.yn.gov.cn/zwgk/zfxxgk/fdzdgknr/tjsj/tjgb/202105/t20210517_1204456.html','saved; age counts and resident definition reviewed'),
 ('population_yunnan2020_pdf_entry.html','https://stats.yn.gov.cn/zwgk/zfxxgk/fdzdgknr/tjsj/tjnj/202402/t20240204_1202761.html','saved; PDF archive link inspected'),
 ('population_yunnan2020_pdf.rar','https://stats.yn.gov.cn/zwgk/zfxxgk/fdzdgknr/tjsj/tjnj/202402/P020251011359712733211.rar','truncated/unusable; excluded as evidence')]:
 p=ROOT/'raw'/name;sources.append(dict(url=url,path='raw/'+name,bytes=p.stat().st_size,sha256=sha(p),status=status))
for r in sources:r.setdefault('status','first pilot original source saved')
for p in sorted(yn_dir.glob('*.xlsx')):
 sources.append(dict(url=yn_url,path=str(p.relative_to(ROOT)).replace('\\','/'),bytes=p.stat().st_size,sha256=sha(p),status='archive member; only 3-1 family consumed in numerical checks'))
assert all(sha(ROOT/r['path'])==r['sha256'] for r in sources)
save('population_census2020_expanded_sources.csv',sources)
protected=json.loads((DATA/'province_finance_next_protected_baseline.json').read_text(encoding='utf-8'))
assert all(sha(ROOT/p)==h for p,h in protected.items())
audit=dict(provinces_with_extracted_counts=3,provinces_with_editor_note_reviewed=2,
 provinces_with_editor_note_pending=1,province_age_rows=len(rows),population_count_cells=len(rows)*3,
 arithmetic_checks=len(checks),arithmetic_checks_passed=True,source_records=len(sources),
 coverage_inventory_provinces=len(coverage),not_investigated=sum(c['status'].startswith('not investigated') for c in coverage),
 archived_pdf_unusable=True,annual_population_panel_completed=False,population_to_enrollment_division=False,
 figures_created_or_updated=False,protected_files_unchanged=True,
 limit='selected source coverage, not a representative province sample; Yunnan editor exclusions pending')
(DATA/'population_census2020_expanded_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(audit,ensure_ascii=False));print(json.dumps(composition,ensure_ascii=False))
