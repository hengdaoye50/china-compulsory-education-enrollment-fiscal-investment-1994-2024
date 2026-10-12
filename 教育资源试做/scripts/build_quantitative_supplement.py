"""Concise, illustrated supplement to the existing long-history report."""
from pathlib import Path
import json
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[2];DATA=ROOT/'教育资源试做/data/定量深化_20261011';FIG=ROOT/'教育资源试做/定量深化图组_20261011';OUT=ROOT/'定量补充报告_20261011';OUT.mkdir(exist_ok=True)
test=pd.read_csv(DATA/'four_direct_contrast_tests.csv')
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.8)
sec.header_distance=sec.footer_distance=Inches(.3)
def font(run,name='宋体',size=11,bold=False):
    run.font.name=name;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),name)
for sn,size in [('Normal',11),('Title',19),('Subtitle',10),('Heading 1',13),('Caption',9)]:
    st=doc.styles[sn];st.font.name='宋体' if sn not in ['Title','Heading 1'] else '黑体';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),st.font.name)
    st.paragraph_format.line_spacing=1.2;st.paragraph_format.space_after=Pt(7)
    if sn=='Heading 1':st.paragraph_format.keep_with_next=True
for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
header=sec.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.RIGHT;font(header.add_run('义务教育经费长期研究  定量补充'),size=8)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
def para(t,small=False):
    p=doc.add_paragraph(t);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if small:
        for r in p.runs:font(r,size=9)
    return p
def heading(t):doc.add_heading(t,1)
def fig(name,caption,width=6.6):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(FIG/f'{name}.png'),width=Inches(width))
    doc.add_paragraph(caption,'Caption')
def page():doc.add_page_break()
doc.add_paragraph('义务教育支出调整定量补充报告','Title')
doc.add_paragraph('2008—2023年省级面板分析    2026年10月11日','Subtitle')
para('本轮将三十年历史描述推进到支出内部调整的关系检验。最清楚的结果是：初中公用支出与区域学生变化的关联强于个人支出；小学方向相同但证据较弱。地方自身收入的支出差异仍缺乏足够证据。结果为探索性条件关联，不能解释为政策因果效果。')
heading('问题与方法')
para('“学生减少而经费增长”本身不能区分最低办学成本、人员存量约束和配置变化。我们直接问：控制省份稳定差异与年度共同变化后，个人、公用支出的学生响应和自身收入响应是否不同？')
para('历史窗口仍为1995—2024年，1994暂估只作历史补充。回归以2007年为基期、分析2008—2023年变化，避开2006→2007分类转换。两学段各472条完整省年、31个省级聚类；各24条相邻金额缺口排除，不插值、不缩尾，也不按支出涨跌筛样本。')
para('分别以个人、公用实际对数变化为结果，同时纳入学生与实际自身收入对数变化、省份效应和年份效应，各省年等权。差异模型D＝公用对数变化−个人对数变化，直接检验两类响应之差。四项主检验采用9999次省级约束野生聚类自助并作Holm校正。')
table=doc.add_table(rows=1,cols=5);table.autofit=False
widths=[.65,1.35,1,1.85,1.05]
headers=['学段','解释变量','响应差','单项95%区间','校正p值']
for i,(cell,txt) in enumerate(zip(table.rows[0].cells,headers)):
    cell.text=txt;cell.width=Inches(widths[i]);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EAEAEA');cell._tc.get_or_add_tcPr().append(sh)
for row in test.itertuples(index=False):
    vals=['小学' if row.stage=='primary' else '初中','学生变化' if row.regressor=='dS' else '自身收入变化',f'{row.coefficient:.3f}',f'{row.WCR_95_CI_low:.3f} 至 {row.WCR_95_CI_high:.3f}',f'{row.Holm_four_p:.4f}']
    cells=table.add_row().cells
    for i,(cell,txt) in enumerate(zip(cells,vals)):cell.text=txt;cell.width=Inches(widths[i])
for row in table.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5)
            for r in p.runs:font(r,size=9)
para('P/C为全隶属、全来源个人/公用支出，不能等同纯工资、纯运行经费或政府预算。学生学校覆盖尚未完全桥接，仅表示区域生源背景。R为自身收入，不代表转移支付后的全部财力；省级CPI只是购买力近似。',small=True)
page();heading('直接差异检验')
fig('图Q1_直接差异与置信区间','图1  点为公用减个人的响应系数差，线为省级约束野生聚类自助检验反演的单项95%区间；Holm p对四项检验校正。区间并非四项同时覆盖区间。')
para('初中学生响应差约0.601，单项95%区间0.229至0.982，校正p＝0.0144。个人与公用的学生系数分别约0.036和0.637。两类支出的调整与生源变化存在不同关联，不能只用总额趋势代替内部配置分析。')
para('以学生变化由0变为−10%作为线性模型的尺度例子，D的拟合差异约−6.33个对数百分点，单项95%区间约−10.35至−2.42。这不表示公用支出必然下降6.33%，也不是学生减少造成的因果效应。真实支出还取决于收入、年度及其他因素。')
para('小学学生响应差约0.695，但校正p＝0.0879。方向与初中一致，当前证据不足以在四项检验族的5%阈值下确认。两个学段各自显著与否不同，不等于学段间差异已经成立。')
para('自身收入响应差在小学、初中分别约0.217和0.164，主自助区间均跨零，校正p分别为0.1388、0.1884。此前持续生源下降样本的H1/H2弱证据仍保留，本轮不能解释为旧假设验证成功。')
page();heading('分项响应与稳定性')
fig('图Q2_支出分项响应','图2  两学段分别估计个人、公用支出模型。误差线为省级聚类CR1标准误×1.96的正态近似95%区间；不能比较两个区间是否重叠来替代图1直接检验。',width=6.5)
para('初中学生差系数在预定样本敏感性中约0.506至0.671；逐省删除约0.528至0.701、逐年删除约0.439至0.746。单一省份或年份没有翻转方向。这些是点估计稳定性诊断，不能称为所有规格都通过相同自助检验。')
para('剔除2015/2016财政分类边界附近变化后，初中自身收入差从0.164降至0.025，财政关系对边界较敏感。加入一期滞后后，初中学生当期与滞后系数之和约0.459、CR1标准误0.203；小学累计约1.054、标准误0.350。调整时间值得研究，但滞后不构成因果识别。')
para('将学生增加与下降幅度分开后，初中下降幅度的D系数为−0.462、CR1标准误0.275，正态近似区间跨零。因此主结果是增长与下降合并样本中的平均线性响应差，不能专门宣称生源下降时公用挤压已获证实。')
page();heading('观测支持与后续研究')
fig('图Q3_偏回归散点','图3  全部省年观测。两轴均剔除省份效应、年份效应及另一个解释变量；线斜率等于主回归系数。无分箱、无删极值；纵轴按面板观测范围独立绘制。',width=6.1)
para('固定效应不消除地方工资政策、转移支付、建设项目及经济变化等时变遗漏因素；省级聚类也不能完全处理跨省相关。当前结果与人员相关支出较难随生源同步调整的解释相容，但没有证明刚性机制，更不能评价支出效率。')
para('报告主线宜明确为：长期总量保障背景下，支出结构如何调整，哪些项目与生源更同步，哪些更受自身收入条件影响。下一步优先核准转移支付及教师、学校、班级存量，检验竞争解释；这些可能处于调整路径中，不能机械加入所有主模型。')
para('数据、主检验、全部敏感性、逐省逐年删除、图表数据及28项复算审校见“教育资源试做/data/定量深化_20261011”；完整方法与来源说明见该目录《定量深化研究说明》。三十年历史报告和旧模型均保留。',small=True)
para('聚类自助方法参考 Cameron、Gelbach 与 Miller（2008），Review of Economics and Statistics，90(3)：414—427。本轮为探索扩展，未完成学校覆盖与财力定义的全部桥接。',small=True)
path=OUT/'义务教育支出调整定量补充报告.docx';doc.save(path)
check=Document(path);assert len(check.inline_shapes)==3 and len(check.tables)==1
(OUT/'报告制作记录.json').write_text(json.dumps({'figures':3,'tables':1,'body_source':'定量深化研究说明.md','numerical_source':'four_direct_contrast_tests.csv','scope':'exploratory supplement; existing report retained'},ensure_ascii=False,indent=2),encoding='utf8')
print(path)
