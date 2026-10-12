"""Retain public finance candidates and arithmetic issues without joining denominators."""
import json,re,hashlib
from decimal import Decimal as D
from docx import Document
from parse import ROOT,read,save_csv
manifests=['finance_gap_sources.json','finance_gap_followup_sources.json','finance_gap_attachment_sources.json','finance_gap_classification_sources.json']
sources=[s for name in manifests for s in json.loads((ROOT/name).read_text(encoding='utf-8'))]
bykey={s['key']:s for s in sources};text={}
for s in sources:
    if s['status']!='downloaded':continue
    p=ROOT/s['path'];assert hashlib.sha256(p.read_bytes()).hexdigest()==s['sha256']
    if p.suffix=='.html':text[s['key']]=' '.join(read(p).text)
    elif p.suffix=='.docx':
        doc=Document(p);text[s['key']]='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in row.cells) for tab in doc.tables for row in tab.rows)
    (ROOT/'data'/f"{s['key']}_text.txt").write_text(text[s['key']],encoding='utf-8')
cells=[];checks=[]
def add(key,year,metric,value,quote,stage='all'):
    normalized=re.sub(r'\s','',text[key]);assert re.sub(r'\s','',quote) in normalized,(key,quote)
    cells.append({'year':year,'region':'威海' if 'weihai' in key else '江阴','stage':stage,'metric':metric,'value':value,'unit':'亿元','evidence_excerpt':quote,'source':bykey[key]['url'],'path':bykey[key]['path'],'sha256':bykey[key]['sha256'],'eligibility':'candidate_only; scope/denominator unresolved'})
for y,vals in [(2022,('118.27','117.85','63.36','37.00','26.36')),(2023,('121.56','120.53','64.88','38.75','26.13'))]:
    key=f'finance_gap_weihai_{y}'
    career_quote=f"教育事业费支出{'，' if y==2022 else ''}{vals[1]}亿元"
    for m,v,q,s in [('education_total_expenditure',vals[0],f'总支出为{vals[0]}亿元','all'),('caption_budget_career_total',vals[1],career_quote,'all'),('caption_budget_career_compulsory',vals[2],f'义务教育支出{vals[2]}亿元','compulsory'),('caption_budget_career_junior',vals[3],f'普通初中支出{vals[3]}亿元','junior'),('caption_budget_career_primary',vals[4],f'普通小学支出{vals[4]}亿元','primary')]:add(key,y,m,v,q,s)
    a=sum(D(v) for v in vals[3:]);checks.append({'region':'威海','year':y,'check':'primary_plus_junior_vs_compulsory','computed':str(a),'reported':vals[2],'difference_100m':str(a-D(vals[2])),'status':'pass'})
for y,parts,total in [(2022,['15.01','10.00','11.66','63.36','11.53','0.60','5.69'],'117.85'),(2023,['16.22','8.76','13.45','64.88','12.17','0.55','5.52'],'120.53')]:
    for p in parts:assert p in re.sub(r'\s','',text[f'finance_gap_weihai_{y}'])
    value=sum(D(p) for p in parts);diff=value-D(total)
    checks.append({'region':'威海','year':y,'check':'education_category_subtotals_vs_caption_budget_career_total','computed':str(value),'reported':total,'difference_100m':str(diff),'status':'pass' if abs(diff)<=D('.01') else 'unresolved'})
checks.append({'region':'威海','year':2023,'check':'total_growth_baseline_vs_2022_page','computed':str(D('121.56')-D('3.71')),'reported':'118.27','difference_100m':str(D('121.56')-D('3.71')-D('118.27')),'status':'unresolved; prior-year revision not explained'})
checks.append({'region':'威海','year':2023,'check':'career_growth_baseline_vs_2022_page','computed':str(D('120.53')-D('3.10')),'reported':'117.85','difference_100m':str(D('120.53')-D('3.10')-D('117.85')),'status':'unresolved; prior-year revision not explained'})
key='finance_gap_jiangyin_2024_attachment';t=re.sub(r'\s','',text[key])
sections=[('budget_per_student','2.生均一般公共预算教育经费增长情况'),('career_per_student','3.生均一般公共预算教育事业费增长情况'),('public_per_student','4．生均一般公共预算公用经费支出增长情况')]
for metric,start in sections:
    pos=t.index(start);end=t.find('三、',pos) if metric=='public_per_student' else t.index(sections[[m for m,_ in sections].index(metric)+1][1],pos)
    chunk=t[pos:end]
    for st,label in [('primary','普通小学'),('junior','普通初中')]:
        match=re.search(label+r'生均一般公共预算.*?([\d.]+)元，比上年([\d.]+)元.*?(增长|下降)([\d.]+)%',chunk)
        assert match,(metric,st)
        current,previous,direction,growth=match.groups();growth=D(growth)*(1 if direction=='增长' else -1)
        computed=100*(D(current)/D(previous)-1);assert abs(computed-growth)<D('.006')
        for year,value in [(2023,previous),(2024,current)]:
            cells.append({'year':year,'region':'江阴','stage':st,'metric':metric,'value':value,'unit':'元/生','evidence_excerpt':match[0],'source':bykey[key]['url'],'path':bykey[key]['path'],'sha256':bykey[key]['sha256'],'eligibility':'published per-student observation; total expenditure and denominator not matched'})
        checks.append({'region':'江阴','year':2024,'stage':st,'metric':metric,'check':'per_student_growth_rounding','computed':str(computed),'reported':str(growth),'difference_pp':str(computed-growth),'status':'pass'})
save_csv(ROOT/'data/finance_gap_candidate_cells.csv',cells)
save_csv(ROOT/'data/finance_gap_quality_checks.csv',checks)
receipt=[]
for s in sources:receipt.append({k:s.get(k) for k in ['key','url','status','path','bytes','sha256','error']})
save_csv(ROOT/'data/finance_gap_source_receipt.csv',receipt)
key='finance_gap_pku_classification'
assert '农村指镇区和乡村' in text[key]
assert '这里使用的教育事业数据来自' in text[key]
save_csv(ROOT/'data/finance_gap_scope_evidence.csv',[{
    'issue':'finance rural vs resource village category',
    'source':bykey[key]['url'],'path':bykey[key]['path'],'sha256':bykey[key]['sha256'],
    'locator':'section 3 subsection 3(1), footnote 6',
    'finding':'Research describes finance rural as town plus village after the 2014 classification change; footnote 6 uses finance-statistics database student counts.',
    'evidence_level':'original institutional research, not annual official reporting instructions',
    'eligibility':'scope lead only; no category join or denominator substitution authorized by evidence'
}])
validation={'download_records':len(sources),'downloaded_records':sum(s['status']=='downloaded' for s in sources),'failed_records':[s['key'] for s in sources if s['status']!='downloaded'],'candidate_cells':len(cells),'checks':checks,'same_scope_expenditure_denominator_pairs':0,'national_finance_panel_modified':False,'finding':'Public local reports can contain stage totals; this round did not establish a compatible expenditure/denominator pair. No reconstructed expenditure or denominator.'}
(ROOT/'data/finance_gap_audit.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(validation,ensure_ascii=False))
