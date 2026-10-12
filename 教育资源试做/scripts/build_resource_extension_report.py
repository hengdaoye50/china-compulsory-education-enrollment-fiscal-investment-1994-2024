"""Illustrated research extension; preserves all earlier reports."""
from pathlib import Path
import json
import pandas as pd
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[2];DATA=ROOT/'教育资源试做/data/资源解释拓展_20261011';FIG=ROOT/'教育资源试做/资源解释图组_20261011';OUT=ROOT/'资源解释拓展报告_20261011';OUT.mkdir(exist_ok=True)
tt=pd.read_csv(DATA/'four_resource_contrast_tests.csv')
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.8);sec.header_distance=sec.footer_distance=Inches(.3)
def font(run,size=11,bold=False):
    run.font.name='宋体';run.font.size=Pt(size);run.font.bold=bold;run.font.italic=False;run.font.color.rgb=RGBColor(0,0,0)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
for name,size in [('Normal',11),('Title',19),('Subtitle',10),('Heading 1',13),('Caption',9)]:
    st=doc.styles[name];st.font.name='黑体' if name in ['Title','Heading 1'] else '宋体';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.font.italic=False
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),st.font.name);st.paragraph_format.line_spacing=1.2;st.paragraph_format.space_after=Pt(7)
    if name=='Heading 1':st.paragraph_format.keep_with_next=True
for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
header=sec.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.RIGHT;font(header.add_run('义务教育经费长期研究  资源解释拓展'),size=8)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
def para(text,small=False):
    p=doc.add_paragraph(text);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if small:
        for r in p.runs:font(r,size=9)
    return p
def heading(text):doc.add_heading(text,1)
def figure(name,text,width=6.5):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(FIG/f'{name}.png'),width=Inches(width));doc.add_paragraph(text,'Caption')
doc.add_paragraph('义务教育办学资源与支出调整拓展报告','Title')
doc.add_paragraph('2014—2023年省级数据    2026年10月11日','Subtitle')
para('本轮追问初中公用与个人支出的生源响应差异能否从教师、班级和学校变化中获得解释。初中同样本模型的学生响应差，在加入班级后变小，加入教师后则没有同样变化；但资源主检验校正后均证据不足。办学组织提供了统计线索，尚不能宣布人员刚性或班级配置机制成立。')
heading('问题与方法')
para('研究仍以三十年经费保障与内部调整为主轴。本轮为上一轮关系检验的解释性拓展，保留此前模型。教师、班级和经费可能同时调整，因此资源变量只作为区域背景，加入它们不构成中介识别。')
para('回读2013—2023年66张教育部原表、2816个单元格，教师、班级、学校和学生数与既有整理层逐项一致。变化窗口为2014—2023年，两学段各310个潜在省年，各10条财政相邻缺口排除，分别300条进入模型。学生变化在全部600条连接中吻合；不补零、不插值、不缩尾。')
para('先在同样本中估计原学生、收入模型，再分别加入专任教师、班级及两者共同变化。分别估计个人、公用实际支出及差D＝公用对数变化−个人对数变化，均含省份、年份效应，各省年等权。4项主检验为两学段的教师、班级差系数，采用9999次省级约束野生聚类自助并作Holm校正。')
table=doc.add_table(rows=1,cols=5);table.autofit=False
headers=['学段','资源变量','公用减个人','单项95%区间','校正p值'];widths=[.65,1.15,1.1,1.95,1.05]
for i,(cell,text) in enumerate(zip(table.rows[0].cells,headers)):
    cell.text=text;cell.width=Inches(widths[i]);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EAEAEA');cell._tc.get_or_add_tcPr().append(sh)
for row in tt.itertuples(index=False):
    values=['小学' if row.stage=='primary' else '初中','专任教师' if row.regressor=='dT' else '班级',f'{row.coefficient:.3f}',f'{row.WCR_95_CI_low:.3f} 至 {row.WCR_95_CI_high:.3f}',f'{row.Holm_four_p:.4f}']
    cells=table.add_row().cells
    for i,(cell,value) in enumerate(zip(cells,values)):cell.text=value;cell.width=Inches(widths[i])
for row in table.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5)
            for run in p.runs:font(run,size=9)
para('经费为全隶属全来源个人/公用支出，非纯工资、纯运行或财政预算；资源与经费学校覆盖尚未完全桥接。专任教师不等于全部人员、学校数不等于所有教学网点。自身收入不代表补助后的全部财力。',small=True)
doc.add_page_break();heading('区分时期变化与资源变量变化')
figure('图E1_同样本模型比较','图1  原窗口为2008—2023年，其余模型使用同一2014—2023年省年。点为学生响应差，线为省级CR1及t(30)近似95%区间。模型步骤不表示中介解释比例。')
para('初中原窗口学生响应差为0.601；缩至资源共同窗口后为0.473。这个变化首先包括时期范围差异，不能归因于资源变量。共同样本中，加入教师后为0.566、加入班级后为0.151，同时加入后为0.073、CR1标准误0.582。')
para('点估计变化更多出现在班级进入时，但区间同时变宽。不能将0.473到0.073的变化写成“解释了85%”，也不能认定班级为中介。生源与班级共同变化、资源和经费同时决定以及过度控制，都可能改变系数。')
para('小学没有相同模式：共同样本基准为0.538，同时加入教师班级后为0.787、标准误0.690。不能由初中结果概括两学段共同机制，也不能根据各自显著性不同宣称学段差异成立。')
para('固定效应残差化后的VIF小学约1.03—2.04、初中1.01—3.19，未见极端数值共线性，但不等于机制可以分别识别。最大杠杆约0.27，保留影响诊断，不删除极值以美化关系。')
doc.add_page_break();heading('资源关联与证据强度')
figure('图E2_资源关联直接差异','图2  同时纳入学生、实际自身收入、教师、班级变化及省份/年份效应。线为单项约束野生聚类自助检验反演95%区间；Holm p对4项校正，区间不是家族同时区间。')
para('初中教师差系数为−0.836，表示在既有条件下，教师增长与个人支出相对于公用的增长更强相关。未经校正自助p为0.0415，四项校正后为0.1660；不能挑选未经校正结果宣布人员机制成立。班级差为0.800，区间宽、校正p为0.6387，同样证据不足。')
para('初中教师差逐省删除约−1.038至−0.687、逐年删除−1.345至−0.394，班级差分别约0.523至1.037、0.385至1.274，方向没有被单一省份或年份翻转。方向稳定不能替代不足的主推断。小学教师逐省、班级逐年删除存在符号翻转，更不宜概括为稳定关系。')
para('资源背景模型的学生系数：班级在两学段均约0.657，教师小学约0.233、初中0.149。这是分模型点估计对照，未列入主检验族。学生减少的77条小学省年中，35条教师增加、12条班级增加；初中103条中分别为52条、24条。资源并非随生源机械变化，但这些并存现象不是低效率证据。')
doc.add_page_break();heading('分项支出与下一步解释研究')
figure('图E3_资源与分项支出','图3  与图2相同样本与解释变量，分别估计个人、公用实际支出。线为CR1及t(30)近似95%区间；不能用两个分模型区间是否重叠替代图2直接差异检验。')
para('初中教师对个人、公用的分模型系数约0.379、−0.457，班级约0.199、0.999。两类资源可能涉及不同支出关联，但当前精度和学校范围限制不足以支持单一路径。个人不等于全部教师工资，公用也包含部分资本性项目。')
para('加入学校、剔除边界年份或疫情年份以及滞后资源的结果均完整保留。初中滞后教师、班级差约−0.462、0.287；小学约−0.776、0.034，且滞后样本为269条。不能用当期与滞后符号变化证明时间机制，也不挑选学校辅助系数成为新核心发现。')
para('本轮把解释范围进一步收窄：初中差异不能简单归因于教师数量，班级组织值得跟踪，但明确机制尚未建立。下一步优先核准转移支付、工资政策与经费细项，区分自身收入与补助后财力，并寻找人员成本和运行项目的独立证据。固定效应和数据回读不能消除反向因果或时变遗漏因素。')
para('完整方案、600条模型输入、20条连接排除、2816条原表回读、主检验、全部敏感性、影响诊断、图注与48项复算检查均另存于“教育资源试做/data/资源解释拓展_20261011”。旧报告与旧数据保留；本轮为有明确口径限制的探索扩展。',small=True)
path=OUT/'义务教育办学资源与支出调整拓展报告.docx';doc.save(path)
assert len(Document(path).inline_shapes)==3
(OUT/'报告制作记录.json').write_text(json.dumps({'figures':3,'tables':1,'numerical_source':'four_resource_contrast_tests.csv','analysis_validation_checks':48,'old_reports_preserved':True},ensure_ascii=False,indent=2),encoding='utf8')
print(path)
