"""Create a concise source-backed Chinese research report with editable Word text."""
from pathlib import Path
import json,hashlib
import pandas as pd
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'正式分析报告_20261011';OUT.mkdir(exist_ok=True)
FIG=ROOT/'教育资源试做/正式学术图组_20261011'
DATA=ROOT/'教育资源试做/data'
F=pd.read_csv(DATA/'早期财政重建_20261011/national_fiscal_1994_2024_report_final.csv')
S=pd.read_csv(FIG/'figure_data/fig01.csv')
P=pd.read_csv(FIG/'figure_data/fig04.csv')
W=pd.read_csv(FIG/'figure_data/fig03.csv')
G=pd.read_csv(FIG/'figure_data/fig05.csv')
PAIR=pd.read_csv(FIG/'figure_data/fig06.csv')
REG=pd.read_csv(FIG/'figure_data/fig07.csv')
D=pd.read_csv(FIG/'figure_data/fig08.csv')
names={'primary':'小学','junior':'初中'}
def finance(stage,year):return F[(F.stage==stage)&(F.year==year)].iloc[0]
def student(stage,year):return S[(S.stage==stage)&(S.year==year)].iloc[0].students/10000
def personal(stage,year):return P[(P.stage==stage)&(P.year==year)].iloc[0].personal_share_of_PC_pct
def weight(stage,year):return W[(W.stage==stage)&(W.year==year)].iloc[0].income_weight_pct
def median(stage,start):return G[(G.stage==stage)&(G.start_year==start)].PC_real_annualized_pct.median()
def decomposition(stage,start):return D[(D.stage==stage)&(D.start_year==start)].iloc[0]
def number(x):return f'{x:,.2f}'

doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.9);sec.bottom_margin=Cm(1.9)
sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)
sec.header_distance=Cm(.8);sec.footer_distance=Cm(.8)
def setfont(obj,name,size,bold=False):
    obj.font.name=name;obj.font.size=Pt(size);obj.font.bold=bold;obj.font.italic=False;obj.font.underline=False;obj.font.color.rgb=RGBColor(0,0,0)
    rpr=obj._element.get_or_add_rPr();f=rpr.find(qn('w:rFonts'))
    if f is None:f=OxmlElement('w:rFonts');rpr.insert(0,f)
    for k in ['ascii','hAnsi','eastAsia','cs']:f.set(qn('w:'+k),name)
    for a in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        f.attrib.pop(qn('w:'+a),None)
    for tag in ['iCs','spacing']:
        for e in list(rpr.findall(qn('w:'+tag))):rpr.remove(e)
for style,font,size,bold in [('Normal','宋体',11,False),('Title','黑体',20,True),('Subtitle','宋体',10,False),('Heading 1','黑体',13,True),('Heading 2','黑体',11.5,True),('Caption','宋体',9,False)]:
    st=doc.styles[style];setfont(st,font,size,bold)
    st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.line_spacing=1.25
    st.paragraph_format.widow_control=True
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Title'].paragraph_format.space_after=Pt(8)
for st in ['Heading 1','Heading 2']:
    doc.styles[st].paragraph_format.keep_with_next=True
    doc.styles[st].paragraph_format.space_before=Pt(5)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(7)
for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
setting=OxmlElement('w:doNotCompressPictures');doc.settings.element.append(setting)
header=sec.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
setfont(header.add_run('义务教育经费长期演进分析报告'),'宋体',8)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
setfont(footer.add_run('第 '),'宋体',9)
run=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)
setfont(footer.add_run(' 页'),'宋体',9)
md=[]
def para(text,small=False):
    p=doc.add_paragraph(text);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if small:
        for r in p.runs:setfont(r,'宋体',9)
        p.paragraph_format.line_spacing=1.2;p.paragraph_format.space_after=Pt(5)
    md.append(text+'\n');return p
def heading(text,level=1):doc.add_heading(text,level);md.append('#'*(level+1)+' '+text+'\n')
def page():doc.add_page_break();md.append('\n---\n')
def figure(n,width=16.5):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(5);p.paragraph_format.line_spacing=1
    pic=p.add_run().add_picture(str(FIG/'png'/f'fig{n:02d}.png'),width=Cm(width))
    meta=json.loads((FIG/'figure_manifest.json').read_text(encoding='utf8'))['figures'][n-1]
    pic._inline.docPr.set('descr',meta['title']+'。'+''.join(meta['notes']))
    md.append(f'![图{n} {meta["title"]}](../教育资源试做/正式学术图组_20261011/png/fig{n:02d}.png)\n')
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for cell,v in zip(t.rows[0].cells,headers):cell.text=v
    for row in rows:
        for cell,value in zip(t.add_row().cells,row):cell.text=str(value)
    for i,row in enumerate(t.rows):
        if i==0:
            repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        for j,c in enumerate(row.cells):
            c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for side,value in [('top',75),('bottom',75),('left',95),('right',95)]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),str(value));e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEDED');pr.append(sh)
            for p in c.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.1
                for r in p.runs:setfont(r,'宋体',9,i==0)
    md.append('| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+''.join('| '+' | '.join(map(str,r))+' |\n' for r in rows))
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

title='生源变动与义务教育经费的长期演进'
doc.add_paragraph(title,style='Title')
doc.add_paragraph('1995至2024年分析报告',style='Subtitle')
doc.add_paragraph('2026年10月11日',style='Subtitle')
md.append('# '+title+'\n\n1995至2024年分析报告\n')
heading('摘要')
para('三十年资料表明，小学与初中的生源转折不同步；财政性经费名义金额总体扩张，其收入占比却受到学校举办者构成影响；地方隶属学校个人与公用支出总体增长，后期增速放缓，部分地区公用支出出现实际下降。判断义务教育保障状况，需要同时观察学段需求、资金金额、收入来源和支出结构。本文提供长期描述与构成分解，不识别政策因果效应。')
heading('一 问题与方法')
para('人口和学生规模变化使义务教育资源配置面临持续调整。核心问题是：经费增长是否与生源变化同步，财政性收入占比能否代表保障强度，全国增长是否掩盖地区和支出项目的差异？单一年份或单一比例难以回答这些问题。')
para('研究主体为1995至2024年普通小学与初中，1994年仅作前置暂估补充。学生数、经费及物价资料来自教育统计、教育经费统计与综合统计年鉴，按实际数据年整理。全国收入涵盖不同举办者；支出分析限定为地方隶属学校，两类资料分别解释。')
para('方法包括长期趋势描述、分段端点比较、地区配对及举办者收入构成的中点恒等式分解。地区实际增长率用本地区居民消费价格指数（CPI）扣除一般物价变化；年均增长率采用复合计算。固定地区比较采用相同的28个省级行政单位，后期地区分布采用31个。')
para('收入、支出与学校覆盖范围存在跨期变化，曲线在重要定义边界断开；不跨边界计算统一长期增速，也不将不同范围的经费除以学生数计算生均投入。个人支出不等同于教师工资，公用支出包含原表列示的资本性等项目。',small=True)

page();heading('二 分析结果');heading('生源规模存在学段错位',2)
para(f'小学在1997年达到观察期峰值约{student("primary",1997):,.0f}万人，至2013年降为{student("primary",2013):,.0f}万人，随后回升。初中峰值在2003年，低点在2015年，分别约{student("junior",2003):,.0f}万人与{student("junior",2015):,.0f}万人。两学段经历扩张、收缩和回升的时间不同，不能用某一年的下降概括整个时期。')
figure(1)
table(['学段','2024年在校生 万人','较2023年变化'],[[names[s],number(student(s,2024)),f'{S[(S.stage==s)&(S.year==2024)].iloc[0].student_change_pct:+.2f}%'] for s in names],[3,7,6.5])
para('2024年小学规模再次下降，初中仍在增长。全国总量呈现的是学段需求变化；学生流动、学校类型与统计时点也会影响这些数字，不能直接解释为出生人口变化。',small=True)

page();heading('财政性收入金额与占比并不总是同向',2)
fp13=finance('primary',2013);fj13=finance('junior',2013)
fp24=finance('primary',2024);fj24=finance('junior',2024)
para(f'财政性收入占比在2013年达到2005至2024年观察窗口的高点，小学为{fp13.fiscal_share_pct:.2f}%，初中为{fj13.fiscal_share_pct:.2f}%。至2024年分别降为{fp24.fiscal_share_pct:.2f}%与{fj24.fiscal_share_pct:.2f}%，但同期财政性收入名义金额分别增长{100*(fp24.fiscal_yi/fp13.fiscal_yi-1):.2f}%与{100*(fj24.fiscal_yi/fj13.fiscal_yi-1):.2f}%。比例下降不能直接解释为财政性投入减少。')
figure(2)
table(['学段','2024年财政性收入 亿元','占当年收入比例'],[[names[s],number(finance(s,2024).fiscal_yi),f'{finance(s,2024).fiscal_share_pct:.2f}%'] for s in names],[3,8,5.5])
para('金额为名义值。1994年为空心暂估点，1995至1997年为财政性支出，1998年起为收入；断线表示收支或分类边界，不能视为缺失，更不能将全段当作单一可比指标。',small=True)

page();heading('学校举办者的收入构成发生变化',2)
para(f'民办学校收入占同学段全国收入的比例，在2021年达到2007至2024年观察窗口的峰值：小学{weight("primary",2021):.2f}%，初中{weight("junior",2021):.2f}%；2024年回落至{weight("primary",2024):.2f}%与{weight("junior",2024):.2f}%。这一变化会影响全国财政性收入占比的分母及学校类别权重。')
figure(3)
para('这些比例是收入权重，不能当作民办学生人数比例。举办者分类变动、学校转制与各类学校收入增长，都可能改变权重；年鉴分类资料尚不足以分别识别这些过程。全国占比的构成分解见图8。')

page();heading('支出持续增长并伴随内部构成变化',2)
para(f'2007至2013年，地方隶属学校个人支出占个人与公用合计的比例，小学由{personal("primary",2007):.2f}%降为{personal("primary",2013):.2f}%，初中由{personal("junior",2007):.2f}%降为{personal("junior",2013):.2f}%；2013至2023年又升为{personal("primary",2023):.2f}%与{personal("junior",2023):.2f}%。后期合计增长与个人支出占比回升并存。')
figure(4)
para('固定28个地区后，个人与公用合计的年均实际增长率中位数如下。数值由各地区分别计算后取中位数，不是全国总额的加权增长率。')
table(['比较时期','小学','初中'],[[f'{a}至{b}年',f'{median("primary",a):.2f}%',f'{median("junior",a):.2f}%'] for a,b in [(1995,2006),(2007,2013),(2013,2023)]],[7.5,4.5,4.5])
para('个人与公用合计不含另列基本建设支出。构成变化不能单独证明公用支出受到挤压，也不足以评价教师工资水平或办学效率。',small=True)

page();heading('后期增长放缓具有较广的地区分布',2)
counts={s:int((PAIR[PAIR.stage==s].change_annualized_percentage_points<0).sum()) for s in names}
assert counts=={'primary':28,'junior':26}
para(f'比较2007至2013年与2013至2023年，同一批地区中，小学{counts["primary"]}个、初中{counts["junior"]}个地区的合计年均实际增速放缓。增速较低仍可能伴随金额增长。',small=True)
figure(5,15.0);figure(6,15.0)

page();heading('全国增长中已经出现局部公用支出收缩',2)
declines={s:int((REG[REG.stage==s].C_real_change_pct<0).sum()) for s in names}
assert declines=={'primary':8,'junior':7}
para('2017至2023年，小学8个、初中7个地区公用支出实际金额下降。图7呈现各地区两个学段的差异；全国增长不能替代对局部运行与资本性支出的核查。',small=True)
figure(7,16.0)

page();heading('全国收入占比受到举办者权重变化影响',2)
dp=decomposition('primary',2014);dj=decomposition('junior',2014)
para(f'2014至2021年，小学和初中财政性收入占比分别下降{-dp.observed_change_pp:.2f}、{-dj.observed_change_pp:.2f}个百分点。其中，举办者收入权重变化的贡献分别为{dp.composition_pp:.2f}、{dj.composition_pp:.2f}个百分点。这解释了全国比例变化的重要构成，但属于算术分解，不能解释为政策净效应。')
figure(8)
para('2021至2024年，民办收入权重回落对应的向上贡献，被各类学校内部财政性收入比例下降所抵消。初中全国比例近乎不变，仍包含方向相反的内部变化。组内变化也不意味着追踪了同一批学校。')
heading('三 结论与政策含义')
para('义务教育经费的长期演进表现为生源转折错位、财政性资金扩张、举办者收入构成变化，以及支出增长放缓和局部项目收缩。评价保障状况应同时报告金额与占比，并区分学段、举办者和支出项目。')
para('经费配置宜结合分学段需求预测安排调整过渡，重点核查公用支出累计实际下降的地区及其项目构成。学校规模缩小仍可能保留人员、校舍和基本运行成本；现有资料能够指出核查重点，尚不能确定最优投入、教育质量或资源充足程度。')

page();heading('四 方法补充与资料来源')
heading('1994年报告采用值',2)
para('1994年采用当年全国经费来源构成及义务教育预算拨款占比53.41%进行换算，借用1995年学段分配比例，并将特殊教育保留在预算分配分母中。最终采用值如下，全部标记为暂估。[4][6]')
table(['学段','财政性支出 亿元','总经费 亿元','财政性占比'],[[names[s],number(finance(s,1994).fiscal_yi),number(finance(s,1994).total_yi),f'{finance(s,1994).fiscal_share_pct:.2f}%'] for s in names],[3,4.5,4.5,4.5])
para('同一方案改用1996年、1997年分配比例，小学财政性情景范围为397.60至399.27亿元，初中为257.35至262.49亿元。范围只表示分配假设敏感性，不是置信区间。官方白皮书载小学总经费594亿元，与本方案501.95亿元存在未解释的范围差异，因此1994年不用于精确同比或转折判断。[7]',small=True)
heading('计算与比较边界',2)
para('实际累计增长率为“期末金额÷期初金额÷累计物价系数−1”；年均增长率为前述金额比值的“1÷间隔年数”次方再减1。累计物价系数由期间逐年CPI连乘得到。固定28地区排除四川、重庆和西藏，分别涉及区划变化、独立起点缺失及早期CPI链缺失。')
para('举办者分解设收入权重为w、组内财政性比例为f：全国比例变化＝ΣΔw×平均f＋ΣΔf×平均w。两部分分别是权重与组内比例贡献。2006年单位经全国合计交叉核实为千元；1998、2002、2007年收入指标及2007、2014、2024年支出分类边界分别保留。CPI并非教育专属成本指数，不将跨期共变解释为因果效应。[5]',small=True)
heading('资料来源',2)
refs=[
 '[1] 中国教育统计年鉴，按1995至2024年实际数据年整理。用于全国与省级普通小学、初中在校生人数。',
 '[2] 中国教育经费统计年鉴，1996至2025年卷。用于经费来源、举办者收入及地方隶属学校支出；卷年与数据年分别登记。',
 '[3] 中国统计年鉴及地区物价表。用于省级居民消费价格指数。',
 '[4] 国家统计局，中国统计年鉴2001，20-36表。1994年全国来源锚点。https://www.stats.gov.cn/sj/ndsj/2001c/t2036c.htm',
 '[5] 国家统计局，中国统计年鉴2012，20-39表。2006年单位交叉核实。https://www.stats.gov.cn/sj/ndsj/2012/html/U2039C.HTM',
 '[6] 中国教育经费投入的“八五”回顾与“九五”展望，1996年年度报告部分内容，表1-4。https://www.cnsaes.org.cn/homepage/Upfile/2010612/2010061208642783.pdf',
 '[7] 国务院新闻办公室，中国的儿童状况，1996年，外交部转载。用于核查1994年小学总经费594亿元记载。https://www.fmprc.gov.cn/ziliao_674904/zt_674979/ywzt_675099/wzzt_675579/2296_675789/200207/t20020705_7961975.shtml',
]
for ref in refs:para(ref,small=True)

doc.core_properties.title=title
doc.core_properties.subject='1995至2024年普通小学与初中生源及教育经费长期分析'
doc.core_properties.author='';doc.core_properties.last_modified_by=''
target=OUT/'义务教育经费长期演进分析报告.docx'
doc.save(target)
(OUT/'报告正文.md').write_text('\n'.join(md),encoding='utf8')
inputs=[DATA/'早期财政重建_20261011/national_fiscal_1994_2024_report_final.csv']+[FIG/'figure_data'/f'fig{n:02d}.csv' for n in range(1,9)]
receipt={'created_on':'2026-10-11','report_data_scope':'1995-2024 main; 1994 conditional supplement','selected_1994_method':'compulsory_budget_constraint','planned_pages':9,'figures':8,'tables':4,'regional_public_decline_counts':declines,'regional_slowdown_counts':counts,'inputs':{str(p.relative_to(ROOT)).replace('\\','/'):hashlib.sha256(p.read_bytes()).hexdigest() for p in inputs},'docx_sha256':hashlib.sha256(target.read_bytes()).hexdigest(),'render_verified':False}
(OUT/'报告制作记录.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf8')
print('Created:',target)
