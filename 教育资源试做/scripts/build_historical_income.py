"""Extract original national income tables for a segmented historical account."""
from pathlib import Path
import zipfile,io,re,sys,json,hashlib
import openpyxl
import pandas as pd
ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'教育资源试做/data/历史趋势描述_20261010'
sys.path.insert(0,str(ROOT/'教育资源试做/.tmp/yearbook_read_deps'))
if hasattr(sys.stdout,'reconfigure'):sys.stdout.reconfigure(encoding='utf8')
def clean(x):return re.sub(r'\s+','',str(x or ''))
def decoded(i):
    n=i.filename
    if not i.flag_bits&0x800:
        try:n=n.encode('cp437').decode('gbk')
        except UnicodeError:pass
    return n
def rows_from(b):
    if b[:2]==b'PK':
        w=openpyxl.load_workbook(io.BytesIO(b),read_only=True,data_only=True)
        rows=list(w.worksheets[0].iter_rows(values_only=True));w.close();return rows
    import xlrd
    s=xlrd.open_workbook(file_contents=b).sheet_by_index(0)
    return [s.row_values(k) for k in range(s.nrows)]
def num(x):
    s=str(x).strip().replace(',','') if x is not None else ''
    return float(s) if re.fullmatch(r'[+-]?\d+(?:\.\d+)?',s) else None
def stage(x):
    s=clean(x)
    if '普通小学' in s:return 'primary'
    if ('初中' in s or '初级中学' in s) and not any(q in s for q in ['职业','成人']):return 'junior'
    return None
def main():
    OUT.mkdir(exist_ok=True)
    evidence=[];issues=[];records=[]
    for volume in range(1996,2026):
        paths=list((ROOT/'中国教育经费统计年鉴').glob(f'*{volume}*.zip'))
        if paths:
            z=zipfile.ZipFile(paths[0]);entries=[(decoded(i),z.read(i)) for i in z.infolist() if not i.is_dir() and re.match(r'^1[-—_ ]',decoded(i).replace('\\','/').split('/')[-1]) and '收入' in decoded(i)]
            archive=str(paths[0].relative_to(ROOT))
        else:
            folder=ROOT/f'教育资源试做/用户提供资料/年鉴解压/卷{volume}'
            entries=[(str(p.relative_to(ROOT)),p.read_bytes()) for p in folder.rglob('*.xlsx') if re.match(r'^1[-—_ ]',p.name) and '收入' in p.name]
            archive=f'中国教育经费统计年鉴/中国教育经费统计年鉴{volume}（excel）.rar; existing extracted files'
        selected=[]
        for name,b in entries:
            base=name.replace('\\','/').split('/')[-1]
            if '收入' not in base or any(w in base for w in ['中央','地方','企业','民办','部门','构成','生均','支出','收支']):continue
            rows=rows_from(b)
            sr=[(k,stage(r[0]),r) for k,r in enumerate(rows) if r and stage(r[0])]
            if not sr:continue
            selected.append((name,b,rows,sr))
        if len(selected)!=1:
            issues.append(dict(year=volume-1,issue='table_selection_count',count=len(selected),tables=[q[0] for q in selected]));continue
        name,b,rows,sr=selected[0]
        first=next(k for k,r in enumerate(rows) if len(r)>1 and num(r[1]) is not None)
        headers=rows[:first]
        ncols=max(map(len,rows));labels=['|'.join(clean(r[c]) for r in headers if c<len(r) and clean(r[c])) for c in range(ncols)]
        unit='千元' if any('千元' in clean(v) for r in headers for v in r) else 'unit_pending'
        sid=hashlib.sha256((archive+'|'+name).encode()).hexdigest()[:24]
        e=dict(source_id=sid,volume_year=volume,data_year=volume-1,archive=archive,table=name,table_sha256=hashlib.sha256(b).hexdigest(),unit=unit,headers=headers,column_labels=labels,stage_rows=[dict(row=k+1,stage=s,values=r) for k,s,r in sr],scope='national all affiliations original income; yearly school definitions pending')
        evidence.append(e)
        def find(pattern):
            return next((c for c,v in enumerate(labels) if re.search(pattern,v.replace('|',''))),None)
        fiscal=find(r'^国家.*财政性.*教育经费')
        budget=find(r'财政预算内拨款|预算内.*教育.*经费|预算内事业性经费|公共财政.*预算|一般公共预算安排')
        fund=find(r'政府性.*基金预算')
        tuition=find(r'#.*学费|#.*学杂费|学费、杂费|学费,杂费|学杂费')
        business=10 if 1998<=volume-1<=2004 else find(r'事业收入')
        donation=find(r'捐.*(经费|收入)|捐集资')
        organizer=find(r'民办学校.*举办者|社会团体.*公民个人')
        other=find(r'^其他(收入|教育经费)$')
        for k,s,r in sr:
            if len([q for q in sr if q[1]==s])!=1:
                issues.append(dict(year=volume-1,stage=s,issue='multiple_stage_rows'));continue
            record=dict(year=volume-1,stage=s,source_id=sid,source_row=k+1,original_stage_label=r[0],unit=unit,income_total_thousand_yuan=num(r[1]),school_scope=e['scope'],annual_comparability='pending; segmented original category description only',per_student_calculated=False)
            for metric,col in [('national_fiscal',fiscal),('budget',budget),('government_fund',fund),('tuition',tuition),('business',business),('donation',donation),('organizer',organizer),('other',other)]:
                record[metric+'_thousand_yuan']=num(r[col]) if col is not None and col<len(r) else None
                record[metric+'_source_cell']=f'{openpyxl.utils.get_column_letter(col+1)}{k+1}' if col is not None else ''
                record[metric+'_original_header']=labels[col] if col is not None else ''
            y=volume-1
            components=[2,5,8,9,10,11,13] if y<=1997 else [2,6,10,12,13,15,16] if y<=2004 else [2,4,5,6,8] if y<=2006 else [fiscal,organizer,donation,business,other]
            assert all(c is not None for c in components),(y,labels)
            values=[num(r[c]) for c in components]
            record['top_level_component_columns']=';'.join(openpyxl.utils.get_column_letter(c+1) for c in components)
            record['component_original_blanks']=sum(v is None for v in values)
            record['total_minus_observed_components']=record['income_total_thousand_yuan']-sum(v for v in values if v is not None)
            record['component_check']='observed_sum_matches_with_original_blanks_retained' if abs(record['total_minus_observed_components'])<=8 else 'pending_difference'
            records.append(record)
    (OUT/'national_income_original_evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf8')
    (OUT/'income_extraction_issues.json').write_text(json.dumps(issues,ensure_ascii=False,indent=2),encoding='utf8')
    f=pd.DataFrame(records).sort_values(['stage','year'])
    for metric in ['national_fiscal','budget','government_fund','tuition','business','donation','organizer','other']:
        f[metric+'_share_pct']=100*f[metric+'_thousand_yuan']/f.income_total_thousand_yuan
    f.to_csv(OUT/'national_income_original_series.csv',index=False,encoding='utf-8-sig')
    print('Records',len(f),'Evidence',len(evidence),'Issues',issues)
    print(f[['year','stage','unit','income_total_thousand_yuan','national_fiscal_share_pct','budget_share_pct','tuition_share_pct']].to_string(index=False))
    print('BUDGET HEADERS',f.groupby('budget_original_header').year.agg(['min','max','count']).to_string())
if __name__=='__main__':main()
