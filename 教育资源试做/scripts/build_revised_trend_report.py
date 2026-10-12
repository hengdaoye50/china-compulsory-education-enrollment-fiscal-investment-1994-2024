"""Produce the formal trend report from verified, immutable research outputs."""
from pathlib import Path
import json,re,hashlib
import pandas as pd
from PIL import Image
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[2];BASE=ROOT/'教育资源试做';DATA=BASE/'data'
OUT=ROOT/'义务教育趋势分析报告_审校修订版_20261012';OUT.mkdir(exist_ok=True)
NAME='中国义务教育生源规模与财政投入趋势分析报告（1994—2024）'
old=ROOT/'最终综合研究报告_20261011'
oldmeta=json.loads((old/'报告合稿记录.json').read_text(encoding='utf8'))
protected={p:hashlib.sha256((ROOT/p).read_bytes()).hexdigest() for p in oldmeta['input_sha256']}
text=(old/'最终报告正文.md').read_text(encoding='utf8')
sections={}
for part in re.split(r'^## ',text,flags=re.M)[1:]:
    title,body=part.split('\n',1);sections[title]=body.strip()
def oldtable(section,k=0):
    b=sections[section];tables=re.findall(r'表\d+[^\n]*\n\n(\|[^\n]*\n\|[^\n]*\n(?:\|[^\n]*(?:\n|$))+)',b)
    lines=tables[k].strip().splitlines()
    values=[[v.strip() for v in line.strip('|').split('|')] for line in lines]
    return values[0],values[2:]
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.78);sec.bottom_margin=Inches(.72);sec.left_margin=sec.right_margin=Inches(.82);sec.header_distance=Inches(.33);sec.footer_distance=Inches(.33);sec.different_first_page_header_footer=True
for name,size,font,bold in [('Normal',11.5,'宋体',False),('Title',23,'黑体',True),('Subtitle',14,'宋体',False),('Heading 1',16,'黑体',True),('Heading 2',12.5,'黑体',True),('Caption',10,'宋体',False)]:
    st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.bold=bold;st.font.italic=False;st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),font)
    pf=st.paragraph_format;pf.space_after=Pt(7);pf.line_spacing=1.35;pf.widow_control=True
    if name.startswith('Heading'):pf.keep_with_next=True;pf.space_before=Pt(12)
    if name=='Normal':pf.first_line_indent=Pt(23)
    if name=='Caption':pf.first_line_indent=Pt(0);pf.line_spacing=1.15
for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
hp=sec.header.paragraphs[0];hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=hp.add_run('中国义务教育生源规模与财政投入趋势分析报告');r.font.size=Pt(8);r.font.color.rgb=RGBColor(0,0,0)
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
fp.add_run('— ').font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld);fp.add_run(' —').font.size=Pt(9)
doc.settings.element.append(OxmlElement('w:doNotAutoCompressPictures'))
blocks=[];figures=[];tables=[];headings=[]
def para(t,small=False):
    p=doc.add_paragraph(t);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if small:
        p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.15
        for r in p.runs:r.font.size=Pt(9)
    blocks.append(dict(type='paragraph',text=t));return p
def h(t,level=2,newpage=False):
    p=doc.add_heading(t,level)
    if newpage:p.paragraph_format.page_break_before=True
    headings.append(t);blocks.append(dict(type='heading',level=level,text=t));return p
def tab(title,headers,rows,note=None):
    n=len(tables)+1;cap=f'表{n}  {title}';p=doc.add_paragraph(cap,'Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
    tb=doc.add_table(rows=1,cols=len(headers));tb.autofit=False
    weights=[1]*len(headers)
    if len(headers)==3:weights=[1.1,2.4,2.3]
    if len(headers)==5:weights=[.8,1.7,1.1,2.1,1.1]
    for rownum,vals in enumerate([headers]+rows):
        rr=tb.rows[0] if rownum==0 else tb.add_row()
        pr=rr._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if rownum==0:pr.append(OxmlElement('w:tblHeader'))
        for j,(cell,val) in enumerate(zip(rr.cells,vals)):
            cell.text=str(val);cell.width=Inches(6.86*weights[j]/sum(weights));cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            cp.append(borders);mar=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');mar.append(el)
            cp.append(mar)
            if rownum==0:
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E8E8E8');cp.append(fill)
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next=len(rows)<=6 and rownum<len(rows);p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.15;p.alignment=WD_ALIGN_PARAGRAPH.CENTER if len(str(val))<30 else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:r.font.size=Pt(9.5);r.font.bold=rownum==0
    if note:
        para('注：'+note,True);blocks.pop()
    tables.append(dict(number=n,title=title,headers=headers,rows=rows));blocks.append(dict(type='table',**tables[-1]))
    if note:blocks.append(dict(type='paragraph',text='注：'+note))
def fig(path,title,note,width=6.55):
    path=Path(path);replacement=OUT/'图表精修'/path.name
    if replacement.exists():path=replacement
    assert path.exists(),path
    n=len(figures)+1;p=doc.add_paragraph(f'图{n}  {title}','Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
    ip=doc.add_paragraph();ip.alignment=WD_ALIGN_PARAGRAPH.CENTER;ip.paragraph_format.first_line_indent=Pt(0);ip.paragraph_format.space_after=Pt(4);ip.paragraph_format.keep_with_next=True
    im=Image.open(path);w=min(width,6.65,6.15/(im.height/im.width));ip.add_run().add_picture(str(path),width=Inches(w))
    np=para('注：'+note,True);np.paragraph_format.space_after=Pt(9);blocks.pop()
    item=dict(number=n,title=title,path=str(path.relative_to(ROOT)),caption=note,sha256=hashlib.sha256(path.read_bytes()).hexdigest());figures.append(item);blocks.append(dict(type='figure',**item))
def oldfig(n,title,note,width=6.55):fig(ROOT/oldmeta['figure_map'][n-1]['path'],title,note,width)
def F(name,title,note):fig(BASE/'支出细项与补助图组_20261011'/f'{name}.png',title,note)
# Cover: no institutional authorship is invented.
p=doc.add_paragraph();p.paragraph_format.space_after=Pt(90)
p=doc.add_paragraph('中国义务教育生源规模与财政投入\n趋势分析报告','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1.7
p=doc.add_paragraph('（1994—2024）','Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(70)
p=doc.add_paragraph('全国趋势　区域差异　支出结构');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
p=doc.add_paragraph('2026年10月');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(80)
blocks.append(dict(type='title',text=NAME))
h('摘要',1,True)
para('本报告考察一个贯穿长期变化的问题：在生源规模多轮变化的过程中，义务教育经费扩张在学段、地区和支出项目之间呈现怎样的不同步。财政性收入用于观察投入来源，全来源支出用于分析学校经费使用的构成，二者分别解释。主体观察期为1995—2024年，1994年经费条件估算作为补充基期。分析以普通小学、初中为对象，结合全国年度资料与省级面板，依次开展趋势描述、构成分解和支出项目响应比较。')
para('生源规模具有反复变化和学段错位特征。小学在1997年达到观察期峰值，2013年降至低点；初中的对应年份为2003年和2015年。财政性经费名义金额总体扩张，财政性收入占比自2013年高点回落，两种变化同时发生。全国比例还受到不同举办者收入权重的影响，不能单独作为财政保障强弱的判据。')
para('支出扩张的速度和内部结构存在明显差异。在固定28个省级地区的比较中，2013—2023年个人与公用支出合计的实际增速低于2007—2013年的地区，小学为28个、初中为26个。2017—2023年，小学8个、初中7个地区公用支出累计实际下降。支出增长放缓与局部项目收缩并存，要求在全国总量之外监测地区和项目变化。')
para('2008—2023年省级面板分析显示，商品服务支出相对工资福利支出的生源响应差，小学为0.766、初中为0.506，均通过相应检验族的多重检验校正。该结果表明，在控制自身财政收入和省份、年份效应后，两类项目与区域生源变化的关联不同。教师、班级和中央补助的拓展分析尚未提供独立机制证据，初中后期估计也不足以确认全窗口关系持续稳定。')
para('三十年资料表明，生源变化并非单向收缩，经费扩张也并非各类项目同步增长。全国总额增长与地方、项目调整可以并存，需求与投入的关系应分学段、分地区和分项目观察。财政保障监测应同时考察金额、来源占比、运行相关项目和资本性项目，并纳入中央补助及其用途约束。报告结论属于趋势描述、会计构成解释和探索性条件关联，不用于认定改革净效应或评价学校支出效率。')
para('关键词：义务教育；生源规模；财政投入；区域差异；支出结构',True)
h('目录',1,True)
chapter_titles=['一  分析框架与资料范围','二  生源规模的长期变化','三  财政性经费与收入构成','四  支出扩张的区域差异','五  支出项目的结构调整','六  生源变化与支出响应','七  办学资源与财政补助','八  主要结论与政策启示','附录一  指标定义与统计范围','附录二  估计方法与稳健性分析','附录三  早期经费与跨期衔接','附录四  分省结果与资料管理','资料来源与参考文献']
pagination=json.loads((OUT/'目录页码.json').read_text(encoding='utf8')) if (OUT/'目录页码.json').exists() else {}
for title in chapter_titles:
    p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(12);p.paragraph_format.tab_stops.add_tab_stop(Inches(6.3),WD_ALIGN_PARAGRAPH.RIGHT)
    p.add_run(title+'\t'+str(pagination.get(title,'')))
h(chapter_titles[0],1,True)
h('（一）从总量变化转向配置结构')
para('生源变化并不会按同一比例转化为各类教育支出的变化。学生规模影响运行需求，学校、人员与项目安排又可能延续存量。本报告考察：生源多轮变化过程中，经费扩张在学段、地区和支出项目之间呈现怎样的不同步？')
para('现有研究关注师资需求、人员配置和小规模学校调整。[5—8] 本报告进一步细化到支出项目，连接全国来源构成、地区趋势和项目关联。其贡献在于长期资料及账面结构的综合解释，不以省级相关性替代学校层面的机制证据。')
h('（二）资金口径与分析层次')
para('资料来自教育、教育经费、财政和综合统计年鉴，按实际数据年整理。[1—4] 全国趋势与内地31个地区的比较分别使用全国行、分省行。历史支出采用地方隶属口径，回归采用全隶属口径；固定地区比较包括28个地区。')
para('财政性收入用于考察财政性金额及其占同范围总收入的比例。全来源支出用于观察个人、公用及内部项目，包含不同来源形成的学校支出，不能单独归因于财政资金。中央补助为财政决算总项，并非义务教育专项补助。')
para('分析依次采用趋势描述、构成分解和控制省份、年份效应的条件关联模型；实际增长以地区CPI调整。1994年属于条件估算，未进入回归。生源与学校财务覆盖尚未完全对应，不计算严格同范围的生均支出。完整规则和样本列于附录。')
h('（三）支出调整的制度背景')
para('2015年义务教育经费保障文件涉及小规模学校保底及寄宿、特殊教育等安排，2018年教育经费使用文件兼顾投入和教师待遇保障。[16—17] 学校需求因而不能只按省级学生总量等比例理解。2019年财政事权方案区分中央与地方分担，支出规模与本级财力负担亦应分别观察。[18]')
para('工资与养老列支、项目完结和采购时点可能改变账面构成。这些制度与组织条件为解释项目差异提供背景，当前资料未分离其作用，不能据政策规定认定实际执行效果。')
h(chapter_titles[1],1,True)
h('（一）小学与初中的峰谷不同步')
para('小学在校生于1997年达到观察期峰值，约13,995万人，2013年降至约9,361万人后回升；初中峰值为2003年的约6,618万人，低点为2015年的约4,312万人。两学段峰值相差六年、低点相差两年（图1），需求变化呈现学段时序差异。')
fig(OUT/'图表/全国生源长期趋势.png','全国小学与初中在校生规模（1995—2024年）','来源：《中国教育统计年鉴》。两面板采用同一纵轴；标记为观察期峰谷及期末值。在校生变化同时受人口、流动、入学及统计范围影响。',6.4)
s=pd.read_csv(BASE/'正式学术图组_20261011/figure_data/fig01.csv')
tab('2024年在校生规模及年度变化',['学段','在校生（万人）','较2023年变化'],[[l,f'{s[(s.stage==st)&(s.year==2024)].students.iloc[0]/10000:,.2f}',f'{s[(s.stage==st)&(s.year==2024)].student_change_pct.iloc[0]:+.2f}%'] for st,l in [('primary','小学'),('junior','初中')]])
h('（二）区域收缩具有反复性')
para('以连续至少两年在校生减少识别下降轮次，观察期内共记录125轮，其中小学63轮、初中62轮；小学21个、初中18个地区出现至少两轮。三年阈值下，两学段仍分别有15个、11个地区反复下降。生源收缩发生在不同时期和地区，资源规划应结合学段队列和区域流动，避免直接从全国总量推导各地配置需求。区划及观察期截断规则见附录。')
h(chapter_titles[2],1,True)
h('（一）财政性金额扩张与占比回落并存')
para('财政性经费名义金额总体扩张。在2005—2024年公布值窗口内，财政性收入占比于2013年达到高点：小学96.11%、初中94.68%；2024年分别为91.36%、89.37%。2013—2024年，财政性收入名义金额分别增长107.31%、111.27%。图2将金额与占比分开呈现，表2列示期末水平。')
fig(OUT/'图表/财政性经费长期趋势.png','全国小学与初中财政性教育经费及其占比（1994—2024年）','来源：《中国教育经费统计年鉴》及早期来源组成表。1994年为空心估算点，1998、2002年处分段；2007年来源组成调整以竖线标示，跨该年连接仅作历史描述。金额为名义值；下图纵轴为65%—100%。估算与衔接详见附录三。',6.4)
fin=pd.read_csv(DATA/'早期财政重建_20261011/national_fiscal_1994_2024_report_final.csv')
tab('2024年财政性教育经费收入',['学段','财政性金额（亿元）','占同范围总收入'],[[l,f'{fin[(fin.stage==st)&(fin.year==2024)].fiscal_yi.iloc[0]:,.2f}',f'{fin[(fin.stage==st)&(fin.year==2024)].fiscal_share_pct.iloc[0]:.2f}%'] for st,l in [('primary','小学'),('junior','初中')]])
para('财政性收入占比下降，表示财政性与其他来源收入的相对关系变化。其保障含义需要结合金额和学校类别判断。1995—1997年支出与后期收入、2002年企业来源定义分别处理；2007年民办来源项目改名并未使总收入只纳入举办者投入，组成核算及跨期边界详见附录三。')
h('（二）举办者构成改变全国比例')
para('2007—2024年，民办学校收入占全国同学段收入的权重于2021年达到窗口峰值，小学7.87%、初中11.21%；2024年分别回落至6.62%、9.20%。这些比例反映收入构成，并不等于民办学生人数占比。')
fig(OUT/'图表/举办者收入构成分解.png','财政性收入占比变化的举办者构成分解','来源：年鉴举办者分类收入表。紫色为举办者收入权重变化贡献，金色为类别内部财政性比例变化贡献，黑色为全国占比变化；单位为百分点。采用中点恒等式分解。',6.4)
para('2014—2021年，两学段全国财政性比例下降主要对应举办者收入权重变化。2021—2024年，权重回落的向上贡献受到组内比例下降的抵消，初中全国比例仅增加约0.006个百分点（图3）。接近不变的全国比例，可以包含方向相反的结构调整。')
tab('财政性收入占比的举办者分解（百分点）',*oldtable('六 举办者构成怎样改变全国比例'),note='分解为会计构成解释，不表示政策净效应；不同举办者收入权重不等于学生权重。')
h(chapter_titles[3],1,True)
h('（一）实际增速放缓具有广泛区域分布')
para('固定28个省级地区的端点比较显示，1995—2006年个人与公用支出合计的年均实际增长率中位数为11.70%（小学）、12.11%（初中）；2007—2013年升至14.54%、14.47%；2013—2023年为6.16%、6.05%。各地区分别计算后取中位数，这一指标不等于全国金额加权增速。')
oldfig(5,'相同28个省级地区的支出实际增速跨期比较','来源：地方隶属学校支出与省级CPI。每点对应同一地区，虚线表示两个窗口增速相同；两轴同尺度。四川、重庆和西藏未纳入，原因见附录一。')
para('逐省配对后，小学28个、初中26个地区后期增速低于前期（图4）。固定地区比较表明，在相同观察对象下仍有广泛的增速放缓。固定样本控制了地区构成变化，跨期分类边界则单独保留；该比较属于原表端点的描述，不用于识别制度转折。较低增速仍可能对应金额增长，应与实际下降区别判断。')
h('（二）公用支出收缩呈现空间差异')
para('2017—2023年，小学8个、初中7个地区的公用支出累计实际下降。小学吉林、辽宁、黑龙江的个人与公用合计也实际下降，初中两项合计在31个地区均增长。图5、图6采用相同分组阈值显示空间分布，精确分省幅度见附录四。')
fig(OUT/'图表/小学公用支出空间分布.png','小学公用支出累计实际增长的省级分布（2017—2023年）','地方隶属学校；金色为负增长，紫色色阶表示正增长幅度。使用省级CPI调整；公用包含资本性项目。数据来源：教育经费与综合统计年鉴；边界来源：DataV.GeoAtlas，原始边界保留，详见附录四。')
fig(OUT/'图表/初中公用支出空间分布.png','初中公用支出累计实际增长的省级分布（2017—2023年）','统计范围、分组阈值与图5相同。两图表示期间端点变化，不代表每年均持续增长或下降；港澳台未纳入数据分析。')
para('空间分布揭示全国扩张与局部收缩可以并存。公用金额下降仍需区分资本项目完结、采购时点、学校类型变化和运行需求调整；端点结果不能直接证明教育资源不足或支出效率变化。其政策价值在于识别需要进一步核实的地区和项目。')
h(chapter_titles[4],1,True)
h('（一）个人与公用支出的相对份额变化')
oldfig(4,'地方隶属学校个人支出占两项合计的比例（1995—2024年）','分母为个人与公用合计，不含另列基本建设。个人不等于教师工资，公用包含资本性项目；分类边界分段显示。')
para('2007—2013年，个人支出占两项合计的比例，小学由74.94%降至65.19%，初中由68.41%降至61.67%；2013—2023年又升至75.40%、72.89%（图7）。这一变化与公用前期较快扩张、后期相对放缓一致。份额回升本身不能证明日常运行经费受挤压。')
h('（二）公用内部的运行相关与资本性项目')
F('图F1_支出内部构成','个人及公用支出的内部构成（2007—2023年）','内地31个省级地区名义金额合计，不是官方全国合计序列。左列分解个人支出，右列分解公用支出；四项分别为工资福利、个人家庭补助、商品服务和其他资本性。')
para('2023年，工资福利占个人支出约89.7%（小学）、90.3%（初中），商品服务占公用支出约58.5%、53.4%。公用中仍有相当份额属于其他资本性项目（图8）。工资福利覆盖的人员范围宽于专任教师，商品服务也不等于全部教学运行成本。')
F('图F2_公用增长贡献','商品服务与其他资本性项目的实际增长贡献','分项实际金额变化除以同省前期公用金额，再取年度可用省年算术均值。两柱合计对应黑线，单位为百分点；不是全国加权增长。')
para('2009—2012年，两学段其他资本性的平均增长贡献普遍高于商品服务；部分后期年份则出现商品服务正贡献与其他资本性负贡献并存（图9）。公用总额能够掩盖内部抵消，增长放缓可能涉及项目组合变化。金额贡献描述账面组成，不能据此判断资本投资效益或排除个别学校的运行压力。')
h(chapter_titles[5],1,True)
h('（一）个人与公用具有不同条件响应')
para('这里的响应差，是在控制相同变量后，比较生源变化相同时两类支出增长关联的差异。例如公用减个人的系数为正，表示公用增长与生源增长的平均关联更强，不能据此直接断言减少学生会导致运行经费受挤压。')
para('在2008—2023年每学段472条完整省年记录中，控制学生、实际自身财政收入及省份、年份效应，直接比较公用与个人支出的响应差。初中学生响应差为0.601，四项检验Holm校正p为0.0144；小学为0.695，校正p为0.0879。自身收入响应差在两个学段均未获相应校正检验支持。')
fig(OUT/'图表/个人公用直接响应差.png','公用支出相对个人支出的响应系数差','点为直接差模型系数，线为单项野生聚类自助检验反演95%区间；四项p值按检验族作Holm校正。每学段472条省年、31个省级聚类。区间针对单项检验，p值经过多重检验校正，两者采用不同评价标准。')
para('直接差异检验比比较两个模型各自是否显著更符合研究问题。结果支持初中不同支出类别的平均条件关联存在差异；小学方向一致，但证据较弱。将学生增长和下降分别建模后，初中下降项差系数约−0.462，近似区间跨零，尚不足以专门认定生源收缩引起公用挤压。')
h('（二）商品服务相对工资福利的差异更清晰')
F('图F3_商品服务与工资响应差','商品服务相对工资福利的响应系数差','与个人、公用比较使用相同的472条省年。因变量为两项实际对数变化之差；线为单项自助反演95%区间，四项检验作族内Holm校正。')
para('商品服务相对工资福利的学生响应差，小学为0.766、初中为0.506，相应校正p为0.0008、0.0033（图11）。在模型控制下，生源增长幅度相差约1%，两项支出的增长差平均相差约0.77和0.51个百分点。剥离其他资本性和个人家庭补助后，项目差异仍然存在，宽泛公用与个人的关系因而不能只归因于资本项目。')
para('辅助模型中，商品服务的学生系数为0.589、0.444，工资福利为−0.177、−0.062；后者区间均含零。结果不能解释为生源减少必然推高工资，也不直接证明人员刚性机制。更准确的判断是，不同项目与区域生源变化的平均关联不同。')
h('（三）平均关联具有时期边界')
F('图F7_时期分段响应','商品服务相对工资福利响应差的分时期比较','全窗口472条，2008—2016年258条，2019—2023年153条。各窗口独立估计；线统一为CR1及t分布近似95%区间。分期属于追加敏感性分析，非正式结构突变检验。')
para('剔除2017—2018年后，学生响应差仍为0.816（小学）、0.450（初中）。按时期估计，前期为1.011、0.502，后期为1.128、−0.207。初中后期95%近似区间为−1.339至0.925，估计精度不足（图12）。全窗口关系不能推广为每一阶段均稳定，也不能根据后期负点估计认定关系逆转。')
h(chapter_titles[6],1,True)
h('（一）教师与班级变化的解释范围')
para('同样本加入教师和班级后，学生响应差有所变化，但系数变化不能换算为机制解释率；完整比较见附录二图14。')
para('四项资源直接差异检验均未通过校正后的5%阈值。初中教师差为−0.836、校正p为0.1660，班级差为0.800、p为0.6387（附录二图15）。班级组织和人员存量是有理论依据的解释方向，但当前省级资料尚未识别独立的人员制度或班级调整机制。')
h('（二）自身收入不能完整代表财政背景')
F('图F5_自身收入与补助合计变化','地方自身收入与中央补助合计的实际变化','436个省年，两学段财政背景仅计一次；金色为自身收入实际下降而两项合计未下降的记录。两项合计不是完整可用财力。')
para('在补助共同样本中，84条自身收入实际下降记录有60条的自身收入与中央补助合计未下降，占71.4%（图13）。单以自身收入下降表示财政资源同步收缩，会遗漏部分年度的预算背景。中央补助包含返还及用途约束，其金额变化也可能响应地区困难，不能据该比例认定外生补助缓冲效应。补助回归未确认独立关联，详见附录二图16。')
h(chapter_titles[7],1,True)
h('（一）长期趋势的综合判断')
para('三十年资料没有显示一条从生源变化到经费变化的统一路径。生源先后经历不同学段的峰值、回落和回升，财政性金额总体扩张，学校来源与支出结构则持续变化。将全国总额视为各地需求与供给同步改善的证据，会遗漏地区收缩和项目间抵消。报告的核心判断是：长期扩张必须结合学段时序、地区分布和支出构成理解。')
tab('主要序列的历时变化与判断边界',['观察对象','主要历时变化','由资料支持的判断'],[['全国生源','小学1997年峰值、2013年低点；初中2003年峰值、2015年低点，其后均回升','峰谷存在学段错位；2024年小学下降而初中增长'],['财政性收入','2013—2024年名义金额扩张，而占同范围总收入比例回落','金额与来源比例不同步；收入权重变化参与全国比例变化'],['固定地区支出','2007—2013年实际增速较高，2013—2023年较低；2017—2023年部分公用支出累计下降','扩张速度和局部项目方向不同；跨期分类边界保留']])
para('表内分段分别依据生源峰谷、收入占比和支出比较窗口，不表示同一个历史断点。1994年为条件估算补充，不能用于识别真实转折；统计分类边界也不能直接命名为制度或政策效果。')
h('（二）结构差异与条件关联')
para('项目层面的条件关联是定量分析最明确的发现。商品服务相对工资福利与区域生源变化的响应不同，在两学段全窗口均获得相应校正检验支持。资源和中央补助拓展尚未建立独立机制，后期估计也保留时期边界。因此，报告对趋势和项目差异作明确判断，对制度原因保持与证据相称的解释力度。')
h('（三）政策启示',2,True)
tab('预算监测的对象、触发条件与核查任务',['监测对象','证据与核查触发条件','责任主体与核查任务'],[['分学段生源与学校分布','学段增减错位或地区连续下降','教育部门结合队列、流动及城乡学校分布复核配置需求'],['财政性金额和来源比例','金额与比例不同向、举办者权重改变','教育与财政部门并列核对收入金额、类别权重及可比范围'],['运行与资本性项目','公用合计下降或内部正负贡献抵消','教育与财政部门分查项目完结、采购时点、人员列支与运行支出'],['地区累计下降','端点金额实际下降','地方部门核实学校覆盖和项目构成，再判断是否存在运行不足'],['中央补助及用途限制','自身收入下降而收入加补助未下降','财政部门对应省市范围、上解和专项约束核实可用财力']])
para('上述建议面向预算监测和配置决策的信息基础。省级资料尚不足以评价学习结果、教师学科匹配或学校资金充足程度，不能据此直接制定撤并学校、削减教师或最优支出比例方案。完善学校覆盖、工资养老列支和教育专项补助资料，是进一步机制研究的关键。')
h(chapter_titles[8],1,True)
tab('数据系列、统计对象与使用窗口',*oldtable('二 数据来源与比较范围'))
para('个人支出记为P，公用记为C；工资福利W、个人家庭补助A、商品服务G、其他资本性X满足原表显示精度内P≈W＋A、C≈G＋X。另列基本建设K不重复加入X。自身财政收入记为R，中央补助记为B。年度CPI以上年为100，只用于一般购买力近似，非教育投入品专属价格指数。')
para('长期支出描述采用地方隶属学校，回归采用全隶属、全来源支出。学生、教师和班级的区域统计范围尚未与支出对象完全对应。全国财政性收入与全来源支出分别定义，不混用于生均指标。')
para('固定地区长期比较排除四川、重庆和西藏：四川涉及重庆设市前后的区划变化，重庆缺少1995年独立起点，西藏早期CPI链不完整。两学段每个窗口均采用相同28个地区。2017—2023年地区端点比较覆盖内地31个省级地区。')
tab('回归窗口与共同样本',*oldtable('三 指标构造与推断方法'))
h('（一）生源下降轮次的识别与敏感性',2,True)
para('以1995—2024年分地区、分学段在校生序列为基础，年度人数严格低于上一年才记一次下降。连续至少两次年度下降构成一轮，其基期为第一次下降前一年，末期为最后一次下降年。人数持平或增加终止该轮；一地区可有多轮，轮次数不能解释为独立地区数。')
para('缺失相邻年会打断连续性；四川1997、1998年的相邻比较因区划边界排除。1996年开始的下降轮次标记为观察期左边界，持续到2024年的轮次标记为右边界；满足持续期的边界轮次纳入描述，不能认定其真实起止正好落在观察期边界。重庆只从可得独立地区记录识别。')
sp=pd.read_csv(DATA/'三十年纵向分析_20261010/recurrent_student_decline_spells.csv')
rr=[]
for threshold in [2,3]:
    for st,l in [('primary','小学'),('junior','初中')]:
        z=sp[(sp.stage==st)&(sp.decline_years>=threshold)];g=z.groupby('province').size()
        rr.append([l,str(threshold)+'年',str(len(z)),str(int((g>=2).sum())),str(int(z.left_study_edge.sum()))+'/'+str(int(z.right_study_edge.sum()))])
tab('生源下降轮次及持续期敏感性',['学段','最低持续期','轮次数','至少两轮的地区数','左/右边界轮次数'],rr,note='按年度下降次数判断持续期，边界标记不是完整持续时间。配套轮次明细保留基期、末期及区划终止标记。')
h('（二）跨期统计范围及允许用途')
tab('统计边界、处理方式与使用限制',['序列或窗口','已知边界','处理方式与允许用途'],[['早期财政性经费','1994年估算；1995—1997年支出、后期收入；1998年会计边界','估算单点、边界分段；只描述金额与来源，详见附录三'],['企业及民办来源','2002年企业定义变化；2007年财政性组成列项调整','企业组成重算；列示2007年组成差额，不认定全部范围统一'],['地方隶属支出','2007、2014年分类边界及2024年资本项目名称变化','长期比例分段，2024年单点；固定28地区端点比较不作改革效应解释'],['全隶属支出细项','2017—2018年内部构成变化及学校覆盖限制','以条件关联解释；剔除两年及分期估计针对细项响应，不替代长期端点核验'],['生源与财务覆盖','区域在校生与经费学校对象尚未完全对应','生源为区域背景；不报告严格同范围的生均支出'],['财政背景','自身收入和中央补助总项并非完整净可用财力','仅比较指标和条件关联，不认定外生补助效应']])
para('数值核准、组成恒等式与跨期范围分别记录。原始派生文件中的待核状态表示其生成时的边界，后续来源核准不自动认证所有年度可比性。对已核准金额按相应用途使用，对仍有学校覆盖或分类边界的序列保留表内限制。固定地区样本只控制地区构成变化。',True)
h(chapter_titles[9],1,True)
h('（一）增长、分解与回归设定')
para('实际累计增长＝期末与期初名义金额之比÷区间CPI连乘因子−1；复合年均增长再按间隔年数开方。地区中位数、等权省年均值和全国金额加权增长分别使用，不能互相替代。')
para('实际对数变化为Δln X实际＝ln(X当年/X上年)−ln(CPI当年/100)。支出变化模型为Δln Y实际＝βSΔln S＋βRΔln R实际＋α省＋λ年＋ε。直接差模型以Δln C−Δln P或Δln G−Δln W为因变量，解释变量保持相同；资源模型加入教师、班级，补助模型加入Δln B实际或以R＋B替代R。')
para('全国财政性收入比例写为Σwf，w为举办者收入权重、f为组内财政性比例。两时点变化＝ΣΔw×平均f＋ΣΔf×平均w。公用增长贡献以同省前期C为分母计算G、X的实际金额变化。两种分解均属会计恒等式，不作因果解释。')
para('模型采用省年等权OLS，不插值、补零、缩尾或按结果选择主样本。省级CR1聚类标准误允许同省相关。四项直接差异检验采用9999次零假设约束的Rademacher野生聚类自助及检验反演单项95%区间。[9] 各检验族分别作Holm校正，未统一控制全部探索分析的错误率。辅助图使用明确注明的CR1近似区间。')
para('分析为探索性设计，分期检查属于观察图形后的追加敏感性分析；各时段独立估计，未实施正式结构突变检验。固定效应不能消除时变遗漏、反向因果或跨省相关。置信区间用于估计推断，与1994年条件估算情景范围不同。')
for section,title in [('十 从历史共变到支出响应差异','个人与公用直接响应差检验'),('十四 商品服务与工资福利的响应差','商品服务与工资福利直接响应差检验')]:tab(title,*oldtable(section))
tab('财政变量规格变化的同样本结果',*oldtable('十七 补助进入模型后的变化'),note='每学段同一436条省年。括号为省级CR1标准误；模型变化不表示中介解释比例。')
h('（二）早期收缩假设与不利结果')
para('连续至少两年学生减少的样本用于检验两项探索假设。收入收缩变量Z＝max(0,−实际自身收入对数变化)。第一项以公用减个人增长为结果，预期Z系数为负；第二项以公用实际增长为结果，检验2016年及以后收入收缩斜率相对前期是否提高，允许学生响应分阶段。')
tab('持续生源下降样本的两项探索检验',*oldtable('附录一 早期收缩假设及不利结果'))
para('四项检验校正后均未通过5%阈值。小学持续下降样本157条，实际收入收缩17条，其中2016年前仅2条；初中持续下降202条，收入收缩前期5条。小学第二项反演置信集在设定搜索下界未封闭，不能写成有限区间。同省同年两学段共同下降的98条中，财政收缩仅5条、3省，学段直接差系数约−0.00136、自助p为0.9777。')
para('剔除财政分类调整相关的2015—2016年后，初中第一项点估计由−0.767变为−0.105，小学第二项由−0.412变为1.236。原假设缺乏稳定支持，完整结果保留。全样本和支出细项分析改变了研究对象及估计问题，不能视为原假设获验证。')
h('（三）资源与财政补助的完整拓展结果',2,True)
para('以下模型保留探索分析的完整证据。资源模型采用相同300条省年进行比较，补助模型采用相同436条省年。控制变量进入后的系数变化不表示中介解释率；统计不显著也不等于变量没有作用。')
para('初中同样本基准学生响应差为0.473，分别加入教师、班级及两者后为0.566、0.151、0.073；小学共同加入后为0.787（附录二图14）。窗口变化与控制变量进入必须分开比较。系数变化涉及资源共变、内生性和控制路径，不能换算为机制解释率。')
oldfig(8,'纳入教师和班级变量后的学生响应差','原窗口每学段472条，其余模型均使用2014—2023年同一300条省年；点为公用减个人的学生系数，线为CR1及t分布近似95%区间。',6.0)
oldfig(9,'教师和班级与支出增长差的条件关联','共同控制学生、收入、教师、班级及双向固定效应。每学段300条省年；线为单项自助反演95%区间，四项资源检验作Holm校正。',6.0)
F('图F6_中央补助关联','中央补助与教育支出调整的条件关联','控制学生、实际自身收入及双向固定效应；每学段436条省年，9999次约束野生聚类自助检验与反演单项95%区间。四项补助检验单独作Holm校正。')
para('中央补助对商品服务的系数为0.010（小学）、0.038（初中），对商品服务减工资福利的差系数为−0.072、−0.053；四项校正p均为1.000，区间均跨零（附录二图16）。现有结果未确认独立补助关联，不等于补助无作用。在同一436条省年中改变财政变量规格后，学生响应差仍为正，详见附录二。')
h(chapter_titles[10],1,True)
h('（一）1994年估算与早期资料')
para('1994年全国预算内经费883.98亿元乘当年义务教育预算拨款份额53.41%，得到约472.13亿元约束总额。[10][12] 按1995年小学、初中、特殊教育预算内支出比例分配，特殊教育保留在分母；其他来源用1994年全国来源金额与1995年学段来源权重分配，再加总财政性组成。')
tab('1994年财政性经费条件估算',*oldtable('附录二 早期经费重建与1994年估算'))
para('小学财政性金额为399.27亿元、初中257.35亿元，均为估算。改用1996、1997年分配比例得到表内同方案情景范围，不能理解为置信区间或真实误差上限。官方白皮书1994年小学总经费594亿元与本方案总额501.95亿元仍有范围差异。[13] 1994年未进入回归、精确同比或转折判断。')
para('1995—1997年年鉴另有教育部门及其他部门学校收入表，其小学、初中记录已单独核准；该范围不包含企业及民办全部收入，因此正文全口径系列仍按支出来源组成重建，不能混加企业支出与部门收入。1998—2004年按收入来源重建，2005年起读取公布值。')
h('（二）2002年企业项与2007年来源组成')
para('1998—2001年企业总收入计入原公布财政性金额；企业项调整仅保留企业拨款及产业等用于教育的收入，扣除学杂费和其他收入。1998—2006年候选算法已回核，2002年以后与公布值相符。2001年小学调整为1011.73亿元、初中548.89亿元，较原值分别少11.63、7.92亿元。该调整只统一企业项，未认证全部跨期学校范围。')
para('2006、2007年全国及两学段原表均满足：全部收入＝部门学校收入＋企业学校全部收入＋民办学校全部收入。因此民办来源项目改为“举办者投入”不能解释为总经费分母缩小。2007年财政性分子包含企业、民办学校预算内资金及教育税费；其与沿用旧表可识别组成的差额，小学15.15亿元、初中15.20亿元，分别影响占比0.5138、0.7400个百分点。2006年未独立列项，不认定对应资金为零。[11]')
fig(DATA/'经费断点再核查_20261012/口径差异核查图.png','企业来源及公共资金列项差异的量化比较','原表组成项差额，非年度增长和抽样误差。2001年为企业项调整；2007年为完整财政性组成与旧表可识别算法之差。')
para('2007年小学公布金额增长34.36%，按旧表可识别组成增长33.60%；初中为48.46%和47.16%。这组差异不足以消除当年明显增长，但不能将其称为全部制度变化效应。2006年千元单位已与国家统计局全国合计交叉核准；1998年新会计制度调整的范围边界仍保留。[11][14]')
h(chapter_titles[11],1,True)
oldfig(6,'各省级地区小学与初中公用支出累计实际增长（2017—2023年）','按小学增长排序；线段连接同一地区两学段。零线左侧表示累计实际下降；与正文空间分布图使用同一数据。',5.8)
local=pd.read_csv(BASE/'正式学术图组_20261011/figure_data/fig07.csv')
piv=local.pivot(index='province',columns='stage',values='C_real_change_pct').sort_values('primary')
doc.add_page_break()
vals=[[p,f'{r.primary:.2f}',f'{r.junior:.2f}'] for p,r in piv.iterrows()]
paired_rows=[vals[j]+(vals[j+16] if j+16<len(vals) else ['','','']) for j in range(16)]
tab('分省公用支出累计实际增长（2017—2023年）',['省级地区','小学（%）','初中（%）','省级地区','小学（%）','初中（%）'],paired_rows,note='按小学增幅升序排列；先左列从上到下，再右列从上到下阅读。地方隶属、全来源支出，使用地区CPI调整；端点结果不表示逐年变化均同向。')
para('地图边界来自阿里云DataV.GeoAtlas公开省级GeoJSON，下载日期为2026年10月12日；保存原始文件、来源地址及SHA-256。图示保留原始省界、岛屿和海域线要素，采用主图及南海诸岛附图；两学段使用相同分组阈值，31个地区均与分析数据匹配。港澳台未纳入数据分析，呈灰色。该底图为公开地理数据；公开出版时，地图应按有关要求使用审核底图或完成地图审核。',True)
para('数据按原始资料、提取值、核准值、派生指标、模型输入和图表数据分层保存；修正不覆盖原值。2016年广东初二核准为1,140,343人，2006年全国教育经费单位恢复为千元。原表坐标、单位及文件哈希用于复核。金额加总核对与跨年范围认证分别记录。',True)
para('数值复算包括固定效应模型的独立OLS核算、组成恒等式、CPI转换、校正p值及置信区间反演；细项和资源回读保留原格坐标。支出与学生覆盖、省市财政范围、工资养老列支及专属成本指数的限制集中列于附录，不通过增加当期控制变量代替机制识别。',True)
h(chapter_titles[12],1,True)
refs=[
'[1] 中国教育统计年鉴：1995—2024年实际数据。全国及分省学生、教师、班级和学校统计表。',
'[2] 中国教育经费统计年鉴：1996—2025年卷。收入来源、举办者及支出明细表，分别记录卷年与实际数据年。',
'[3] 中国财政年鉴。地方一般预算及一般公共预算决算表：自身收入与中央补助。',
'[4] 中国统计年鉴及分地区居民消费价格指数表。地区CPI以上年为100。',
'[5] 梁文艳，孙雨婷，杜育红. 教育强国背景下适应人口变动的义务教育师资配置——基于人口增减分化典型省域的研究[J]. 教育研究，2026（2）：40—53。',
'[6] 蔺海沣，马玲. 人口结构变动对公共教育支出的影响效应——基于全球70项定量研究的三水平元分析[J]. 教育研究，2026（2）：54—68。',
'[7] 李涛，黄嘉欣. 学龄人口减少对乡村小规模学校发展的影响及其疏解——基于东北地区人口负增长省份的调查[J]. 教育研究，2025（5）：107—121。',
'[8] 陈坤，秦玉友. 人口变动趋势下基础教育师资配置机制研究——基于8省16县（区）的调查[J]. 西南民族大学学报（人文社科版），2026，47（5）：229—240。',
'[9] Cameron A C, Gelbach J B, Miller D L. Bootstrap-Based Improvements for Inference with Clustered Errors. Review of Economics and Statistics, 2008, 90(3): 414–427.',
'[10] 国家统计局. 中国统计年鉴2001：20-36表. https://www.stats.gov.cn/sj/ndsj/2001c/t2036c.htm',
'[11] 国家统计局. 中国统计年鉴2012：20-39表. https://www.stats.gov.cn/sj/ndsj/2012/html/U2039C.HTM',
'[12] 中国教育经费投入的“八五”回顾与“九五”展望：表1-4. https://www.cnsaes.org.cn/homepage/Upfile/201029/2010020981343441.pdf',
'[13] 国务院新闻办公室. 中国的儿童状况，1996. 外交部转载. https://www.fmprc.gov.cn/ziliao_674904/zt_674979/ywzt_675099/wzzt_675579/2296_675789/200207/t20020705_7961975.shtml',
'[14] 教育部，国家统计局，财政部. 1998年全国教育经费执行情况统计公告. https://hudong.moe.gov.cn/jyb_sjzl/sjzl_jfzxgg/201005/t20100527_88452.html',
'[15] 阿里云DataV.GeoAtlas. 省级行政区边界数据. https://geo.datav.aliyun.com/areas_v3/bound/100000_full.json（2026年10月12日下载）.'
]
policy=pd.read_csv(DATA/'adjustment_policy_evidence_ledger.csv')
for key,num in [('2015_guarantee',16),('2018_efficiency',17),('2019_responsibility',18)]:
    row=policy[policy.raw_path.str.contains(key)].iloc[0]
    assert hashlib.sha256((BASE/row.raw_path).read_bytes()).hexdigest()==row.sha256
    protected[str((BASE/row.raw_path).relative_to(ROOT))]=row.sha256
    refs.append(f'[{num}] '+row.title+'，'+row.document_number+'，'+row.document_date+'。'+row.url)
for r in refs:
    p=para(r,True);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
doc.save(OUT/(NAME+'.docx'))
assert all(hashlib.sha256((ROOT/p).read_bytes()).hexdigest()==v for p,v in protected.items())
md=[]
for b in blocks:
    if b['type']=='title':md.append('# '+b['text'])
    elif b['type']=='heading':md.append('#'*(b['level']+1)+' '+b['text'])
    elif b['type']=='paragraph':md.append(b['text'])
    elif b['type']=='figure':md.append(f'图{b["number"]}  {b["title"]}\n\n![{b["title"]}]({(ROOT/b["path"]).as_posix()})\n\n注：{b["caption"]}')
    else:md.append(f'表{b["number"]}  {b["title"]}\n\n| '+' | '.join(b['headers'])+' |\n| '+' | '.join(['---']*len(b['headers']))+' |\n'+'\n'.join('| '+' | '.join(str(v) for v in r)+' |' for r in b['rows']))
(OUT/'报告正文.md').write_text('\n\n'.join(md),encoding='utf8')
(OUT/'正式报告生成记录.json').write_text(json.dumps(dict(title=NAME,figures=figures,tables=tables,chapters=chapter_titles,input_sha256=protected,new_models=False,original_reports_preserved=True),ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(dict(figures=len(figures),tables=len(tables),path=str(OUT/(NAME+'.docx'))),ensure_ascii=False))

