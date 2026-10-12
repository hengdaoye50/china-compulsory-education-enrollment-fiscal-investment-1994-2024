from pathlib import Path
from copy import deepcopy
import zipfile, json, hashlib, re
from docx import Document
from docx.text.paragraph import Paragraph

NAME='中国义务教育生源规模与财政投入趋势分析报告（1994—2024）'
SRC=Path('<LOCAL_USER>/Desktop')/(NAME+'.docx')
OUT=Path.cwd()/'义务教育趋势分析报告_AI使用说明修订版_20261012'
OUT.mkdir(exist_ok=True)
d=Document(SRC)
hash_before=hashlib.sha256(SRC.read_bytes()).hexdigest()

def write(p,text):
    pr=deepcopy(p.runs[0]._r.rPr) if p.runs else None
    for r in list(p.runs):p._p.remove(r._r)
    r=p.add_run(text)
    if pr is not None:r._r.insert(0,pr)

# Delete only the policy summary table, its caption and its closing explanation.
caption=next(p for p in d.paragraphs if p.text.startswith('表5  资源配置'))
closing=next(p for p in d.paragraphs if p.text.startswith('报告据此将政策关注点'))
table=d.tables[4]
assert table.cell(0,0).text=='政策关注'
for element in [caption._p, table._tbl, closing._p]:element.getparent().remove(element)

# Shift appendix references and table numbers inside runs to retain superscripts/fonts.
allparas=list(d.paragraphs)
for t in d.tables:
    for row in t.rows:
        for cell in row.cells:allparas.extend(cell.paragraphs)
for p in allparas:
    for run in p.runs:
        text=run.text.replace('附录四','附录五').replace('附录三','附录四')
        text=re.sub(r'表\s*(\d+)',lambda m:'表'+str(int(m[1])-1) if int(m[1])>=6 else m[0],text)
        if text!=run.text:run.text=text

items=[
('Heading 1','附录三  AI工具使用说明'),
('Heading 2','（一）主要智能体与研究分工'),
('Normal','本报告主要使用OpenAI Codex AI agent，在本地工作区辅助资料检索、数据整理、分析程序编写、可视化和报告修订。研究团队提出问题、确定时段与对象，审阅阶段性结果，并就数据口径、图示表达和报告内容提出修订要求。AI agent依据这些要求执行具体任务；数据判断以可追溯原始资料和实际程序结果为依据。'),
('Heading 2','（二）技能与配套工具'),
('Normal','技能（skill）为智能体提供任务流程与审校规则。scientific-visualization用于图示设计和审校，检查变量、单位、统计范围、缺失值、分段及不确定性表达；matplotlib用于坐标轴、图例、中文字体、南开紫与金色配色以及静态图输出；statistical-analysis用于梳理检验对象、样本单位、估计区间、多重检验和结果解释。技能建议结合本报告的数据条件采用，未照搬不适用的检验流程。'),
('Normal','profyunv（于文轩学术指导知识库）用于审视问题意识、研究设计、论证和结论边界；academic-human-writing用于检查中文学术表达、证据与判断的对应关系及重复表述。前者属于基于知识库的AI辅助审校，不代表于文轩教授本人参与审阅或认可报告结论。'),
('Normal','documents文档技能及配套转换、渲染工具用于Word修订和PDF排版核验。Codex app工具（codex-app-tools）用于工作区依赖定位和文档展示，联网检索工具用于查找官方统计与政策材料。实际计算和制图由Python及pandas、NumPy、Matplotlib等程序库完成，文档处理使用python-docx等工具，PDF转换使用LibreOffice。程序库和转换软件承担计算、绘制与排版任务，与AI模型的文本生成功能分别记录。'),
('Heading 2','（三）AI辅助研究与可视化的具体流程'),
('Normal','资料阶段，智能体读取年鉴表格、PDF和网页，编写抽取与整理程序，建立分学段、分地区、分年份的数据系列；对异常数值、单位及统计分类变化回查原表，并通过官方网络材料交叉核对。早期经费按明确的来源组成或分配假设重建，条件估算、缺失与口径边界分别保留，未将模型生成的数值作为原始观测。'),
('Normal','分析阶段，智能体编写并执行趋势描述、CPI换算、收入构成分解和省级面板分析程序，比较支出类别的条件关联；进一步进行同样本比较、分期敏感性分析、聚类推断和多重检验校正。独立OLS复算、组成恒等式和原格回读用于核查计算。未获稳定支持的探索假设及拓展结果一并保留，未依据期望结论删改样本或把条件关联改写为因果效应。'),
('Normal','可视化阶段，智能体依据整理后的绘图数据编写Matplotlib程序，绘制全国时序、地区比较、支出构成、系数区间及省级地图。统计图由数据程序生成，未以文生图替代数据曲线或统计地图。图注注明来源、单位、范围、估算和区间含义；地图由省级边界文件与分析数据匹配生成。图示审校同时核对精确数值、文字可读性和Word中的实际显示效果。'),
('Normal','报告阶段，智能体据分析结果辅助起草和修订摘要、正文、结论与政策启示，整理参考文献和引文编号。每轮编辑以团队提供的版本为底稿，保留既有字体和设计，导出PDF后检查页面、图表位置及目录页码；团队反馈用于继续修订。'),
('Heading 2','（四）复核与可追溯材料'),
('Normal','原始资料、来源地址、抽取坐标、处理脚本、绘图数据和核验记录保存在项目工作区；部分输入与输出文件记录SHA-256，统计程序保留样本排除规则、检验参数及环境信息。AI生成的程序、文献著录和文字均需核对，技能或自动检查不能代替原始来源核准。研究团队负责最终审阅、结论取舍及对外提交，公开复现材料时应同步保留上述处理依据。'),
]
anchor=next(p for p in d.paragraphs if p.text=='附录四  早期经费与跨期衔接')
templates={s:next(p for p in d.paragraphs if p.style.name==s and (p.text.startswith('附录一') if s=='Heading 1' else True)) for s in ['Heading 1','Heading 2','Normal']}
templates['Normal']=next(p for p in d.paragraphs if p.text.startswith('个人支出记为'))
for style,text in items:
    element=deepcopy(templates[style]._p)
    anchor._p.addprevious(element)
    p=Paragraph(element,anchor._parent);write(p,text)
    if style=='Heading 2':p.paragraph_format.page_break_before=False

# Insert an entry into the existing contents list using its current typography.
tocanchor=next(p for p in d.paragraphs if p.text.startswith('附录四') and '\t' in p.text)
element=deepcopy(tocanchor._p);tocanchor._p.addprevious(element)
p=Paragraph(element,tocanchor._parent)
# Preserve the tab stop and separate page-number run needed for deterministic update.
for r in list(p.runs):p._p.remove(r._r)
props=deepcopy(tocanchor.runs[0]._r.rPr)
r=p.add_run('附录三  AI工具使用说明');
if props is not None:r._r.insert(0,props)
for text in ['\t','0']:
    r=p.add_run(text)
    if props is not None:r._r.insert(0,deepcopy(props))

dest=OUT/(NAME+'.docx');d.save(dest)
with zipfile.ZipFile(dest) as z:body=z.read('word/document.xml')
tmp=OUT/'preserved.docx'
with zipfile.ZipFile(SRC) as a,zipfile.ZipFile(tmp,'w',compression=zipfile.ZIP_DEFLATED) as b:
    for info in a.infolist():b.writestr(info,body if info.filename=='word/document.xml' else a.read(info.filename))
tmp.replace(dest)
assert hashlib.sha256(SRC.read_bytes()).hexdigest()==hash_before
(OUT/'AI工具使用说明.md').write_text('\n\n'.join(t for _,t in items),encoding='utf-8')
(OUT/'修订记录.json').write_text(json.dumps({'source':str(SRC),'source_sha256':hash_before,'deleted':'原表5、题头及其后总结说明','table_count':len(d.tables),'figure_numbering':'图号保持不变；原表6至15改为表5至14','appendices':'新增附录三，原附录三、四顺延为四、五，正文回指同步更新','original_package_preserved_except_document_body':True},ensure_ascii=False,indent=2),encoding='utf-8')
print(dest)
