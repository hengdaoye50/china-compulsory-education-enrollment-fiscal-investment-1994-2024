"""Reproducible source audit; no reconstructed spending or fiscal student counts.

Fujian numeric cells were read visually from rendered physical pages 11/18/25.
All automation here is same-analyst verification, not independent peer review.
"""
import csv, hashlib, io, json, re, sys
from decimal import Decimal as D
from pathlib import Path
import pdfplumber
from openpyxl import load_workbook
from lxml import html
from parse import ROOT, save_csv
sys.stdout.reconfigure(encoding='utf-8')
DATA=ROOT/'data'
def digest(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def read_csv(p): return list(csv.DictReader(p.open(encoding='utf-8-sig',newline='')))
def text(p):
 d=html.fromstring(p.read_bytes().decode('utf-8'))
 return ' '.join(' '.join(d.xpath('//text()[not(ancestor::script) and not(ancestor::style)]')).split())
protected=['data/finance_long.csv','data/education_resource_panel.csv','data/panel.json','教育资源试做.html']
baseline_path=DATA/'province_finance_next_protected_baseline.json'
if not baseline_path.exists():
 baseline_path.write_text(json.dumps({p:digest(ROOT/p) for p in protected},ensure_ascii=False,indent=2),encoding='utf-8')
baseline=json.loads(baseline_path.read_text(encoding='utf-8'))
assert all(digest(ROOT/p)==h for p,h in baseline.items())
sources=[]
for name in ['province_finance_next_sources','province_finance_next_attachment_sources','province_finance_next_rule_sources','province_finance_next_final_sources']:
 sources.extend(json.loads((ROOT/(name+'.json')).read_text(encoding='utf-8')))
bykey={r['key']:r for r in sources}
receipt=[]
for s in sources:
 row={k:s.get(k,'') for k in ['key','url','status','bytes','sha256','path','error']}
 row['evidence_type']='download attempt failed'
 if s['status']=='downloaded':
  p=ROOT/s['path']; b=p.read_bytes()
  assert len(b)==s['bytes'] and digest(p)==s['sha256']
  # A download endpoint suffix is not the file format.
  ext='.doc' if b.startswith(bytes.fromhex('d0cf11e0a1b11ae1')) else '.xlsx' if b.startswith(b'PK') else None
  if ext and p.suffix!=ext:
   q=p.with_suffix(ext)
   if not q.exists(): q.write_bytes(b)
   assert q.read_bytes()==b
   row['path']=str(q.relative_to(ROOT));row['original_endpoint_saved_path']=s['path']
  row['evidence_type']='official institutional republication' if '_cee' in s['key'] else 'university-hosted provincial notification copy' if 'reporting' in s['key'] else 'government original download'
 receipt.append(row)
save_csv(DATA/'province_finance_next_source_receipt.csv',receipt)

inventory=[]; export=[]
book_keys=['province_finance_next_sichuan_2023_xlsx','province_finance_next_hubei_2023_xlsx','province_finance_next_hunan_2022_attachment','province_finance_next_hunan_2023_attachment']
books={}
for key in book_keys:
 s=bykey[key];w=load_workbook(io.BytesIO((ROOT/s['path']).read_bytes()),data_only=True);books[key]=w
 for ws in w:
  occupied=[(c.coordinate,str(c.value)) for row in ws for c in row if c.value is not None]
  joined=' '.join(v for _,v in occupied)
  inventory.append(dict(source_key=key,sheet=ws.title,rows=ws.max_row,columns=ws.max_column,nonempty_cells=len(occupied),
   headers=' | '.join(f'{c.coordinate}={c.value}' for row in list(ws)[:4] for c in row if c.value is not None),
   student_field_hits=';'.join(sorted(set(re.findall(r'.{0,20}(?:学生数|学生人数|年初学生|年末学生|加权|权重).{0,25}',joined)))),
   stage_spending_total='not found in inspected sheets; stage columns are per-student',fiscal_student_denominator='not found in inspected sheets',weight_rule='not found in inspected sheets'))
  export.extend(dict(source_key=key,sheet=ws.title,cell=a,value=v) for a,v in occupied)
save_csv(DATA/'province_finance_next_sheet_inventory.csv',inventory)
save_csv(DATA/'province_finance_next_xlsx_cells.csv',export)

cells=[]
def add(province,year,stage,metric,value,growth,key,locator,method='direct original table'):
 cells.append(dict(province=province,year=year,stage=stage,metric=metric,value=str(value),published_growth_pct=str(growth),unit='元（原公告生均指标，学生分母未核准）',
  source=bykey[key]['url'] if key in bykey else key,source_key=key,locator=locator,method=method,
  coverage='province aggregate as labelled; stage denominator organizer scope not separately verified',publication_data_year=year,vintage='current-year column of currently available annual publication; identity with first release not separately verified',eligibility='published per-student only; not spending/student pair'))
metrics=['budget_per_student','career_per_student','public_per_student']
key='province_finance_next_sichuan_2023_xlsx';w=books[key]
for index,metric in [(1,metrics[0]),(3,metrics[1]),(5,metrics[2])]:
 ws=w.worksheets[index]
 assert ('生均' in ws['A1'].value) and '普通小学'==ws['D3'].value and '普通初中'==ws['F3'].value
 for stage,col in [('primary',4),('junior',6)]:add('四川',2023,stage,metric,ws.cell(5,col).value,ws.cell(5,col+1).value,key,f'{ws.title}!{ws.cell(5,col).coordinate}; headers A1/D3/F3')
key='province_finance_next_hubei_2023_xlsx';ws=books[key].worksheets[1]
assert '生均一般公共预算教育经费' in ws['A1'].value
for stage,col in [('primary',4),('junior',6)]: add('湖北',2023,stage,metrics[0],ws.cell(5,col).value,ws.cell(5,col+1).value,key,f'{ws.title}!{ws.cell(5,col).coordinate}; A1/D3/F3; body label conflict retained')
for year in [2022,2023]:
 key=f'province_finance_next_hunan_{year}_attachment';w=books[key]
 for i,metric in enumerate(metrics,1):
  ws=w.worksheets[i]
  assert str(year) in ws['F4'].value and str(year) in ws['I4'].value and ws['E3'].value=='普通小学' and ws['H3'].value=='普通初中'
  for stage,col in [('primary',6),('junior',9)]: add('湖南',year,stage,metric,ws.cell(5,col).value,ws.cell(5,col+1).value,key,f'{ws.title}!{ws.cell(5,col).coordinate}; E3/H3; F4/I4')

# DOC conversion puts column headings in preceding paragraphs, not table rows.
sd=ROOT/'converted/province_finance_next_shandong_2023_attachment.html'
doc=html.fromstring(sd.read_bytes().decode('utf-8'));tables=doc.xpath('//table');assert len(tables)==7
sd_pdf=sd.with_suffix('.pdf')
if not sd_pdf.exists():sd_pdf.write_bytes((ROOT/'.tmp/province_finance_next_shandong_2023_attachment.pdf').read_bytes())
with pdfplumber.open(sd_pdf) as pdf:
 sd_pages=[p.extract_text() or '' for p in pdf.pages]
sd_headers=[]
for index,metric,caption in [(1,metrics[0],'生均一般公共预算教育经费'),(3,metrics[1],'生均一般公共预算教育事业费'),(5,metrics[2],'生均一般公共预算公用经费')]:
 # Word floating frames lose some column headings in HTML. Use a separate PDF
 # conversion of the same original for captions/stage labels and cross-read row.
 pnum,page=next((i+1,t) for i,t in enumerate(sd_pages) if caption in re.sub(r'\s','',t) and '普通小学' in t and re.search(r'山东省\s+[\d,]',t))
 header=page.split('山东省')[0]
 assert '普通小学' in header and '普通初中' in header
 sd_headers.append(dict(table_index=index,metric=metric,converted_pdf_physical_page=pnum,headers=header))
 t=tables[index]
 rows=t.xpath('./tr|./tbody/tr');values=[' '.join(c.text_content().split()) for c in rows[0].xpath('./td|./th')]
 assert values[0]=='山东省'
 pdf_values=next(l for l in page.splitlines() if l.startswith('山东省 ')).split()
 assert values==pdf_values
 for stage,col in [('primary',3),('junior',5)]:add('山东',2023,stage,metric,values[col],values[col+1],'province_finance_next_shandong_2023_attachment',f'converted table {index}, first row column {col}; caption in preceding paragraphs','official Word converted with LibreOffice; government file signature verified')
save_csv(DATA/'province_finance_next_shandong_headers.csv',sd_headers)

prior=json.loads((ROOT/'province_finance_attachment_sources.json').read_text(encoding='utf-8'))+json.loads((ROOT/'province_finance_followup_sources.json').read_text(encoding='utf-8'))+json.loads((ROOT/'province_finance_deep_sources.json').read_text(encoding='utf-8'))+json.loads((ROOT/'province_finance_sources.json').read_text(encoding='utf-8'))
old={s['key']:s for s in prior}
for key in ['province_finance_fujian_2023_pdf','province_finance_jilin_2023_pdf']:
 assert digest(ROOT/old[key]['path'])==old[key]['sha256']
jl=ROOT/'raw/province_finance_jilin_2023_pdf.pdf'
with pdfplumber.open(jl) as pdf:
 pages=[p.extract_text() or '' for p in pdf.pages]
for caption,metric in [('表二',metrics[0]),('表三',metrics[1]),('表四',metrics[2])]:
 pnum,txt=next((i+1,t) for i,t in enumerate(pages) if caption in t and re.search(r'吉林省\s+[\d,]',t))
 line=next(l for l in txt.splitlines() if re.match(r'吉林省\s+[\d,]',l));values=line.split()[1:]
 assert '普通小学' in txt and '普通初中' in txt
 for stage,col in [('primary',2),('junior',4)]:add('吉林',2023,stage,metric,values[col].replace(',',''),values[col+1],old['province_finance_jilin_2023_pdf']['url'],f'physical PDF page {pnum}; {caption}; 吉林省 row','PDF text extraction with caption and stage headers checked')

fj_key='province_finance_fujian_2023_pdf';fj=old[fj_key]
visual=[]
for page,metric,values,growths in [(11,metrics[0],['12282.27','19396.20'],['-0.29','0.28']),(18,metrics[1],['11905.67','18706.59'],['-0.39','0.15']),(25,metrics[2],['2399.83','3842.90'],['-6.00','1.10'])]:
 for stage,value,growth in zip(['primary','junior'],values,growths):
  add('福建',2023,stage,metric,value,growth,fj['url'],f'physical page {page} / printed page {page+4}; 福建省 row; 普通小学/普通初中 columns','manual visual transcription from 2x raster; not OCR-derived')
  visual.append(dict(physical_page=page,printed_page=page+4,stage=stage,metric=metric,value=value,growth=growth,source=fj['url'],source_sha256=fj['sha256'],method='manual visual; full 32-page overview plus enlarged first-page headers; not an independent review'))
save_csv(DATA/'province_finance_next_fujian_visual_cells.csv',visual)
save_csv(DATA/'province_finance_next_candidate_cells.csv',cells)

finance=read_csv(DATA/'finance_long.csv');national={(r['province'],r['year'],r['stage'],r['metric']):r for r in finance}
checks=[]
for c in cells:
 n=national[(c['province'],str(c['year']),c['stage'],c['metric'])]
 for field,tol in [('value',D('0.005')),('published_growth_pct',D('0.005'))]:
  diff=D(c[field])-D(n[field]);checks.append(dict(province=c['province'],year=c['year'],stage=c['stage'],metric=c['metric'],field=field,local_value=c[field],national_value=n[field],difference=str(diff),status='pass' if abs(diff)<=tol else 'unresolved',source=c['source'],locator=c['locator']))
save_csv(DATA/'province_finance_next_cross_source_checks.csv',checks)
continuity=[]
for stage in ['primary','junior']:
 for i,metric in enumerate(metrics,1):
  previous=next(c for c in cells if c['province']=='湖南' and c['year']==2022 and c['stage']==stage and c['metric']==metric)
  current=next(c for c in cells if c['province']=='湖南' and c['year']==2023 and c['stage']==stage and c['metric']==metric)
  col=5 if stage=='primary' else 8
  reference=str(books['province_finance_next_hunan_2023_attachment'].worksheets[i].cell(5,col).value)
  calc=(D(current['value'])/D(previous['value'])-1)*100
  continuity.append(dict(province='湖南',year=2023,stage=stage,metric=metric,prior_original_value=previous['value'],current_prior_column=reference,prior_status='pass' if D(reference)==D(previous['value']) else 'unresolved',computed_growth=str(calc),published_growth=current['published_growth_pct'],growth_status='pass' if abs(calc-D(current['published_growth_pct']))<=D('0.005') else 'unresolved',eligibility='per-student vintage consistency only; no fiscal denominator'))
save_csv(DATA/'province_finance_next_hunan_continuity_checks.csv',continuity)

scope=[
 dict(province='福建',year=2023,locator='PDF physical 1-5 / 6-10 / 11-17 / 18-24 / 25-32',finding='32 pages, five attachments: all-stage budget amount (千元), education surcharge input, three per-student families (元). No stage spending total/student-count/weight column seen in reviewed table forms.',boundary='Visual field audit of this attachment only; no claim about all Fujian public sources',source=fj['url']),
 dict(province='湖北',year=2023,locator='notice II(二) heading/body; XLSX 表二 A1,D3,F3,D5,F5',finding='Notice heading says budget education per-student, following sentence says career spending per-student. XLSX caption and national budget values match 12266.01/18936.42; body-label conflict retained.',boundary='Do not silently relabel the original or treat this as a career-spending observation; no main panel correction warranted',source=bykey['province_finance_next_hubei_2023']['url']),
 dict(province='四川',year=2023,locator='notice II(一); XLSX table 1 B4',finding='2015.95亿元 provincial total and 1929.10亿元 local total are distinct; table reports local 1929.1024145亿元.',boundary='Do not substitute totals across coverage or use either for primary/junior totals',source=bykey['province_finance_next_sichuan_2023']['url']),
 dict(province='湖南',year=2023,locator='notification physical pages 1-6; especially II-III and attachment p6',finding='2024 reporting-work notification collects 2023 annual data; includes 205 reconciliation table and broad survey organizer scope; no student-weight coefficients.',boundary='University-hosted notification copy; broad survey objects are not proof of coverage of each per-student monitoring indicator; do not assign annual observations from notification title year',source=bykey['province_finance_next_hunan_reporting_2023_pdf_http']['url']),
]
save_csv(DATA/'province_finance_next_scope_evidence.csv',scope)
bridges=[]
for page_key,attachment_key,token,basis in [
 ('province_finance_next_sichuan_2023','province_finance_next_sichuan_2023_xlsx','2023年全省教育经费执行情况统计表-2024.12.23.xlsx','original relative link; percent-encoded filename'),
 ('province_finance_next_hubei_2023','province_finance_next_hubei_2023_xlsx','P020241226335741972813.xlsx','original attachment link'),
 ('province_finance_next_shandong_2023_cee','province_finance_next_shandong_2023_attachment','5469dddcc5714edc9bc17a676a59e6ff.doc','official institutional republication points to government original'),
 ('province_finance_next_hunan_2022_http','province_finance_next_hunan_2022_attachment','ID=118281','original legacy URL plus same page script mapping govnew.hnedu.cn:8090 to cms.jyt.hunan.gov.cn'),
 ('province_finance_next_hunan_2023_http','province_finance_next_hunan_2023_attachment','ID=125566','original attachment link'),
]:
 assert token in (ROOT/bykey[page_key]['path']).read_bytes().decode('utf-8')
 bridges.append(dict(page_source=bykey[page_key]['url'],attachment_source=bykey[attachment_key]['url'],token=token,basis=basis,status='pass'))
save_csv(DATA/'province_finance_next_publication_bridge.csv',bridges)
coverage=[]
for province,year,depth in [('四川',2023,'official notice + 7-sheet original'),('湖北',2023,'official notice + 2-sheet original; body label conflict'),('山东',2023,'official institutional republication + government Word original; 7 converted tables'),('湖南',2022,'official HTTP notice + 5-sheet original'),('湖南',2023,'official HTTP notice + 5-sheet original'),('吉林',2023,'existing official PDF; full text field inventory'),('福建',2023,'existing 32-page official PDF; completed visual field audit')]:
 coverage.append(dict(province=province,year=year,local_source_depth=depth,stage_actual_spending_verified=False,paired_fiscal_students_verified=False,student_weights_verified=False,decision='published per-student only; formal pairs still zero',coverage_boundary='limited to listed local sources, not all province-year publications'))
save_csv(DATA/'province_finance_next_coverage.csv',coverage)
registry=read_csv(DATA/'province_finance_2017_2023_registry.csv')
for r in registry:
 match=next((c for c in coverage if c['province']==r['province'] and str(c['year'])==r['year']),None)
 if match:
  r['local_source_depth']=match['local_source_depth'];r['note']=(r['note']+'; next-round field audit; see province_finance_next_coverage.csv').strip('; ')
save_csv(DATA/'province_finance_next_2017_2023_registry.csv',registry)
for a,b in [('province_finance_next_sichuan_2023_xlsx','province_finance_next_sichuan_2023_cee_xlsx'),('province_finance_next_hubei_2023_xlsx','province_finance_next_hubei_2023_cee_xlsx')]:assert bykey[a]['sha256']==bykey[b]['sha256']
assert all(digest(ROOT/p)==h for p,h in baseline.items())
audit=dict(download_records=len(sources),saved_download_records=sum(r['status']=='downloaded' for r in sources),unique_saved_contents=len({r['sha256'] for r in sources if r['status']=='downloaded'}),failed_download_records=sum(r['status']!='downloaded' for r in sources),verified_attachment_publication_bridges=len(bridges),new_local_province_years_deepened=6,previous_fujian_visual_depth_updated=1,original_workbook_sheets=len(inventory),original_word_tables=len(tables),fujian_pages_visually_reviewed=32,fujian_enlarged_numeric_cells=len(visual),candidate_cells=len(cells),cross_source_checks=len(checks),cross_source_pass=sum(c['status']=='pass' for c in checks),cross_source_unresolved=[c for c in checks if c['status']!='pass'],hunan_prior_value_checks=len(continuity),hunan_growth_checks=len(continuity),hunan_continuity_all_pass=all(c['prior_status']=='pass' and c['growth_status']=='pass' for c in continuity),registry_rows=len(registry),matched_continuous_stage_spending_student_pairs=0,complete_student_weight_rule_found=False,all_provinces_fully_screened=False,protected_files_unchanged=True,protected_hashes=baseline,review_type='same analyst, alternate source/file/arithmetical verification; not independent peer review')
(DATA/'province_finance_next_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(audit,ensure_ascii=False,indent=2))
